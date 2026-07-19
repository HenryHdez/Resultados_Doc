# -*- coding: utf-8 -*-
"""
Versión con:
1) Verificación estricta de coordenadas dentro del espacio de búsqueda (clamping).
2) Opción de límites personalizados (Xmin/Xmax/Ymin/Ymax) desde la interfaz.
3) Todo el pipeline (mapa de humedad, algoritmo, waypoints, gráfica, exportación)
   usa los límites configurados (no valores fijos).

Requisitos:
pip install pyserial matplotlib openpyxl
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import threading
import time
import math
import random
import json
import os
from datetime import datetime

import serial
import serial.tools.list_ports

from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.ticker import MultipleLocator
from matplotlib.patches import Rectangle

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


# =========================
# Parámetros base (servos)
# =========================
HOME_DX = 512
HOME_LX = 95

DX_MIN, DX_MAX = 0, 1023
LX_MIN, LX_MAX = 0, 240

LX_PER_DX = 0.22
PHI_H_DEG_DEFAULT = 90.0


# =========================
# Utilidades
# =========================
def list_serial_ports():
    return [p.device for p in serial.tools.list_ports.comports()]


def clamp_int(x: float, lo: int, hi: int) -> int:
    return max(lo, min(int(round(x)), hi))


def clamp_float(x: float, lo: float, hi: float) -> float:
    return max(lo, min(float(x), hi))


def deg2rad(deg: float) -> float:
    return deg * math.pi / 180.0


def wrap_pi(a: float) -> float:
    while a > math.pi:
        a -= 2.0 * math.pi
    while a < -math.pi:
        a += 2.0 * math.pi
    return a


def theta_from_alignment(aln: str) -> float:
    if aln == "+X":
        return 0.0
    if aln == "+Y":
        return math.pi / 2.0
    if aln == "-X":
        return math.pi
    if aln == "-Y":
        return -math.pi / 2.0
    return 0.0


def parse_id_list(s: str):
    s = s.strip()
    if not s:
        return []
    out = []
    for part in s.split(","):
        p = part.strip()
        if not p:
            continue
        out.append(int(p))
    return out


def build_manhattan_waypoints(ax, ay, bx, by, step_m, x_first=True, bounds=(0.0, 1.0, 0.0, 1.0)):
    """
    Waypoints Manhattan desde A a B, respetando límites (clamp en cada punto).
    bounds = (X_MIN, X_MAX, Y_MIN, Y_MAX)
    """
    X_MIN, X_MAX, Y_MIN, Y_MAX = bounds

    def clip(x, y):
        return (clamp_float(x, X_MIN, X_MAX), clamp_float(y, Y_MIN, Y_MAX))

    ax, ay = clip(ax, ay)
    bx, by = clip(bx, by)

    if step_m <= 0:
        step_m = 0.01

    wps = []
    x, y = ax, ay

    def walk_axis(cur, target):
        if abs(target - cur) < 1e-12:
            return []
        sgn = 1.0 if target > cur else -1.0
        pts = []
        v = cur
        while True:
            nxt = v + sgn * step_m
            if (sgn > 0 and nxt >= target) or (sgn < 0 and nxt <= target):
                pts.append(target)
                break
            pts.append(nxt)
            v = nxt
        return pts

    if x_first:
        for nx in walk_axis(x, bx):
            x, y = clip(nx, y)
            wps.append((x, y))
        for ny in walk_axis(y, by):
            x, y = clip(x, ny)
            wps.append((x, y))
    else:
        for ny in walk_axis(y, by):
            x, y = clip(x, ny)
            wps.append((x, y))
        for nx in walk_axis(x, bx):
            x, y = clip(nx, y)
            wps.append((x, y))

    if not wps or (wps[-1][0] != bx or wps[-1][1] != by):
        wps.append((bx, by))
    return wps


# =========================
# Morfologías (dibujo)
# =========================
def morphology_positions(name: str, n: int):
    n = max(1, int(n))
    pts = []

    if name == "Serpiente":
        for i in range(n):
            pts.append((i, 0))
        return pts

    if name == "L":
        k = max(2, min(n, n // 2 + 1))
        for i in range(k):
            pts.append((i, 0))
        x0, y0 = pts[-1]
        for j in range(1, n - k + 1):
            pts.append((x0, -j))
        return pts

    if name == "T":
        tv = max(2, min(n, n // 2 + 1))
        th = n - tv
        for j in range(tv):
            pts.append((0, -j))
        if th <= 0:
            return pts
        used = 0
        step = 1
        while used < th:
            pts.append((step, 0))
            used += 1
            if used >= th:
                break
            pts.append((-step, 0))
            used += 1
            step += 1
        return pts[:n]

    if name == "Cuadrúpedo":
        pts.append((0, 0))
        if n == 1:
            return pts
        legs = [(1, 0), (-1, 0), (0, 1), (0, -1)]
        i = 0
        while len(pts) < n and i < len(legs):
            pts.append(legs[i])
            i += 1
        if len(pts) >= n:
            return pts[:n]
        base = [(1, 0), (-1, 0), (0, 1), (0, -1)]
        ext_dirs = [(1, 0), (-1, 0), (0, 1), (0, -1)]
        ext_len = 2
        while len(pts) < n:
            for b, d in zip(base, ext_dirs):
                if len(pts) >= n:
                    break
                bx, by = b
                dx, dy = d
                pts.append((bx + dx * (ext_len - 1), by + dy * (ext_len - 1)))
            ext_len += 1
        return pts[:n]

    for i in range(n):
        pts.append((i, 0))
    return pts


# =========================
# Mapa de humedad (grid) + interpolación
# =========================
class HumidityMap:
    def __init__(self):
        self.nx = 60
        self.ny = 40
        self.seed = 1
        self.noise = 0.05

        self.peak_x = 1.20
        self.peak_y = 0.70
        self.peak_sigma = 0.18

        self.base_ax = 0.10
        self.base_ay = 0.05

        self.xs = []
        self.ys = []
        self.H = []
        self._generated = False

        # límites actuales del mapa
        self.bounds = (0.0, 1.90, 0.0, 1.20)

    def generate(self, nx, ny, seed, noise, peak_x, peak_y, peak_sigma, bounds):
        self.bounds = bounds
        X_MIN, X_MAX, Y_MIN, Y_MAX = bounds

        self.nx = max(10, min(int(nx), 500))
        self.ny = max(10, min(int(ny), 500))
        self.seed = int(seed)
        self.noise = max(0.0, min(float(noise), 1.0))

        self.peak_x = clamp_float(float(peak_x), X_MIN, X_MAX)
        self.peak_y = clamp_float(float(peak_y), Y_MIN, Y_MAX)
        self.peak_sigma = max(0.03, min(float(peak_sigma), 0.80))

        random.seed(self.seed)

        self.xs = [X_MIN + (X_MAX - X_MIN) * i / (self.nx - 1) for i in range(self.nx)]
        self.ys = [Y_MIN + (Y_MAX - Y_MIN) * j / (self.ny - 1) for j in range(self.ny)]

        H = []
        maxv = -1e9
        minv = 1e9

        for y in self.ys:
            row = []
            for x in self.xs:
                dx = (x - self.peak_x)
                dy = (y - self.peak_y)
                g = math.exp(-(dx * dx + dy * dy) / (2.0 * self.peak_sigma * self.peak_sigma))

                base = self.base_ax * (x - X_MIN) / (X_MAX - X_MIN) + self.base_ay * (y - Y_MIN) / (Y_MAX - Y_MIN)
                val = 0.15 + 0.80 * g + 0.20 * base
                val += self.noise * 0.08 * (2.0 * random.random() - 1.0)

                row.append(val)
                maxv = max(maxv, val)
                minv = min(minv, val)
            H.append(row)

        if maxv - minv < 1e-12:
            maxv = minv + 1e-12
        for j in range(self.ny):
            for i in range(self.nx):
                H[j][i] = (H[j][i] - minv) / (maxv - minv)

        self.H = H
        self._generated = True

    def is_ready(self):
        return self._generated and self.H and self.xs and self.ys

    def value_bilinear(self, x, y):
        if not self.is_ready():
            return 0.0

        X_MIN, X_MAX, Y_MIN, Y_MAX = self.bounds

        x = clamp_float(x, X_MIN, X_MAX)
        y = clamp_float(y, Y_MIN, Y_MAX)

        nx = self.nx
        ny = self.ny

        fx = (x - X_MIN) / (X_MAX - X_MIN) * (nx - 1)
        fy = (y - Y_MIN) / (Y_MAX - Y_MIN) * (ny - 1)
        i0 = int(math.floor(fx))
        j0 = int(math.floor(fy))
        i1 = min(i0 + 1, nx - 1)
        j1 = min(j0 + 1, ny - 1)
        i0 = max(0, min(i0, nx - 1))
        j0 = max(0, min(j0, ny - 1))

        tx = fx - i0
        ty = fy - j0

        h00 = self.H[j0][i0]
        h10 = self.H[j0][i1]
        h01 = self.H[j1][i0]
        h11 = self.H[j1][i1]

        hx0 = h00 * (1 - tx) + h10 * tx
        hx1 = h01 * (1 - tx) + h11 * tx
        hxy = hx0 * (1 - ty) + hx1 * ty
        return max(0.0, min(1.0, hxy))

    def argmax_grid(self):
        if not self.is_ready():
            return None
        best = (-1.0, None, None)
        for j, y in enumerate(self.ys):
            for i, x in enumerate(self.xs):
                v = self.H[j][i]
                if v > best[0]:
                    best = (v, x, y)
        return (best[1], best[2], best[0])


# ======================================================
# Ventana de gráficas (pestañas)
# ======================================================
class PlotWindow(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Gráficas")
        self.geometry("1180x780")

        self.nb = ttk.Notebook(self)
        self.nb.pack(fill="both", expand=True)

        self.tab_morph = ttk.Frame(self.nb)
        self.tab_map = ttk.Frame(self.nb)
        self.nb.add(self.tab_morph, text="Morfología")
        self.nb.add(self.tab_map, text="Humedad + Recorrido")

        self.fig_m = Figure(figsize=(8.5, 6.5), dpi=100)
        self.ax_m = self.fig_m.add_subplot(111)
        self.ax_m.set_aspect("equal", adjustable="box")
        self.ax_m.grid(True)
        self.canvas_m = FigureCanvasTkAgg(self.fig_m, master=self.tab_morph)
        self.canvas_m.get_tk_widget().pack(fill="both", expand=True)

        self.fig_h = Figure(figsize=(8.5, 6.5), dpi=100)
        self.ax_h = self.fig_h.add_subplot(111)
        self.ax_h.set_xlabel("X (m)")
        self.ax_h.set_ylabel("Y (m)")
        self.ax_h.grid(True)
        self.canvas_h = FigureCanvasTkAgg(self.fig_h, master=self.tab_map)
        self.canvas_h.get_tk_widget().pack(fill="both", expand=True)

        self.im = None
        self.plan_line, = self.ax_h.plot([], [], linestyle="--", alpha=0.6)
        self.traj_line, = self.ax_h.plot([], [], linewidth=2.0)
        self.robot_tip, = self.ax_h.plot([], [], linestyle="None", marker=(3, 0, 0), markersize=16)
        self.pt_goal, = self.ax_h.plot([], [], marker="x", markersize=10)
        self.pt_best, = self.ax_h.plot([], [], marker="*", markersize=14)
        self.pt_true, = self.ax_h.plot([], [], marker="o", markersize=6)
        self.txt = self.ax_h.text(0.02, 0.98, "", transform=self.ax_h.transAxes, ha="left", va="top", fontsize=10)

        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self._alive = True
        self.after(150, self._refresh)

    def _on_close(self):
        self._alive = False
        self.destroy()

    def _apply_fixed_axes(self):
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.app.get_bounds()

        try:
            tick = float(self.app.tick_var.get())
            if tick <= 0:
                tick = (X_MAX - X_MIN) / 10.0
        except Exception:
            tick = (X_MAX - X_MIN) / 10.0

        self.ax_h.set_xlim(X_MIN, X_MAX)
        self.ax_h.set_ylim(Y_MIN, Y_MAX)
        self.ax_h.xaxis.set_major_locator(MultipleLocator(tick))
        self.ax_h.yaxis.set_major_locator(MultipleLocator(tick))

    def _draw_morphology(self):
        self.ax_m.clear()
        self.ax_m.set_title("Morfología (H/V por módulo)")
        self.ax_m.grid(True)
        self.ax_m.set_aspect("equal", adjustable="box")

        ids = self.app.ids_runtime[:]
        pts = self.app.morph_positions[:]
        if not ids or not pts or len(ids) != len(pts):
            self.canvas_m.draw_idle()
            return

        xs = [p[0] for p in pts]
        ys = [p[1] for p in pts]
        xmin, xmax = min(xs) - 1, max(xs) + 1
        ymin, ymax = min(ys) - 1, max(ys) + 1
        self.ax_m.set_xlim(xmin, xmax)
        self.ax_m.set_ylim(ymin, ymax)

        for rid, (cx, cy) in zip(ids, pts):
            is_h = self.app.horiz_var_by_id.get(rid, tk.BooleanVar(value=False)).get()
            face = "#2b7bff" if is_h else "#ff9b2b"
            r = Rectangle((cx - 0.45, cy - 0.30), 0.90, 0.60, facecolor=face, edgecolor="black", linewidth=1.5)
            self.ax_m.add_patch(r)
            self.ax_m.text(cx, cy, f"ID {rid}\n{'H' if is_h else 'V'}", ha="center", va="center", fontsize=9, color="black")

        self.ax_m.text(0.02, 0.98, f"N módulos: {len(ids)}", transform=self.ax_m.transAxes, ha="left", va="top", fontsize=10)
        self.canvas_m.draw_idle()

    def _draw_humidity(self):
        hm = self.app.hmap
        if not hm.is_ready():
            return

        X_MIN, X_MAX, Y_MIN, Y_MAX = self.app.get_bounds()

        if self.im is None:
            self.im = self.ax_h.imshow(
                hm.H,
                origin="lower",
                extent=(X_MIN, X_MAX, Y_MIN, Y_MAX),
                aspect="auto",
                interpolation="nearest",
                alpha=0.85,
                cmap="jet",
                vmin=0.0, vmax=1.0
            )
            self.fig_h.colorbar(self.im, ax=self.ax_h, fraction=0.046, pad=0.04).set_label("Humedad (0..1)")
        else:
            self.im.set_data(hm.H)
            self.im.set_extent((X_MIN, X_MAX, Y_MIN, Y_MAX))

        tg = hm.argmax_grid()
        if tg:
            xg, yg, _ = tg
            self.pt_true.set_data([xg], [yg])
        else:
            self.pt_true.set_data([], [])

        show_route = (self.app.show_route_var.get() == "Mostrar")
        show_goal = (self.app.show_goal_var.get() == "Mostrar")

        wps = self.app.waypoints_runtime[:]
        if show_route and wps:
            ax0, ay0 = self.app.start_A
            xs_plan = [ax0] + [p[0] for p in wps]
            ys_plan = [ay0] + [p[1] for p in wps]
            self.plan_line.set_data(xs_plan, ys_plan)
        else:
            self.plan_line.set_data([], [])

        if show_route:
            self.traj_line.set_data(self.app.traj_x[:], self.app.traj_y[:])
        else:
            self.traj_line.set_data([], [])

        x_r, y_r, th = self.app.pose_runtime
        ang_deg = math.degrees(th) - 90.0
        self.robot_tip.set_data([x_r], [y_r])
        self.robot_tip.set_marker((3, 0, ang_deg))

        gx, gy = self.app.current_goal
        if show_goal and gx is not None:
            self.pt_goal.set_data([gx], [gy])
        else:
            self.pt_goal.set_data([], [])

        bx, by, bv = self.app.best_found
        if bx is not None:
            self.pt_best.set_data([bx], [by])
        else:
            self.pt_best.set_data([], [])

        algo = self.app.algo_var.get()
        it = self.app.search_iter
        dist = self.app.distance_m
        vavg = self.app.robot_avg_speed_m_s()

        if bx is not None:
            txt = (
                f"Algoritmo: {algo}\n"
                f"Iteración: {it}\n"
                f"Mejor h={bv:.3f} en ({bx:.2f},{by:.2f})\n"
                f"Distancia acumulada: {dist:.3f} m\n"
                f"Velocidad promedio robot: {vavg:.4f} m/s"
            )
        else:
            txt = (
                f"Algoritmo: {algo}\n"
                f"Distancia acumulada: {dist:.3f} m\n"
                f"Velocidad promedio robot: {vavg:.4f} m/s"
            )

        self.txt.set_text(txt)
        self._apply_fixed_axes()
        self.canvas_h.draw_idle()

    def _refresh(self):
        if not self._alive:
            return
        try:
            self._draw_morphology()
            self._draw_humidity()
        except Exception:
            pass
        finally:
            self.after(150, self._refresh)


# ======================================================
# Hilo: movimiento hacia un objetivo (waypoints)
# ======================================================
class NavigateThread(threading.Thread):
    def __init__(self, app, stop_event):
        super().__init__(daemon=True)
        self.app = app
        self.stop_event = stop_event
        self.last_t = time.perf_counter()

    def run(self):
        while not self.stop_event.is_set():
            try:
                if self.app.require_serial() and (not self.app.ser or not self.app.ser.is_open):
                    time.sleep(0.2)
                    continue

                freq = float(self.app.freq_var.get())
                rate = float(self.app.rate_var.get())
                if freq <= 0 or rate <= 0:
                    time.sleep(0.1)
                    continue

                tol = float(self.app.tol_var.get())
                dead = math.radians(float(self.app.deadband_deg_var.get()))

                now = time.perf_counter()
                dt = now - self.last_t
                if dt <= 0:
                    dt = 1.0 / max(1.0, rate)
                self.last_t = now

                kv = float(self.app.kv_var.get())
                kw = float(self.app.kw_var.get())

                if self.app.current_goal[0] is None:
                    g = self.app.algo_next_goal()
                    if g is None:
                        self.app.set_status("Búsqueda finalizada (sin nuevo objetivo).")
                        self.stop_event.set()
                        break
                    self.app.set_new_goal(gx=g[0], gy=g[1])
                    self.app.set_status("Objetivo generado por algoritmo.")
                    time.sleep(0.01)
                    continue

                wp = self.app.current_wp()
                if wp is None:
                    self.app.on_reach_goal()
                    continue

                tx_r, ty_r = wp
                x_r, y_r, th = self.app.pose_runtime

                dist = math.hypot(tx_r - x_r, ty_r - y_r)
                if dist <= tol:
                    self.app.advance_wp()
                    time.sleep(0.005)
                    continue

                desired = math.atan2(ty_r - y_r, tx_r - x_r)
                err = wrap_pi(desired - th)

                turning = abs(err) > dead
                cw = err < 0

                if turning:
                    Av = float(self.app.turn_v_ticks.get())
                    Ah = float(self.app.turn_h_ticks.get())
                    phi_v = float(self.app.turn_phase_v_deg.get())
                    phi_h = float(self.app.turn_phase_h_deg.get())
                else:
                    Av = float(self.app.move_v_ticks.get())
                    Ah = float(self.app.move_h_ticks.get())
                    phi_v = float(self.app.move_phase_v_deg.get())
                    phi_h = float(self.app.move_phase_h_deg.get())

                shift_enable = bool(self.app.turn_phase_shift_enable.get())
                shift_dir = self.app.turn_phase_shift_dir.get().strip()
                shift_deg = float(self.app.turn_phase_shift_deg.get())
                extra_phi_h = 0.0
                if turning and shift_enable:
                    if shift_dir == "CW" and cw:
                        extra_phi_h = shift_deg
                    elif shift_dir == "CCW" and (not cw):
                        extra_phi_h = shift_deg

                t_chain = time.perf_counter() - self.app.t0
                lines, cmd_positions = self.app.build_servo_lines_with_positions(
                    t=t_chain, freq=freq,
                    Av=Av, Ah=Ah,
                    phi_v_deg=phi_v,
                    phi_h_deg=phi_h + extra_phi_h,
                    turning=turning, cw=cw
                )

                self.app.update_speed_history(cmd_positions=cmd_positions, dt=dt)

                if lines and self.app.require_serial():
                    self.app.send_lines_batch(lines)

                self.app.update_pose_model(Av=Av, Ah=Ah, turning=turning, cw=cw, kv=kv, kw=kw, dt=dt)
                self.app.sample_humidity_at_pose()

                self.app.set_status("Navegación + búsqueda activa.")
                time.sleep(1.0 / rate)

            except Exception as e:
                self.app.set_status(f"Detenido por error: {e}")
                self.stop_event.set()


# ======================================================
# Núcleo de algoritmos de búsqueda
# ======================================================
class SearchController:
    def __init__(self, app):
        self.app = app
        self.reset()

    def reset(self):
        self.algo = None
        self.iter = 0

        self.hc_step = 0.10

        self.ga_pop = []
        self.ga_pop_size = 18
        self.ga_elite = 4
        self.ga_mut_sigma = 0.12
        self.ga_gen = 0

        self.cg_k = 20
        self.cg_radius = 0.35

        self.bc_dir = None
        self.bc_run_len = 0
        self.bc_max_run = 10
        self.bc_step = 0.08
        self.bc_last_h = None

        self.sa_T = 1.0
        self.sa_alpha = 0.92
        self.sa_step = 0.18
        self.sa_current = None
        self.sa_current_h = None

        self.rw_step = 0.12

    def set_algo(self, algo_name: str):
        self.reset()
        self.algo = algo_name
        self.iter = 0

        self.hc_step = float(self.app.hc_step_var.get())
        self.ga_pop_size = int(self.app.ga_pop_var.get())
        self.ga_elite = int(self.app.ga_elite_var.get())
        self.ga_mut_sigma = float(self.app.ga_mut_var.get())
        self.cg_k = int(self.app.cg_k_var.get())
        self.cg_radius = float(self.app.cg_radius_var.get())
        self.bc_step = float(self.app.bc_step_var.get())
        self.bc_max_run = int(self.app.bc_maxrun_var.get())
        self.sa_T = float(self.app.sa_T0_var.get())
        self.sa_alpha = float(self.app.sa_alpha_var.get())
        self.sa_step = float(self.app.sa_step_var.get())
        self.rw_step = float(self.app.rw_step_var.get())

        x0, y0, _ = self.app.pose_runtime
        if algo_name == "Algoritmo genético":
            self._ga_init()
        if algo_name == "Recocido simulado":
            self.sa_current = (x0, y0)
            self.sa_current_h = self.app.hmap.value_bilinear(x0, y0)
        if algo_name == "Quimiosíntesis bacteriana":
            self.bc_dir = self._rand_unit()
            self.bc_run_len = 0
            self.bc_last_h = self.app.hmap.value_bilinear(x0, y0)

    def _bounds(self):
        return self.app.get_bounds()

    def _rand_point(self):
        X_MIN, X_MAX, Y_MIN, Y_MAX = self._bounds()
        return (random.uniform(X_MIN, X_MAX), random.uniform(Y_MIN, Y_MAX))

    def _rand_unit(self):
        a = random.uniform(-math.pi, math.pi)
        return (math.cos(a), math.sin(a))

    def _clip(self, x, y):
        X_MIN, X_MAX, Y_MIN, Y_MAX = self._bounds()
        return (clamp_float(x, X_MIN, X_MAX), clamp_float(y, Y_MIN, Y_MAX))

    def _hc_next(self):
        x, y, _ = self.app.pose_runtime
        step = max(0.005, self.hc_step)

        candidates = []
        for dx in (-step, 0.0, step):
            for dy in (-step, 0.0, step):
                if dx == 0.0 and dy == 0.0:
                    continue
                xx, yy = self._clip(x + dx, y + dy)
                hv = self.app.hmap.value_bilinear(xx, yy)
                candidates.append((hv, xx, yy))

        candidates.sort(reverse=True, key=lambda t: t[0])
        best = candidates[0]
        return (best[1], best[2])

    def _ga_init(self):
        self.ga_pop = []
        for _ in range(self.ga_pop_size):
            self.ga_pop.append(self._rand_point())
        self.ga_gen = 0

    def _ga_next(self):
        scored = []
        for (x, y) in self.ga_pop:
            xx, yy = self._clip(x, y)
            hv = self.app.hmap.value_bilinear(xx, yy)
            scored.append((hv, xx, yy))
        scored.sort(reverse=True, key=lambda t: t[0])

        best = scored[0]
        best_point = (best[1], best[2])

        elites = scored[:max(1, min(self.ga_elite, len(scored)))]
        new_pop = [(e[1], e[2]) for e in elites]

        def pick_parent():
            a = random.choice(scored)
            b = random.choice(scored)
            return a if a[0] > b[0] else b

        while len(new_pop) < self.ga_pop_size:
            p1 = pick_parent()
            p2 = pick_parent()
            cx = 0.5 * (p1[1] + p2[1])
            cy = 0.5 * (p1[2] + p2[2])
            cx += random.gauss(0.0, self.ga_mut_sigma)
            cy += random.gauss(0.0, self.ga_mut_sigma)
            cx, cy = self._clip(cx, cy)
            new_pop.append((cx, cy))

        self.ga_pop = new_pop
        self.ga_gen += 1
        return best_point

    def _cg_next(self):
        x, y, _ = self.app.pose_runtime
        k = max(5, self.cg_k)
        R = max(0.02, self.cg_radius)

        samples = []
        for _ in range(k):
            ang = random.uniform(-math.pi, math.pi)
            r = R * math.sqrt(random.random())
            xx = x + r * math.cos(ang)
            yy = y + r * math.sin(ang)
            xx, yy = self._clip(xx, yy)
            hv = self.app.hmap.value_bilinear(xx, yy)
            samples.append((hv, xx, yy))

        p = 2.0
        wsum = 1e-12
        xsum = 0.0
        ysum = 0.0
        for hv, xx, yy in samples:
            w = (hv ** p)
            wsum += w
            xsum += w * xx
            ysum += w * yy

        gx = xsum / wsum
        gy = ysum / wsum
        gx, gy = self._clip(gx, gy)
        return (gx, gy)

    def _bc_next(self):
        x, y, _ = self.app.pose_runtime
        step = max(0.005, self.bc_step)

        if self.bc_dir is None:
            self.bc_dir = self._rand_unit()
            self.bc_run_len = 0
            self.bc_last_h = self.app.hmap.value_bilinear(x, y)

        dx, dy = self.bc_dir
        nx, ny = self._clip(x + step * dx, y + step * dy)
        hnew = self.app.hmap.value_bilinear(nx, ny)

        if self.bc_last_h is None:
            self.bc_last_h = self.app.hmap.value_bilinear(x, y)

        if hnew >= self.bc_last_h:
            self.bc_last_h = hnew
            self.bc_run_len += 1
            if self.bc_run_len >= self.bc_max_run:
                self.bc_dir = self._rand_unit()
                self.bc_run_len = 0
            return (nx, ny)
        else:
            self.bc_dir = self._rand_unit()
            self.bc_run_len = 0
            dx, dy = self.bc_dir
            nx, ny = self._clip(x + step * dx, y + step * dy)
            self.bc_last_h = self.app.hmap.value_bilinear(nx, ny)
            return (nx, ny)

    def _sa_next(self):
        x, y, _ = self.app.pose_runtime

        if self.sa_current is None:
            self.sa_current = (x, y)
            self.sa_current_h = self.app.hmap.value_bilinear(x, y)

        T = max(1e-6, self.sa_T)
        step = max(0.005, self.sa_step)

        ang = random.uniform(-math.pi, math.pi)
        r = step * math.sqrt(random.random())
        nx, ny = self._clip(
            self.sa_current[0] + r * math.cos(ang),
            self.sa_current[1] + r * math.sin(ang)
        )
        hnew = self.app.hmap.value_bilinear(nx, ny)

        d = hnew - self.sa_current_h
        if d >= 0:
            accept = True
        else:
            p = math.exp(d / T)
            accept = (random.random() < p)

        if accept:
            self.sa_current = (nx, ny)
            self.sa_current_h = hnew

        self.sa_T = T * max(0.50, min(self.sa_alpha, 0.999))
        return self.sa_current

    def _rw_next(self):
        x, y, _ = self.app.pose_runtime
        step = max(0.005, float(self.rw_step))
        ang = random.uniform(-math.pi, math.pi)
        r = step * math.sqrt(random.random())
        nx, ny = self._clip(x + r * math.cos(ang), y + r * math.sin(ang))
        return (nx, ny)

    def next_goal(self):
        self.iter += 1
        algo = self.algo
        if algo == "Ascenso a la colina":
            return self._hc_next()
        if algo == "Algoritmo genético":
            return self._ga_next()
        if algo == "Centro de gravedad":
            return self._cg_next()
        if algo == "Quimiosíntesis bacteriana":
            return self._bc_next()
        if algo == "Recocido simulado":
            return self._sa_next()
        if algo == "Caminata aleatoria":
            return self._rw_next()
        return self._hc_next()


# ======================================================
# Aplicación principal
# ======================================================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Búsqueda en mapa de humedad + Morfologías + Histórico (Excel)")
        self.geometry("1600x940")

        # Serial
        self.ser = None
        self.wlock = threading.Lock()

        # Thread
        self.stop_event = threading.Event()
        self.nav_thread = None

        # Plot
        self.plot_win = None

        # Humedad
        self.hmap = HumidityMap()

        # Runtime del robot
        self.ids_runtime = []
        self.morph_positions = []

        self.start_A = (0.0, 0.0)
        self.current_goal = (None, None)
        self.waypoints_runtime = []
        self.wp_index = 0

        self.pose_runtime = (0.0, 0.0, 0.0)  # (x,y,theta)
        self.traj_x = []
        self.traj_y = []
        self.t0 = time.perf_counter()

        self.best_found = (None, None, -1.0)  # (x,y,h)
        self.search_iter = 0

        self.distance_m = 0.0

        # Histórico
        self.fitness_history = []
        self.speed_acc = {}
        self.last_cmd_pos = {}
        self.speed_history = []
        self.run_config_snapshot = None

        self.humidity_series = []
        self.last_humidity = 0.0

        self.robot_series = []

        self._last_iter_dist = 0.0
        self._last_iter_time = 0.0

        self.search = SearchController(self)

        # ========== Variables UI ==========
        self.port_var = tk.StringVar(value="COM7")
        self.baud_var = tk.StringVar(value="115200")
        self.eol_var = tk.StringVar(value="LF")
        self.status_var = tk.StringVar(value="Desconectado")

        self.nmods_var = tk.IntVar(value=5)
        self.ids_var = tk.StringVar(value="3,6,4,5,1")
        self.morph_var = tk.StringVar(value="Cuadrúpedo")

        # --- NUEVO: límites configurables ---
        self.xmin_var = tk.StringVar(value="0.0")
        self.xmax_var = tk.StringVar(value="1.90")
        self.ymin_var = tk.StringVar(value="0.0")
        self.ymax_var = tk.StringVar(value="1.20")

        self.ax_var = tk.StringVar(value="0.000")
        self.ay_var = tk.StringVar(value="1.000")
        self.align_var = tk.StringVar(value="+X")

        self.step_m_var = tk.StringVar(value="0.010")
        self.x_first_var = tk.BooleanVar(value=True)
        self.tol_var = tk.StringVar(value="0.10")
        self.deadband_deg_var = tk.StringVar(value="5")
        self.tick_var = tk.StringVar(value="0.10")

        self.freq_var = tk.StringVar(value="0.40")
        self.rate_var = tk.StringVar(value="10")

        self.move_v_ticks = tk.StringVar(value="210")
        self.move_h_ticks = tk.StringVar(value="10")
        self.move_phase_v_deg = tk.StringVar(value="0")
        self.move_phase_h_deg = tk.StringVar(value=str(PHI_H_DEG_DEFAULT))

        self.turn_v_ticks = tk.StringVar(value="180")
        self.turn_h_ticks = tk.StringVar(value="210")
        self.turn_phase_v_deg = tk.StringVar(value="0")
        self.turn_phase_h_deg = tk.StringVar(value=str(PHI_H_DEG_DEFAULT))

        self.turn_phase_shift_enable = tk.BooleanVar(value=False)
        self.turn_phase_shift_dir = tk.StringVar(value="CW")
        self.turn_phase_shift_deg = tk.StringVar(value="45")

        self.kv_var = tk.StringVar(value="0.00010")
        self.kw_var = tk.StringVar(value="0.01")

        self.send_serial_var = tk.BooleanVar(value=True)

        self.show_route_var = tk.StringVar(value="Mostrar")
        self.show_goal_var = tk.StringVar(value="Mostrar")

        # H/V por ID
        self.horiz_var_by_id = {}

        # -------- Mapa humedad UI --------
        self.hm_nx_var = tk.StringVar(value="60")
        self.hm_ny_var = tk.StringVar(value="40")
        self.hm_seed_var = tk.StringVar(value="1")
        self.hm_noise_var = tk.StringVar(value="0.05")
        self.hm_peakx_var = tk.StringVar(value="1.20")
        self.hm_peaky_var = tk.StringVar(value="0.70")
        self.hm_sigma_var = tk.StringVar(value="0.18")

        # -------- Algoritmos UI --------
        self.algo_var = tk.StringVar(value="Ascenso a la colina")
        self.max_iters_var = tk.StringVar(value="120")
        self.stop_h_thresh_var = tk.StringVar(value="0.995")

        self.hc_step_var = tk.StringVar(value="0.10")
        self.ga_pop_var = tk.StringVar(value="18")
        self.ga_elite_var = tk.StringVar(value="4")
        self.ga_mut_var = tk.StringVar(value="0.12")
        self.cg_k_var = tk.StringVar(value="20")
        self.cg_radius_var = tk.StringVar(value="0.35")
        self.bc_step_var = tk.StringVar(value="0.08")
        self.bc_maxrun_var = tk.StringVar(value="10")
        self.sa_T0_var = tk.StringVar(value="1.0")
        self.sa_alpha_var = tk.StringVar(value="0.92")
        self.sa_step_var = tk.StringVar(value="0.18")
        self.rw_step_var = tk.StringVar(value="0.12")

        # construir UI
        self._build_ui()
        self.protocol("WM_DELETE_WINDOW", self.on_close)

        self._rebuild_hv_checks()
        self.update_morphology_runtime()

        # Inicializar con límites válidos + mapa
        self.apply_bounds_and_reinit()
        self.generate_humidity_map()

    # ================== Límites (NUEVO) ==================
    def get_bounds(self):
        """
        Retorna bounds = (X_MIN, X_MAX, Y_MIN, Y_MAX) garantizando coherencia:
        - Si el usuario pone min>=max, se corrige.
        - Si hay errores de parseo, se usan valores por defecto.
        """
        try:
            xmin = float(self.xmin_var.get())
            xmax = float(self.xmax_var.get())
            ymin = float(self.ymin_var.get())
            ymax = float(self.ymax_var.get())
        except Exception:
            xmin, xmax, ymin, ymax = 0.0, 1.90, 0.0, 1.20

        # Correcciones mínimas para evitar rangos degenerados
        if xmax <= xmin:
            xmax = xmin + 1e-6
        if ymax <= ymin:
            ymax = ymin + 1e-6

        return (xmin, xmax, ymin, ymax)

    def clip_xy(self, x, y):
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.get_bounds()
        return (clamp_float(x, X_MIN, X_MAX), clamp_float(y, Y_MIN, Y_MAX))

    def apply_bounds_and_reinit(self):
        """
        Aplica límites:
        - Clampea A.x/A.y dentro del rango.
        - Ajusta pico del mapa dentro del rango.
        - Re-centra pose/trayectoria si estaban fuera.
        """
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.get_bounds()

        # Clamp de A
        try:
            ax = float(self.ax_var.get())
            ay = float(self.ay_var.get())
        except Exception:
            ax, ay = X_MIN, Y_MIN
        ax, ay = self.clip_xy(ax, ay)
        self.ax_var.set(f"{ax:.3f}")
        self.ay_var.set(f"{ay:.3f}")

        # Clamp de pico (si estaba fuera)
        try:
            px = float(self.hm_peakx_var.get())
            py = float(self.hm_peaky_var.get())
        except Exception:
            px, py = (X_MIN + X_MAX) / 2.0, (Y_MIN + Y_MAX) / 2.0
        px, py = self.clip_xy(px, py)
        self.hm_peakx_var.set(f"{px:.3f}")
        self.hm_peaky_var.set(f"{py:.3f}")

        # Clamp pose actual
        x, y, th = self.pose_runtime
        x, y = self.clip_xy(x, y)
        self.pose_runtime = (x, y, th)
        if not self.traj_x:
            self.traj_x = [x]
            self.traj_y = [y]
        else:
            self.traj_x[-1] = x
            self.traj_y[-1] = y

        # Clamp objetivo actual (si existe)
        gx, gy = self.current_goal
        if gx is not None:
            gx, gy = self.clip_xy(gx, gy)
            self.current_goal = (gx, gy)

        # Clamp mejor encontrado
        bx, by, bv = self.best_found
        if bx is not None:
            bx, by = self.clip_xy(bx, by)
            self.best_found = (bx, by, bv)

    # ================= UI =================
    def _build_ui(self):
        root = ttk.Frame(self, padding=10)
        root.pack(fill="both", expand=True)
        for i in range(3):
            root.columnconfigure(i, weight=1)

        col1 = ttk.Frame(root)
        col2 = ttk.Frame(root)
        col3 = ttk.Frame(root)
        col1.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        col2.grid(row=0, column=1, sticky="nsew", padx=(8, 8))
        col3.grid(row=0, column=2, sticky="nsew", padx=(8, 0))

        self._ui_serial(col1)
        self._ui_modules(col1)
        self._ui_morphology(col1)
        self._ui_humidity(col1)

        self._ui_navigation(col2)
        self._ui_motion(col2)
        self._ui_algo(col2)

        self._ui_exec(col3)
        self._ui_status(col3)
        self._ui_history(col3)

    def _ui_serial(self, parent):
        f = ttk.LabelFrame(parent, text="Puerto Serial", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="COM").grid(row=0, column=0, sticky="w")
        self.port_cb = ttk.Combobox(f, textvariable=self.port_var, values=list_serial_ports(), width=16)
        self.port_cb.grid(row=0, column=1, padx=6, sticky="w")
        ttk.Button(f, text="Refrescar", command=self._refresh_ports).grid(row=0, column=2, padx=6)

        ttk.Label(f, text="Baudios").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.baud_var, width=10).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))

        ttk.Label(f, text="EOL").grid(row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Combobox(f, textvariable=self.eol_var, values=["LF", "CRLF"], width=6, state="readonly")\
            .grid(row=1, column=3, padx=6, sticky="w", pady=(6, 0))

        ttk.Button(f, text="Conectar", command=self.connect).grid(row=2, column=1, padx=6, pady=(8, 0), sticky="w")
        ttk.Button(f, text="Desconectar", command=self.disconnect).grid(row=2, column=2, padx=6, pady=(8, 0), sticky="w")

        ttk.Checkbutton(f, text="Enviar comandos al robot (Serial)", variable=self.send_serial_var)\
            .grid(row=3, column=0, columnspan=4, sticky="w", pady=(10, 0))

    def _ui_modules(self, parent):
        f = ttk.LabelFrame(parent, text="Módulos", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="Cantidad").grid(row=0, column=0, sticky="w")
        ttk.Combobox(f, textvariable=self.nmods_var, values=list(range(1, 25)), width=6, state="readonly")\
            .grid(row=0, column=1, padx=6, sticky="w")
        ttk.Button(f, text="Aplicar", command=self._apply_nmods).grid(row=0, column=2, padx=6, sticky="w")

        ttk.Label(f, text="IDs (coma)").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.ids_var, width=26).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0), columnspan=2)

        self.hv_frame = ttk.Frame(f)
        self.hv_frame.grid(row=2, column=0, columnspan=3, sticky="ew", pady=(10, 0))

        ttk.Label(f, text="Nota: H/V define qué onda aplica a cada ID.").grid(
            row=3, column=0, columnspan=3, sticky="w", pady=(8, 0)
        )

    def _ui_morphology(self, parent):
        f = ttk.LabelFrame(parent, text="Morfología", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="Tipo").grid(row=0, column=0, sticky="w")
        ttk.Combobox(f, textvariable=self.morph_var,
                     values=["Serpiente", "L", "T", "Cuadrúpedo"], width=12, state="readonly")\
            .grid(row=0, column=1, padx=6, sticky="w")
        ttk.Button(f, text="Actualizar dibujo", command=self.update_morphology_runtime)\
            .grid(row=0, column=2, padx=6, sticky="w")

    def _ui_humidity(self, parent):
        f = ttk.LabelFrame(parent, text="Mapa de humedad", padding=10)
        f.pack(fill="x", pady=(0, 10))

        # Límites (NUEVO)
        lim = ttk.LabelFrame(f, text="Límites del espacio (NUEVO)", padding=8)
        lim.grid(row=0, column=0, columnspan=4, sticky="ew")

        ttk.Label(lim, text="Xmin").grid(row=0, column=0, sticky="w")
        ttk.Entry(lim, textvariable=self.xmin_var, width=8).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(lim, text="Xmax").grid(row=0, column=2, sticky="w")
        ttk.Entry(lim, textvariable=self.xmax_var, width=8).grid(row=0, column=3, padx=6, sticky="w")

        ttk.Label(lim, text="Ymin").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(lim, textvariable=self.ymin_var, width=8).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Label(lim, text="Ymax").grid(row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(lim, textvariable=self.ymax_var, width=8).grid(row=1, column=3, padx=6, sticky="w", pady=(6, 0))

        ttk.Button(lim, text="Aplicar límites", command=self.on_apply_limits)\
            .grid(row=2, column=0, columnspan=4, sticky="w", pady=(8, 0))

        # Parámetros del mapa
        ttk.Label(f, text="Nx").grid(row=1, column=0, sticky="w", pady=(10, 0))
        ttk.Entry(f, textvariable=self.hm_nx_var, width=6).grid(row=1, column=1, padx=6, sticky="w", pady=(10, 0))
        ttk.Label(f, text="Ny").grid(row=1, column=2, sticky="w", pady=(10, 0))
        ttk.Entry(f, textvariable=self.hm_ny_var, width=6).grid(row=1, column=3, padx=6, sticky="w", pady=(10, 0))

        ttk.Label(f, text="Seed").grid(row=2, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.hm_seed_var, width=6).grid(row=2, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Label(f, text="Ruido (0..1)").grid(row=2, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.hm_noise_var, width=6).grid(row=2, column=3, padx=6, sticky="w", pady=(6, 0))

        ttk.Label(f, text="Pico x").grid(row=3, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.hm_peakx_var, width=8).grid(row=3, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Label(f, text="Pico y").grid(row=3, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.hm_peaky_var, width=8).grid(row=3, column=3, padx=6, sticky="w", pady=(6, 0))

        ttk.Label(f, text="Sigma").grid(row=4, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.hm_sigma_var, width=8).grid(row=4, column=1, padx=6, sticky="w", pady=(6, 0))

        ttk.Button(f, text="Generar mapa", command=self.generate_humidity_map)\
            .grid(row=5, column=0, columnspan=4, sticky="w", pady=(10, 0))

    def _ui_navigation(self, parent):
        f = ttk.LabelFrame(parent, text="Navegación (pose inicial)", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="A.x").grid(row=0, column=0, sticky="w")
        ttk.Entry(f, textvariable=self.ax_var, width=10).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(f, text="A.y").grid(row=0, column=2, sticky="w")
        ttk.Entry(f, textvariable=self.ay_var, width=10).grid(row=0, column=3, padx=6, sticky="w")

        ttk.Label(f, text="Alineación").grid(row=0, column=4, sticky="w")
        ttk.Combobox(f, textvariable=self.align_var, values=["+X", "+Y", "-X", "-Y"], width=6, state="readonly")\
            .grid(row=0, column=5, padx=6, sticky="w")

        ttk.Label(f, text="Paso (m)").grid(row=1, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(f, textvariable=self.step_m_var, width=10).grid(row=1, column=1, padx=6, sticky="w", pady=(8, 0))
        ttk.Checkbutton(f, text="Primero X luego Y", variable=self.x_first_var)\
            .grid(row=1, column=2, columnspan=2, sticky="w", pady=(8, 0))

        ttk.Label(f, text="Tol (m)").grid(row=2, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(f, textvariable=self.tol_var, width=10).grid(row=2, column=1, padx=6, sticky="w", pady=(8, 0))
        ttk.Label(f, text="Deadband (°)").grid(row=2, column=2, sticky="w", pady=(8, 0))
        ttk.Entry(f, textvariable=self.deadband_deg_var, width=10).grid(row=2, column=3, padx=6, sticky="w", pady=(8, 0))

        ttk.Label(f, text="Tick ejes (m)").grid(row=3, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(f, textvariable=self.tick_var, width=10).grid(row=3, column=1, padx=6, sticky="w", pady=(8, 0))

        ttk.Label(f, text="Ejes se ajustan a límites configurados.").grid(
            row=3, column=2, columnspan=4, sticky="w", pady=(8, 0)
        )

    def _ui_motion(self, parent):
        f = ttk.LabelFrame(parent, text="Movimiento", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="Frecuencia (Hz)").grid(row=0, column=0, sticky="w")
        ttk.Entry(f, textvariable=self.freq_var, width=10).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(f, text="Tasa envío (Hz)").grid(row=0, column=2, sticky="w")
        ttk.Entry(f, textvariable=self.rate_var, width=10).grid(row=0, column=3, padx=6, sticky="w")

        f1 = ttk.LabelFrame(f, text="Avance (recto)", padding=8)
        f1.grid(row=1, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Label(f1, text="V(A ticks)").grid(row=0, column=0, sticky="w")
        ttk.Entry(f1, textvariable=self.move_v_ticks, width=10).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(f1, text="H(A ticks)").grid(row=0, column=2, sticky="w")
        ttk.Entry(f1, textvariable=self.move_h_ticks, width=10).grid(row=0, column=3, padx=6, sticky="w")

        ttk.Label(f1, text="V(φ°)").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f1, textvariable=self.move_phase_v_deg, width=10).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Label(f1, text="H(φ°)").grid(row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(f1, textvariable=self.move_phase_h_deg, width=10).grid(row=1, column=3, padx=6, sticky="w", pady=(6, 0))

        f2 = ttk.LabelFrame(f, text="Giro", padding=8)
        f2.grid(row=2, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Label(f2, text="V(A ticks)").grid(row=0, column=0, sticky="w")
        ttk.Entry(f2, textvariable=self.turn_v_ticks, width=10).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(f2, text="H(A ticks)").grid(row=0, column=2, sticky="w")
        ttk.Entry(f2, textvariable=self.turn_h_ticks, width=10).grid(row=0, column=3, padx=6, sticky="w")

        ttk.Label(f2, text="V(φ°)").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f2, textvariable=self.turn_phase_v_deg, width=10).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Label(f2, text="H(φ°)").grid(row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(f2, textvariable=self.turn_phase_h_deg, width=10).grid(row=1, column=3, padx=6, sticky="w", pady=(6, 0))

        f3 = ttk.LabelFrame(f, text="Giro: cambio de fase SOLO en un sentido", padding=8)
        f3.grid(row=3, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Checkbutton(f3, text="Habilitar", variable=self.turn_phase_shift_enable).grid(row=0, column=0, sticky="w")
        ttk.Label(f3, text="Sentido").grid(row=0, column=1, sticky="w")
        ttk.Combobox(f3, textvariable=self.turn_phase_shift_dir, values=["CW", "CCW"], width=6, state="readonly")\
            .grid(row=0, column=2, padx=6, sticky="w")
        ttk.Label(f3, text="Δφ (°) aplicado a H").grid(row=0, column=3, sticky="w")
        ttk.Entry(f3, textvariable=self.turn_phase_shift_deg, width=10).grid(row=0, column=4, padx=6, sticky="w")

        f4 = ttk.LabelFrame(f, text="Modelo (pose) para la gráfica", padding=8)
        f4.grid(row=4, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Label(f4, text="k_v").grid(row=0, column=0, sticky="w")
        ttk.Entry(f4, textvariable=self.kv_var, width=10).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(f4, text="k_w").grid(row=0, column=2, sticky="w")
        ttk.Entry(f4, textvariable=self.kw_var, width=10).grid(row=0, column=3, padx=6, sticky="w")

    def _ui_algo(self, parent):
        f = ttk.LabelFrame(parent, text="Búsqueda (selección de algoritmo)", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="Algoritmo").grid(row=0, column=0, sticky="w")
        ttk.Combobox(
            f,
            textvariable=self.algo_var,
            values=[
                "Ascenso a la colina",
                "Algoritmo genético",
                "Centro de gravedad",
                "Quimiosíntesis bacteriana",
                "Recocido simulado",
                "Caminata aleatoria"
            ],
            width=24,
            state="readonly"
        ).grid(row=0, column=1, padx=6, sticky="w")

        ttk.Label(f, text="Máx iteraciones").grid(row=0, column=2, sticky="w")
        ttk.Entry(f, textvariable=self.max_iters_var, width=8).grid(row=0, column=3, padx=6, sticky="w")

        ttk.Label(f, text="Parar si h ≥").grid(row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.stop_h_thresh_var, width=8).grid(row=1, column=3, padx=6, sticky="w", pady=(6, 0))

        box = ttk.Frame(f)
        box.grid(row=2, column=0, columnspan=4, sticky="ew", pady=(10, 0))

        hf = ttk.LabelFrame(box, text="Hill", padding=8)
        hf.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        ttk.Label(hf, text="step").grid(row=0, column=0, sticky="w")
        ttk.Entry(hf, textvariable=self.hc_step_var, width=8).grid(row=0, column=1, padx=6, sticky="w")

        gf = ttk.LabelFrame(box, text="GA", padding=8)
        gf.grid(row=0, column=1, sticky="ew", padx=(0, 8))
        ttk.Label(gf, text="pop").grid(row=0, column=0, sticky="w")
        ttk.Entry(gf, textvariable=self.ga_pop_var, width=6).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(gf, text="elite").grid(row=0, column=2, sticky="w")
        ttk.Entry(gf, textvariable=self.ga_elite_var, width=6).grid(row=0, column=3, padx=6, sticky="w")
        ttk.Label(gf, text="mutσ").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(gf, textvariable=self.ga_mut_var, width=6).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))

        cf = ttk.LabelFrame(box, text="Centroide", padding=8)
        cf.grid(row=0, column=2, sticky="ew", padx=(0, 8))
        ttk.Label(cf, text="k").grid(row=0, column=0, sticky="w")
        ttk.Entry(cf, textvariable=self.cg_k_var, width=6).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(cf, text="R").grid(row=0, column=2, sticky="w")
        ttk.Entry(cf, textvariable=self.cg_radius_var, width=6).grid(row=0, column=3, padx=6, sticky="w")

        bf = ttk.LabelFrame(box, text="Bacteria", padding=8)
        bf.grid(row=1, column=0, sticky="ew", padx=(0, 8), pady=(8, 0))
        ttk.Label(bf, text="step").grid(row=0, column=0, sticky="w")
        ttk.Entry(bf, textvariable=self.bc_step_var, width=8).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(bf, text="maxrun").grid(row=0, column=2, sticky="w")
        ttk.Entry(bf, textvariable=self.bc_maxrun_var, width=6).grid(row=0, column=3, padx=6, sticky="w")

        sf = ttk.LabelFrame(box, text="SA", padding=8)
        sf.grid(row=1, column=1, sticky="ew", padx=(0, 8), pady=(8, 0))
        ttk.Label(sf, text="T0").grid(row=0, column=0, sticky="w")
        ttk.Entry(sf, textvariable=self.sa_T0_var, width=6).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(sf, text="α").grid(row=0, column=2, sticky="w")
        ttk.Entry(sf, textvariable=self.sa_alpha_var, width=6).grid(row=0, column=3, padx=6, sticky="w")
        ttk.Label(sf, text="step").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(sf, textvariable=self.sa_step_var, width=6).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))

        rf = ttk.LabelFrame(box, text="Random walk", padding=8)
        rf.grid(row=1, column=2, sticky="ew", padx=(0, 8), pady=(8, 0))
        ttk.Label(rf, text="step").grid(row=0, column=0, sticky="w")
        ttk.Entry(rf, textvariable=self.rw_step_var, width=8).grid(row=0, column=1, padx=6, sticky="w")

    def _ui_exec(self, parent):
        f = ttk.LabelFrame(parent, text="Ejecución", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Button(f, text="Abrir gráficas", command=self.open_plot_window).pack(side="left", padx=6)
        ttk.Button(f, text="Actualizar Morfología", command=self.update_morphology_runtime).pack(side="left", padx=6)
        ttk.Button(f, text="HOME", command=self.home_all).pack(side="left", padx=6)

        ttk.Button(f, text="Iniciar búsqueda", command=self.start_search).pack(side="left", padx=6)
        ttk.Button(f, text="Detener", command=self.stop).pack(side="left", padx=6)

        f2 = ttk.Frame(parent)
        f2.pack(fill="x", pady=(0, 10))
        box = ttk.LabelFrame(f2, text="Vista (gráfica)", padding=10)
        box.pack(fill="x")

        ttk.Label(box, text="Ruta").grid(row=0, column=0, sticky="w")
        ttk.Combobox(box, textvariable=self.show_route_var, values=["Mostrar", "Ocultar"], width=10, state="readonly")\
            .grid(row=0, column=1, padx=6, sticky="w")

        ttk.Label(box, text="Punto final").grid(row=0, column=2, sticky="w")
        ttk.Combobox(box, textvariable=self.show_goal_var, values=["Mostrar", "Ocultar"], width=10, state="readonly")\
            .grid(row=0, column=3, padx=6, sticky="w")

    def _ui_status(self, parent):
        f = ttk.LabelFrame(parent, text="Estado", padding=10)
        f.pack(fill="x")
        ttk.Label(f, textvariable=self.status_var, wraplength=520, justify="left").pack(anchor="w")

    def _ui_history(self, parent):
        f = ttk.LabelFrame(parent, text="Exportación", padding=10)
        f.pack(fill="x", pady=(10, 0))

        ttk.Button(f, text="Exportar a Excel (XLSX + JSON)", command=self.export_history).pack(side="left", padx=6)
        ttk.Button(f, text="Limpiar histórico", command=self.clear_history).pack(side="left", padx=6)

    # ================== Eventos límites (NUEVO) ==================
    def on_apply_limits(self):
        # Validar / corregir, re-clampear y regenerar mapa para consistencia
        self.apply_bounds_and_reinit()
        self.generate_humidity_map()
        self.set_status("Límites aplicados. Mapa regenerado y coordenadas validadas.")

    # ================== H/V dinámico ==================
    def _apply_nmods(self):
        self._rebuild_hv_checks()
        self.update_morphology_runtime()

    def _rebuild_hv_checks(self):
        for w in self.hv_frame.winfo_children():
            w.destroy()

        n = int(self.nmods_var.get())
        try:
            ids = parse_id_list(self.ids_var.get())
        except Exception:
            ids = []

        if len(ids) != n:
            if len(ids) > n:
                ids = ids[:n]
            else:
                used = set(ids)
                nxt = 1
                while len(ids) < n:
                    if nxt not in used:
                        ids.append(nxt)
                        used.add(nxt)
                    nxt += 1
            self.ids_var.set(",".join(str(i) for i in ids))

        self.horiz_var_by_id = {}
        for rid in ids:
            self.horiz_var_by_id[rid] = tk.BooleanVar(value=False)

        ttk.Label(self.hv_frame, text="Marcar IDs horizontales (H):").grid(row=0, column=0, columnspan=6, sticky="w")
        for i, rid in enumerate(ids):
            r = 1 + (i // 6)
            c = i % 6
            ttk.Checkbutton(self.hv_frame, text=f"ID {rid}", variable=self.horiz_var_by_id[rid])\
                .grid(row=r, column=c, padx=(0, 10), pady=(6, 0), sticky="w")

    def horizontal_ids_set(self):
        return {rid for rid, var in self.horiz_var_by_id.items() if var.get()}

    # ================== Morfología runtime ==================
    def update_morphology_runtime(self):
        try:
            ids = parse_id_list(self.ids_var.get())
        except Exception:
            ids = []
        if not ids:
            self.set_status("IDs vacíos.")
            return
        self.ids_runtime = ids[:]
        self.morph_positions = morphology_positions(self.morph_var.get().strip(), len(ids))
        if self.plot_win and self.plot_win.winfo_exists():
            self.plot_win.lift()

    # ================= Humedad =================
    def generate_humidity_map(self):
        try:
            nx = int(self.hm_nx_var.get())
            ny = int(self.hm_ny_var.get())
            seed = int(self.hm_seed_var.get())
            noise = float(self.hm_noise_var.get())
            px = float(self.hm_peakx_var.get())
            py = float(self.hm_peaky_var.get())
            sg = float(self.hm_sigma_var.get())
        except Exception as e:
            messagebox.showerror("Humedad", f"Parámetros inválidos: {e}")
            return

        bounds = self.get_bounds()
        # Asegurar pico dentro de límites (doble verificación)
        px, py = self.clip_xy(px, py)
        self.hm_peakx_var.set(f"{px:.3f}")
        self.hm_peaky_var.set(f"{py:.3f}")

        self.hmap.generate(nx=nx, ny=ny, seed=seed, noise=noise, peak_x=px, peak_y=py, peak_sigma=sg, bounds=bounds)
        mx = self.hmap.argmax_grid()
        if mx:
            xg, yg, vg = mx
            self.set_status(f"Mapa generado. Máximo grid h={vg:.3f} en ({xg:.2f},{yg:.2f}).")
        else:
            self.set_status("Mapa generado.")
        if self.plot_win and self.plot_win.winfo_exists():
            self.plot_win.lift()

    # ================= Plot =================
    def open_plot_window(self):
        if self.plot_win and self.plot_win.winfo_exists():
            self.plot_win.lift()
            return
        self.plot_win = PlotWindow(self)

    # ================= Serial =================
    def _refresh_ports(self):
        self.port_cb["values"] = list_serial_ports()

    def _eol_bytes(self) -> bytes:
        return b"\r\n" if self.eol_var.get() == "CRLF" else b"\n"

    def send_lines_batch(self, lines):
        if not self.ser or not self.ser.is_open:
            raise RuntimeError("Maestro no conectado")
        eol = self._eol_bytes()
        payload = b"".join([ln.strip().encode("utf-8") + eol for ln in lines])
        with self.wlock:
            self.ser.write(payload)

    def connect(self):
        port = self.port_var.get().strip()
        if not port:
            messagebox.showwarning("Serial", "Seleccione el COM.")
            return
        try:
            baud = int(self.baud_var.get().strip())
        except ValueError:
            messagebox.showwarning("Serial", "Baudios inválidos.")
            return

        try:
            self.ser = serial.Serial(port, baud, timeout=0.2)
            time.sleep(1.2)
        except Exception as e:
            messagebox.showerror("Serial", f"No fue posible abrir el puerto: {e}")
            self.ser = None
            return

        self.status_var.set(f"Conectado a {port} @ {baud} | EOL={self.eol_var.get()}")

    def disconnect(self):
        self.stop()
        try:
            if self.ser and self.ser.is_open:
                self.ser.close()
        except Exception:
            pass
        self.ser = None
        self.status_var.set("Desconectado")

    def set_status(self, s: str):
        self.status_var.set(s)

    def require_serial(self):
        return bool(self.send_serial_var.get())

    # ================= HOME =================
    def home_all(self):
        if self.require_serial() and (not self.ser or not self.ser.is_open):
            messagebox.showwarning("HOME", "Conecte el Maestro o desactive 'Enviar comandos al robot'.")
            return

        self.stop_event.set()

        ids = parse_id_list(self.ids_var.get())
        if not ids:
            messagebox.showwarning("HOME", "IDs vacíos.")
            return

        if self.require_serial():
            lines = []
            for rid in ids:
                if rid <= 3:
                    lines.append(f"N {rid} POS {HOME_DX}")
                else:
                    lines.append(f"N {rid} LX {HOME_LX}")
            try:
                self.send_lines_batch(lines)
            except Exception as e:
                self.set_status(f"Error HOME: {e}")
                return

        # reset pose y búsqueda (respetando límites)
        self.apply_bounds_and_reinit()
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.get_bounds()

        ax = float(self.ax_var.get())
        ay = float(self.ay_var.get())
        th0 = theta_from_alignment(self.align_var.get().strip())
        self.pose_runtime = (ax, ay, th0)

        self.traj_x = [ax]
        self.traj_y = [ay]
        self.t0 = time.perf_counter()

        self.distance_m = 0.0
        self._last_iter_dist = 0.0
        self._last_iter_time = 0.0

        hv0 = self.hmap.value_bilinear(ax, ay) if self.hmap.is_ready() else 0.0
        self.best_found = (ax, ay, hv0)
        self.search_iter = 0

        self.current_goal = (None, None)
        self.waypoints_runtime = []
        self.wp_index = 0

        self.last_cmd_pos = {}
        self.speed_acc = {}
        self.speed_history = []
        self.humidity_series = []
        self.robot_series = []
        self.last_humidity = hv0

        self.set_status("HOME enviado (si aplica) y pose/búsqueda reiniciadas (con límites).")

    # ================= Histórico: configuración =================
    def snapshot_config(self):
        ids = parse_id_list(self.ids_var.get())
        horiz = sorted(list(self.horizontal_ids_set()))
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.get_bounds()
        cfg = {
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "bounds": {"xmin": X_MIN, "xmax": X_MAX, "ymin": Y_MIN, "ymax": Y_MAX},
            "serial": {
                "port": self.port_var.get(),
                "baud": self.baud_var.get(),
                "eol": self.eol_var.get(),
                "send_enabled": bool(self.send_serial_var.get())
            },
            "modules": {
                "n": int(self.nmods_var.get()),
                "ids": ids,
                "horizontal_ids": horiz,
                "morphology": self.morph_var.get()
            },
            "navigation": {
                "A": {"x": float(self.ax_var.get()), "y": float(self.ay_var.get())},
                "alignment": self.align_var.get(),
                "step_m": float(self.step_m_var.get()),
                "x_first": bool(self.x_first_var.get()),
                "tol_m": float(self.tol_var.get()),
                "deadband_deg": float(self.deadband_deg_var.get()),
                "tick_axes_m": float(self.tick_var.get())
            },
            "motion": {
                "freq_hz": float(self.freq_var.get()),
                "rate_hz": float(self.rate_var.get()),
                "move": {"Av": float(self.move_v_ticks.get()), "Ah": float(self.move_h_ticks.get()),
                         "phiV_deg": float(self.move_phase_v_deg.get()), "phiH_deg": float(self.move_phase_h_deg.get())},
                "turn": {"Av": float(self.turn_v_ticks.get()), "Ah": float(self.turn_h_ticks.get()),
                         "phiV_deg": float(self.turn_phase_v_deg.get()), "phiH_deg": float(self.turn_phase_h_deg.get())},
                "phase_shift": {"enable": bool(self.turn_phase_shift_enable.get()),
                                "dir": self.turn_phase_shift_dir.get(),
                                "deg": float(self.turn_phase_shift_deg.get())},
                "pose_model": {"kv": float(self.kv_var.get()), "kw": float(self.kw_var.get())}
            },
            "humidity_map": {
                "nx": int(self.hm_nx_var.get()), "ny": int(self.hm_ny_var.get()),
                "seed": int(self.hm_seed_var.get()), "noise": float(self.hm_noise_var.get()),
                "peak_x": float(self.hm_peakx_var.get()), "peak_y": float(self.hm_peaky_var.get()),
                "sigma": float(self.hm_sigma_var.get())
            },
            "search": {
                "algorithm": self.algo_var.get(),
                "max_iters": int(self.max_iters_var.get()),
                "stop_threshold": float(self.stop_h_thresh_var.get()),
                "params": {
                    "hc_step": float(self.hc_step_var.get()),
                    "ga_pop": int(self.ga_pop_var.get()),
                    "ga_elite": int(self.ga_elite_var.get()),
                    "ga_mut_sigma": float(self.ga_mut_var.get()),
                    "cg_k": int(self.cg_k_var.get()),
                    "cg_radius": float(self.cg_radius_var.get()),
                    "bc_step": float(self.bc_step_var.get()),
                    "bc_maxrun": int(self.bc_maxrun_var.get()),
                    "sa_T0": float(self.sa_T0_var.get()),
                    "sa_alpha": float(self.sa_alpha_var.get()),
                    "sa_step": float(self.sa_step_var.get()),
                    "rw_step": float(self.rw_step_var.get())
                }
            }
        }
        return cfg

    # ================= Movimiento: líneas + posiciones =================
    def build_servo_lines_with_positions(self, t, freq, Av, Ah, phi_v_deg, phi_h_deg, turning, cw):
        sign_v = -1.0
        if turning:
            sign_h = sign_v if cw else -sign_v
        else:
            sign_h = sign_v

        Av_lx = Av * LX_PER_DX
        Ah_lx = Ah * LX_PER_DX

        ids = self.ids_runtime
        n = len(ids)
        if n <= 0:
            return [], {}

        horizontals = self.horizontal_ids_set()
        phase_origin_deg = 0.0

        out = []
        posmap = {}

        for i, rid in enumerate(ids):
            phi_idx_deg = phase_origin_deg + (360.0 * i / n)

            if rid in horizontals:
                phase = 2 * math.pi * freq * t + deg2rad(phi_h_deg + sign_h * phi_idx_deg)
                ss = math.sin(phase)
                if rid <= 3:
                    pos = clamp_int(HOME_DX + Ah * ss, DX_MIN, DX_MAX)
                    out.append(f"N {rid} POS {pos}")
                    posmap[rid] = float(pos)
                else:
                    ang = clamp_int(HOME_LX + Ah_lx * ss, LX_MIN, LX_MAX)
                    out.append(f"N {rid} LX {ang}")
                    posmap[rid] = float(ang)
            else:
                phase = 2 * math.pi * freq * t + deg2rad(phi_v_deg + sign_v * phi_idx_deg)
                ss = math.sin(phase)
                if rid <= 3:
                    pos = clamp_int(HOME_DX + Av * ss, DX_MIN, DX_MAX)
                    out.append(f"N {rid} POS {pos}")
                    posmap[rid] = float(pos)
                else:
                    ang = clamp_int(HOME_LX + Av_lx * ss, LX_MIN, LX_MAX)
                    out.append(f"N {rid} LX {ang}")
                    posmap[rid] = float(ang)

        return out, posmap

    # ================= Histórico: velocidades por módulo =================
    def init_speed_acc_if_needed(self):
        for rid in self.ids_runtime:
            if rid not in self.speed_acc:
                self.speed_acc[rid] = {"sum_abs": 0.0, "sum_dt": 0.0, "avg": 0.0}

    def update_speed_history(self, cmd_positions: dict, dt: float):
        if dt <= 0:
            return
        self.init_speed_acc_if_needed()

        for rid, pos in cmd_positions.items():
            if rid in self.last_cmd_pos:
                d = abs(pos - self.last_cmd_pos[rid])
                acc = self.speed_acc.setdefault(rid, {"sum_abs": 0.0, "sum_dt": 0.0, "avg": 0.0})
                acc["sum_abs"] += d
                acc["sum_dt"] += dt
                if acc["sum_dt"] > 1e-12:
                    acc["avg"] = acc["sum_abs"] / acc["sum_dt"]
            self.last_cmd_pos[rid] = pos

        now = time.perf_counter()
        if not hasattr(self, "_last_speed_snap_t"):
            self._last_speed_snap_t = now
        if (now - self._last_speed_snap_t) >= 0.5:
            self._last_speed_snap_t = now
            avg_by_id = {rid: round(self.speed_acc[rid]["avg"], 6) for rid in sorted(self.speed_acc.keys())}
            self.speed_history.append({
                "t": round(now - self.t0, 6),
                "avg_by_id": avg_by_id,
                "distance_m": round(self.distance_m, 6)
            })

    # ================= Pose + distancia + serie robot =================
    def update_pose_model(self, Av, Ah, turning, cw, kv, kw, dt):
        v = kv * max(0.0, Av - Ah)
        w = 0.0
        if turning:
            w = kw * max(0.0, Ah - Av) * (-1.0 if cw else +1.0)

        x, y, th = self.pose_runtime
        x_prev, y_prev = x, y

        th = th + w * dt
        x = x + v * math.cos(th) * dt
        y = y + v * math.sin(th) * dt

        # clamp dentro de límites configurados
        x, y = self.clip_xy(x, y)

        dd = math.hypot(x - x_prev, y - y_prev)
        self.distance_m += dd

        self.pose_runtime = (x, y, th)
        self.traj_x.append(x)
        self.traj_y.append(y)

        now = time.perf_counter()
        if not hasattr(self, "_last_robot_snap_t"):
            self._last_robot_snap_t = now
        if (now - self._last_robot_snap_t) >= 0.5:
            self._last_robot_snap_t = now
            t = now - self.t0
            v_inst = dd / max(1e-9, dt)
            self.robot_series.append({"t": round(t, 6), "distance_m": round(self.distance_m, 6), "v_m_s": round(v_inst, 6)})

    def sample_humidity_at_pose(self):
        if not self.hmap.is_ready():
            return
        x, y, _ = self.pose_runtime
        hv = self.hmap.value_bilinear(x, y)
        self.last_humidity = hv

        bx, by, bv = self.best_found
        if hv > bv:
            self.best_found = (x, y, hv)

        now = time.perf_counter()
        if not hasattr(self, "_last_hum_snap_t"):
            self._last_hum_snap_t = now
        if (now - self._last_hum_snap_t) >= 0.5:
            self._last_hum_snap_t = now
            self.humidity_series.append({"t": round(now - self.t0, 6), "x": round(x, 6), "y": round(y, 6),
                                         "h": round(hv, 6), "distance_m": round(self.distance_m, 6)})

    # ================= Ruta hacia objetivo =================
    def set_new_goal(self, gx, gy):
        gx, gy = self.clip_xy(gx, gy)
        self.current_goal = (gx, gy)

        x, y, _ = self.pose_runtime
        step_m = float(self.step_m_var.get())
        x_first = bool(self.x_first_var.get())
        self.start_A = (x, y)

        bounds = self.get_bounds()
        self.waypoints_runtime = build_manhattan_waypoints(x, y, gx, gy, step_m, x_first=x_first, bounds=bounds)
        self.wp_index = 0

    def current_wp(self):
        if 0 <= self.wp_index < len(self.waypoints_runtime):
            return self.waypoints_runtime[self.wp_index]
        return None

    def advance_wp(self):
        if self.wp_index + 1 < len(self.waypoints_runtime):
            self.wp_index += 1
            return True
        self.wp_index = len(self.waypoints_runtime)
        return False

    # ================= Fitness por iteración (histórico) =================
    def record_fitness_iteration(self):
        bx, by, bv = self.best_found
        if bx is None:
            return

        elapsed_s = float(time.perf_counter() - self.t0)
        dist_total = float(self.distance_m)

        dt = max(1e-9, elapsed_s - float(getattr(self, "_last_iter_time", 0.0)))
        dd = dist_total - float(getattr(self, "_last_iter_dist", 0.0))
        v_robot = dd / dt

        self._last_iter_time = elapsed_s
        self._last_iter_dist = dist_total

        self.fitness_history.append({
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "iter": int(self.search_iter),
            "algorithm": self.algo_var.get(),
            "best_humidity": float(bv), "best_x": float(bx), "best_y": float(by),
            "humidity_now": float(self.last_humidity),
            "distance_m": float(dist_total), "elapsed_s": float(elapsed_s),
            "delta_distance_m": float(dd), "delta_time_s": float(dt), "robot_speed_m_s": float(v_robot)
        })

    def on_reach_goal(self):
        gx, gy = self.current_goal
        if gx is None:
            return

        self.sample_humidity_at_pose()
        self.record_fitness_iteration()

        self.search_iter += 1
        try:
            max_iters = int(self.max_iters_var.get())
        except Exception:
            max_iters = 120

        try:
            h_stop = float(self.stop_h_thresh_var.get())
        except Exception:
            h_stop = 0.995

        bx, by, bv = self.best_found
        if bv >= h_stop or self.search_iter >= max_iters:
            self.set_status(f"Búsqueda finalizada. Mejor h={bv:.3f} en ({bx:.2f},{by:.2f}).")
            self.current_goal = (None, None)
            self.waypoints_runtime = []
            self.wp_index = 0
            self.stop_event.set()
            return

        self.current_goal = (None, None)
        self.waypoints_runtime = []
        self.wp_index = 0

    # ================= Control del algoritmo =================
    def algo_next_goal(self):
        if not self.hmap.is_ready():
            self.set_status("Mapa no está listo.")
            return None

        # Validación: asegurar límites aplicados antes de usar el objetivo
        self.apply_bounds_and_reinit()

        self.update_morphology_runtime()

        if self.search.algo != self.algo_var.get():
            self.search.set_algo(self.algo_var.get())

        g = self.search.next_goal()
        if g is None:
            return None

        # Clamping final (doble barrera)
        gx, gy = self.clip_xy(g[0], g[1])
        return (gx, gy)

    # ================= Ejecución =================
    def start_search(self):
        if self.require_serial() and (not self.ser or not self.ser.is_open):
            messagebox.showwarning("Inicio", "Conecte el Maestro o desactive 'Enviar comandos al robot'.")
            return
        if not self.hmap.is_ready():
            messagebox.showwarning("Inicio", "Genere primero el mapa de humedad.")
            return

        self.update_morphology_runtime()

        # Validación/clamping de todo el estado según límites actuales
        self.apply_bounds_and_reinit()

        self.run_config_snapshot = self.snapshot_config()

        self.fitness_history = []
        self.speed_acc = {}
        self.last_cmd_pos = {}
        self.speed_history = []
        self.humidity_series = []
        self.robot_series = []
        self.distance_m = 0.0
        self._last_iter_dist = 0.0
        self._last_iter_time = 0.0

        # reset pose desde A (ya clampeado)
        x0 = float(self.ax_var.get())
        y0 = float(self.ay_var.get())
        th0 = theta_from_alignment(self.align_var.get().strip())
        self.pose_runtime = (x0, y0, th0)
        self.traj_x = [x0]
        self.traj_y = [y0]
        self.t0 = time.perf_counter()

        hv0 = self.hmap.value_bilinear(x0, y0)
        self.last_humidity = hv0
        self.best_found = (x0, y0, hv0)
        self.search_iter = 0

        self.current_goal = (None, None)
        self.waypoints_runtime = []
        self.wp_index = 0

        self.search.set_algo(self.algo_var.get())

        if self.nav_thread and self.nav_thread.is_alive():
            return

        self.stop_event.clear()
        self.nav_thread = NavigateThread(self, self.stop_event)
        self.nav_thread.start()
        self.set_status("Búsqueda iniciada.")

        if not (self.plot_win and self.plot_win.winfo_exists()):
            self.open_plot_window()
        else:
            self.plot_win.lift()

    def stop(self):
        self.stop_event.set()
        self.set_status("Detenido")

    def robot_avg_speed_m_s(self):
        dur = time.perf_counter() - self.t0
        if dur <= 1e-9:
            return 0.0
        return self.distance_m / dur

    # ================= Exportar / limpiar histórico =================
    def clear_history(self):
        self.fitness_history = []
        self.speed_history = []
        self.speed_acc = {}
        self.last_cmd_pos = {}
        self.humidity_series = []
        self.robot_series = []
        self.run_config_snapshot = None
        self.set_status("Histórico limpiado.")

    def _autosize_ws(self, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    v = "" if cell.value is None else str(cell.value)
                    max_len = max(max_len, len(v))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(60, max(10, max_len + 2))

    def export_history(self):
        if self.run_config_snapshot is None and not self.fitness_history and not self.speed_history:
            messagebox.showinfo("Exportar", "No hay histórico para exportar.")
            return

        folder = filedialog.askdirectory(title="Seleccione carpeta para exportar")
        if not folder:
            return

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        base = os.path.join(folder, f"historico_busqueda_{ts}")

        data = {
            "config": self.run_config_snapshot,
            "fitness_history": self.fitness_history,
            "speed_history": self.speed_history,
            "speed_avg_final_by_id": {str(rid): self.speed_acc[rid]["avg"] for rid in self.speed_acc} if self.speed_acc else {},
            "distance_m_final": self.distance_m,
            "robot_avg_speed_m_s": self.robot_avg_speed_m_s(),
            "best_found": {"x": self.best_found[0], "y": self.best_found[1], "h": self.best_found[2]},
            "humidity_series": self.humidity_series,
            "robot_series": self.robot_series
        }
        json_path = base + ".json"
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

        xlsx_path = base + ".xlsx"
        wb = Workbook()

        ws = wb.active
        ws.title = "Config"
        ws["A1"] = "config_json"
        ws["A2"] = json.dumps(self.run_config_snapshot, ensure_ascii=False, indent=2) if self.run_config_snapshot else ""
        self._autosize_ws(ws)

        ws_it = wb.create_sheet("Iteraciones")
        headers_it = [
            "timestamp", "iter", "algorithm",
            "best_humidity", "best_x", "best_y",
            "humidity_now",
            "distance_m", "elapsed_s",
            "delta_distance_m", "delta_time_s", "robot_speed_m_s"
        ]
        ws_it.append(headers_it)
        for r in self.fitness_history:
            ws_it.append([r.get(h, None) for h in headers_it])
        self._autosize_ws(ws_it)

        ws_sp = wb.create_sheet("SpeedSeries")
        ids = sorted(self.speed_acc.keys()) if self.speed_acc else []
        headers_sp = ["t", "distance_m"] + [f"id_{rid}_avg_ticks_per_s" for rid in ids]
        ws_sp.append(headers_sp)
        for s in self.speed_history:
            row = [s.get("t", None), s.get("distance_m", None)]
            avg_by_id = s.get("avg_by_id", {})
            for rid in ids:
                row.append(avg_by_id.get(rid, None))
            ws_sp.append(row)
        self._autosize_ws(ws_sp)

        ws_sf = wb.create_sheet("SpeedAvgFinal")
        ws_sf.append(["id", "avg_ticks_per_s"])
        for rid in ids:
            ws_sf.append([rid, float(self.speed_acc[rid]["avg"])])
        self._autosize_ws(ws_sf)

        ws_rs = wb.create_sheet("RobotSeries")
        ws_rs.append(["t", "distance_m", "v_m_s"])
        for s in self.robot_series:
            ws_rs.append([s.get("t"), s.get("distance_m"), s.get("v_m_s")])
        self._autosize_ws(ws_rs)

        ws_hs = wb.create_sheet("HumiditySeries")
        ws_hs.append(["t", "x", "y", "h", "distance_m"])
        for s in self.humidity_series:
            ws_hs.append([s.get("t"), s.get("x"), s.get("y"), s.get("h"), s.get("distance_m")])
        self._autosize_ws(ws_hs)

        ws_sum = wb.create_sheet("Resumen")
        bx, by, bh = self.best_found
        duration = time.perf_counter() - self.t0
        robot_avg_speed = (self.distance_m / duration) if duration > 1e-9 else 0.0
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.get_bounds()

        ws_sum.append(["metric", "value"])
        ws_sum.append(["bounds_xmin", X_MIN])
        ws_sum.append(["bounds_xmax", X_MAX])
        ws_sum.append(["bounds_ymin", Y_MIN])
        ws_sum.append(["bounds_ymax", Y_MAX])
        ws_sum.append(["algorithm", self.algo_var.get()])
        ws_sum.append(["iterations", self.search_iter])
        ws_sum.append(["best_humidity", bh])
        ws_sum.append(["best_x", bx])
        ws_sum.append(["best_y", by])
        ws_sum.append(["distance_m_total", self.distance_m])
        ws_sum.append(["duration_s", duration])
        ws_sum.append(["robot_avg_speed_m_s", robot_avg_speed])
        ws_sum.append(["num_modules", len(self.ids_runtime)])
        self._autosize_ws(ws_sum)

        wb.save(xlsx_path)
        self.set_status(f"Exportado: {os.path.basename(xlsx_path)} (+ JSON).")

    # ================= Serial / cierre =================
    def on_close(self):
        try:
            self.disconnect()
        except Exception:
            pass
        self.destroy()


if __name__ == "__main__":
    App().mainloop()
