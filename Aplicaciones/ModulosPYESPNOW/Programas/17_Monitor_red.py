# -*- coding: utf-8 -*-
"""
Monitor experimental ESP-NOW -> HUB -> USB/Serial -> PC.

Dependencias:
    pip install pyserial openpyxl

Formatos recibidos admitidos:
    id,humedad
    id,prox_f,prox_i,prox_d
    DATA,run_id,node_id,seq,tx_us,rx_hub_us,tipo,v1,v2,v3,bytes,rssi,status
    ACK,run_id,node_id,seq,pc_tx_ns[,node_rx_us]

El formato DATA permite calcular pérdida por saltos de secuencia. El formato ACK
permite calcular RTT sin sincronizar relojes. Con los formatos heredados solo se
calculan tasa de recepción, periodo entre llegadas y disponibilidad.

Comando PING generado (requiere soporte en el firmware del HUB/nodo):
    PING,run_id,node_id,seq,pc_tx_ns
"""

from __future__ import annotations

import math
import queue
import statistics
import threading
import time
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

import serial
import serial.tools.list_ports
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass
class Packet:
    run_id: str
    node_id: int
    seq: Optional[int]
    pc_rx_ns: int
    elapsed_s: float
    interarrival_ms: Optional[float]
    tx_us: Optional[int] = None
    hub_rx_us: Optional[int] = None
    packet_type: str = "LEGACY"
    v1: Optional[float] = None
    v2: Optional[float] = None
    v3: Optional[float] = None
    payload_bytes: Optional[int] = None
    rssi_dbm: Optional[float] = None
    status: str = "RX"
    rtt_ms: Optional[float] = None
    raw: str = ""


@dataclass
class Event:
    elapsed_s: float
    level: str
    node_id: Optional[int]
    event: str
    detail: str


def fnum(value: str) -> Optional[float]:
    value = value.strip()
    return None if value == "" else float(value)


def inum(value: str) -> Optional[int]:
    value = value.strip()
    return None if value == "" else int(value)


def mean(values: List[float]) -> Optional[float]:
    return statistics.fmean(values) if values else None


def stdev(values: List[float]) -> Optional[float]:
    return statistics.stdev(values) if len(values) >= 2 else None


def percentile(values: List[float], p: float) -> Optional[float]:
    if not values:
        return None
    data = sorted(values)
    pos = (len(data) - 1) * p
    lo, hi = math.floor(pos), math.ceil(pos)
    if lo == hi:
        return data[lo]
    return data[lo] + (data[hi] - data[lo]) * (pos - lo)


class NetworkMeasurement:
    def __init__(self):
        self.lock = threading.Lock()
        self.running = False
        self.started_ns = 0
        self.stopped_ns = 0
        self.config: Dict[str, object] = {}
        self.packets: List[Packet] = []
        self.events: List[Event] = []
        self.last_rx_ns: Dict[int, int] = {}
        self.sequences: Dict[int, set] = defaultdict(set)
        self.ping_sent_ns: Dict[tuple, int] = {}

    def start(self, config: Dict[str, object]):
        with self.lock:
            self.running = True
            self.started_ns = time.perf_counter_ns()
            self.stopped_ns = 0
            self.config = dict(config)
            self.packets.clear()
            self.events.clear()
            self.last_rx_ns.clear()
            self.sequences.clear()
            self.ping_sent_ns.clear()

    def stop(self):
        with self.lock:
            self.running = False
            self.stopped_ns = time.perf_counter_ns()

    def elapsed(self, now_ns: Optional[int] = None) -> float:
        if not self.started_ns:
            return 0.0
        end = now_ns or self.stopped_ns or time.perf_counter_ns()
        return max(0.0, (end - self.started_ns) / 1e9)

    def add_event(self, level: str, event: str, detail: str, node_id=None):
        with self.lock:
            self.events.append(Event(self.elapsed(), level, node_id, event, detail))

    def register_ping(self, run_id: str, node_id: int, seq: int, sent_ns: int):
        with self.lock:
            self.ping_sent_ns[(run_id, node_id, seq)] = sent_ns

    def ingest(self, line: str):
        if not self.running:
            return
        now_ns = time.perf_counter_ns()
        parts = [x.strip() for x in line.strip().split(",")]
        try:
            if parts[0].upper() == "DATA" and len(parts) >= 13:
                packet = Packet(
                    run_id=parts[1] or str(self.config.get("run_id", "")),
                    node_id=int(parts[2]), seq=inum(parts[3]),
                    pc_rx_ns=now_ns, elapsed_s=self.elapsed(now_ns), interarrival_ms=None,
                    tx_us=inum(parts[4]), hub_rx_us=inum(parts[5]), packet_type=parts[6],
                    v1=fnum(parts[7]), v2=fnum(parts[8]), v3=fnum(parts[9]),
                    payload_bytes=inum(parts[10]), rssi_dbm=fnum(parts[11]),
                    status=parts[12], raw=line,
                )
            elif parts[0].upper() == "ACK" and len(parts) >= 5:
                run_id, node_id, seq = parts[1], int(parts[2]), int(parts[3])
                sent_ns = int(parts[4])
                rtt_ms = (now_ns - sent_ns) / 1e6
                packet = Packet(run_id, node_id, seq, now_ns, self.elapsed(now_ns), None,
                                packet_type="ACK", status="ACK", rtt_ms=rtt_ms, raw=line)
            elif len(parts) in (2, 4):
                node_id = int(parts[0])
                vals = [fnum(x) for x in parts[1:]]
                packet = Packet(str(self.config.get("run_id", "")), node_id, None,
                                now_ns, self.elapsed(now_ns), None,
                                packet_type="HUM" if len(parts) == 2 else "PROX",
                                v1=vals[0], v2=vals[1] if len(vals) > 1 else None,
                                v3=vals[2] if len(vals) > 2 else None, raw=line)
            else:
                self.add_event("WARN", "LINEA_INVALIDA", line)
                return
        except (ValueError, IndexError) as exc:
            self.add_event("WARN", "ERROR_PARSE", f"{line} | {exc}")
            return

        with self.lock:
            previous = self.last_rx_ns.get(packet.node_id)
            if previous is not None:
                packet.interarrival_ms = (now_ns - previous) / 1e6
            self.last_rx_ns[packet.node_id] = now_ns

            if packet.seq is not None:
                if packet.seq in self.sequences[packet.node_id]:
                    self.events.append(Event(packet.elapsed_s, "WARN", packet.node_id,
                                             "DUPLICADO", f"Secuencia {packet.seq}"))
                self.sequences[packet.node_id].add(packet.seq)
            self.packets.append(packet)

    def snapshot(self):
        with self.lock:
            return list(self.packets), list(self.events), dict(self.config)

    def summaries(self):
        packets, _, cfg = self.snapshot()
        duration = self.elapsed()
        expected_period_ms = float(cfg.get("period_ms", 0) or 0)
        expected_duration = float(cfg.get("duration_s", duration) or duration)
        by_node = defaultdict(list)
        for p in packets:
            by_node[p.node_id].append(p)

        rows = []
        for node_id, items in sorted(by_node.items()):
            items.sort(key=lambda p: p.pc_rx_ns)
            intervals = [p.interarrival_ms for p in items if p.interarrival_ms is not None]
            rtts = [p.rtt_ms for p in items if p.rtt_ms is not None]
            seqs = sorted({p.seq for p in items if p.seq is not None and p.packet_type != "ACK"})
            if seqs:
                expected = seqs[-1] - seqs[0] + 1
                received = len(seqs)
                lost = max(0, expected - received)
            elif expected_period_ms > 0:
                expected = max(1, round(expected_duration * 1000 / expected_period_ms))
                received = len([p for p in items if p.packet_type != "ACK"])
                lost = None  # estimación temporal no equivale a pérdida verificable
            else:
                expected, received, lost = None, len(items), None

            data_items = [p for p in items if p.packet_type != "ACK"]
            first_last_s = ((data_items[-1].pc_rx_ns - data_items[0].pc_rx_ns) / 1e9
                            if len(data_items) >= 2 else None)
            frequency = ((len(data_items) - 1) / first_last_s
                         if first_last_s and first_last_s > 0 else None)
            useful_bits = sum((p.payload_bytes or 0) * 8 for p in data_items)
            throughput = useful_bits / first_last_s if first_last_s and useful_bits else None
            jitter = mean([abs(intervals[i] - intervals[i - 1])
                           for i in range(1, len(intervals))])
            pdr = received / expected if expected and lost is not None else None
            rows.append({
                "run_id": cfg.get("run_id"), "node_id": node_id,
                "expected": expected, "received": received, "lost": lost, "pdr": pdr,
                "period_mean_ms": mean(intervals), "period_sd_ms": stdev(intervals),
                "period_min_ms": min(intervals) if intervals else None,
                "period_max_ms": max(intervals) if intervals else None,
                "jitter_mean_ms": jitter, "frequency_hz": frequency,
                "rtt_mean_ms": mean(rtts), "rtt_sd_ms": stdev(rtts),
                "rtt_p95_ms": percentile(rtts, 0.95),
                "rtt_max_ms": max(rtts) if rtts else None,
                "throughput_bps": throughput,
                "rssi_mean_dbm": mean([p.rssi_dbm for p in data_items if p.rssi_dbm is not None]),
            })
        return rows


class SerialWorker:
    def __init__(self, measurement: NetworkMeasurement, log_q: queue.Queue):
        self.measurement = measurement
        self.log_q = log_q
        self.ser: Optional[serial.Serial] = None
        self.stop_evt = threading.Event()
        self.thread: Optional[threading.Thread] = None

    def open(self, port: str, baud: int):
        self.close()
        self.ser = serial.Serial(port, baud, timeout=0.2, write_timeout=0.5)
        self.stop_evt.clear()
        self.thread = threading.Thread(target=self._read_loop, daemon=True)
        self.thread.start()
        self.log_q.put(f"Conectado a {port} @ {baud}")

    def close(self):
        self.stop_evt.set()
        if self.thread and self.thread.is_alive():
            self.thread.join(timeout=1)
        if self.ser and self.ser.is_open:
            self.ser.close()
        self.ser = None

    def send(self, text: str):
        if not self.ser or not self.ser.is_open:
            raise RuntimeError("El puerto serial no está conectado")
        self.ser.write((text.rstrip() + "\n").encode("utf-8"))

    def _read_loop(self):
        while not self.stop_evt.is_set() and self.ser:
            try:
                raw = self.ser.readline()
                if raw:
                    line = raw.decode("utf-8", errors="replace").strip()
                    if line.startswith("#"):
                        self.log_q.put(line)
                    elif line:
                        self.measurement.ingest(line)
            except Exception as exc:
                self.log_q.put(f"ERROR serial: {exc}")
                time.sleep(0.2)


def style_sheet(ws):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        width = min(28, max(10, max(len(str(c.value or "")) for c in column) + 2))
        ws.column_dimensions[get_column_letter(column[0].column)].width = width


def export_excel(path: Path, measurement: NetworkMeasurement):
    packets, events, cfg = measurement.snapshot()
    summary = measurement.summaries()
    wb = Workbook()
    ws_cfg = wb.active
    ws_cfg.title = "Configuracion"
    ws_cfg.append(["Parametro", "Valor"])
    for key, value in cfg.items():
        ws_cfg.append([key, value])
    ws_cfg.append(["duration_effective_s", measurement.elapsed()])

    ws = wb.create_sheet("Paquetes")
    packet_headers = list(asdict(Packet("", 0, None, 0, 0, None)).keys())
    ws.append(packet_headers)
    for packet in packets:
        ws.append([asdict(packet)[h] for h in packet_headers])

    ws_sum = wb.create_sheet("Resumen_nodo")
    summary_headers = list(summary[0].keys()) if summary else ["run_id", "node_id"]
    ws_sum.append(summary_headers)
    for row in summary:
        ws_sum.append([row[h] for h in summary_headers])

    ws_evt = wb.create_sheet("Eventos")
    event_headers = list(asdict(Event(0, "", None, "", "")).keys())
    ws_evt.append(event_headers)
    for event in events:
        ws_evt.append([asdict(event)[h] for h in event_headers])

    if len(packets) > 1:
        ws_chart = wb.create_sheet("Graficas")
        ws_chart.append(["Muestra", "Interarribo_ms", "RTT_ms"])
        for i, p in enumerate(packets, 1):
            ws_chart.append([i, p.interarrival_ms, p.rtt_ms])
        chart = LineChart()
        chart.title = "Comportamiento temporal de la comunicación"
        chart.y_axis.title = "Tiempo (ms)"
        chart.x_axis.title = "Muestra"
        chart.add_data(Reference(ws_chart, min_col=2, max_col=3, min_row=1,
                                 max_row=ws_chart.max_row), titles_from_data=True)
        chart.set_categories(Reference(ws_chart, min_col=1, min_row=2,
                                       max_row=ws_chart.max_row))
        chart.height, chart.width = 10, 20
        ws_chart.add_chart(chart, "E2")

    for sheet in wb.worksheets:
        style_sheet(sheet)
    wb.save(path)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Medición de red ESP-NOW")
        self.geometry("1100x720")
        self.measurement = NetworkMeasurement()
        self.log_q = queue.Queue()
        self.serial = SerialWorker(self.measurement, self.log_q)
        self.ping_seq = 0
        self.ping_job = None
        self._variables()
        self._ui()
        self._refresh_ports()
        self.after(250, self._tick)
        self.protocol("WM_DELETE_WINDOW", self._close)

    def _variables(self):
        self.port = tk.StringVar()
        self.baud = tk.IntVar(value=115200)
        self.run_id = tk.StringVar(value=datetime.now().strftime("R%Y%m%d_%H%M%S"))
        self.modules = tk.StringVar(value="1,2,3")
        self.morphology = tk.StringVar(value="Cadena")
        self.period_ms = tk.DoubleVar(value=50)
        self.duration_s = tk.DoubleVar(value=60)
        self.payload_bytes = tk.IntVar(value=32)
        self.channel = tk.IntVar(value=1)
        self.distance_m = tk.DoubleVar(value=1.0)
        self.state = tk.StringVar(value="Detenido")
        self.repetition = tk.IntVar(value=1)
        self.enable_ping = tk.BooleanVar(value=False)

    def _ui(self):
        conn = ttk.LabelFrame(self, text="Conexión serial", padding=8)
        conn.pack(fill="x", padx=10, pady=6)
        self.port_cb = ttk.Combobox(conn, textvariable=self.port, width=18)
        self.port_cb.grid(row=0, column=0, padx=4)
        ttk.Entry(conn, textvariable=self.baud, width=10).grid(row=0, column=1, padx=4)
        ttk.Button(conn, text="Refrescar", command=self._refresh_ports).grid(row=0, column=2, padx=4)
        ttk.Button(conn, text="Conectar", command=self._connect).grid(row=0, column=3, padx=4)
        ttk.Button(conn, text="Desconectar", command=self.serial.close).grid(row=0, column=4, padx=4)

        cfg = ttk.LabelFrame(self, text="Configuración experimental", padding=8)
        cfg.pack(fill="x", padx=10, pady=6)
        fields = [
            ("Run ID", self.run_id), ("IDs módulos", self.modules),
            ("Morfología", self.morphology), ("Periodo (ms)", self.period_ms),
            ("Duración (s)", self.duration_s), ("Payload (bytes)", self.payload_bytes),
            ("Canal", self.channel), ("Distancia (m)", self.distance_m),
            ("Estado", self.state), ("Repetición", self.repetition),
        ]
        for i, (label, var) in enumerate(fields):
            row, col = divmod(i, 5)
            box = ttk.Frame(cfg)
            box.grid(row=row, column=col, padx=6, pady=4, sticky="w")
            ttk.Label(box, text=label).pack(anchor="w")
            if label == "Morfología":
                ttk.Combobox(box, textvariable=var, width=17,
                             values=["Cadena", "L", "T", "Cruz", "Cuadrúpedo"]).pack()
            elif label == "Estado":
                ttk.Combobox(box, textvariable=var, width=17,
                             values=["Detenido", "Movimiento"]).pack()
            else:
                ttk.Entry(box, textvariable=var, width=19).pack()

        actions = ttk.Frame(self)
        actions.pack(fill="x", padx=10, pady=6)
        ttk.Checkbutton(actions, text="Enviar PING periódico", variable=self.enable_ping).pack(side="left")
        ttk.Button(actions, text="Iniciar medición", command=self._start).pack(side="left", padx=8)
        ttk.Button(actions, text="Detener", command=self._stop).pack(side="left", padx=8)
        ttk.Button(actions, text="Exportar Excel", command=self._export).pack(side="left", padx=8)
        self.status = ttk.Label(actions, text="Sin medición")
        self.status.pack(side="right")

        self.tree = ttk.Treeview(self, columns=("node", "rx", "pdr", "period", "jitter", "rtt", "freq"),
                                 show="headings", height=12)
        headings = [("node", "Nodo"), ("rx", "Recibidos"), ("pdr", "PDR (%)"),
                    ("period", "Periodo medio (ms)"), ("jitter", "Jitter (ms)"),
                    ("rtt", "RTT medio (ms)"), ("freq", "Frecuencia (Hz)")]
        for key, label in headings:
            self.tree.heading(key, text=label)
            self.tree.column(key, width=145, anchor="center")
        self.tree.pack(fill="both", expand=True, padx=10, pady=6)

        logf = ttk.LabelFrame(self, text="Registro", padding=6)
        logf.pack(fill="both", padx=10, pady=6)
        self.log = tk.Text(logf, height=8, state="disabled")
        self.log.pack(fill="both")

    def _refresh_ports(self):
        ports = [p.device for p in serial.tools.list_ports.comports()]
        self.port_cb["values"] = ports
        if ports and not self.port.get():
            self.port.set(ports[0])

    def _connect(self):
        try:
            self.serial.open(self.port.get(), int(self.baud.get()))
        except Exception as exc:
            messagebox.showerror("Conexión", str(exc))

    def _module_ids(self):
        return [int(x.strip()) for x in self.modules.get().split(",") if x.strip()]

    def _config(self):
        ids = self._module_ids()
        return {
            "run_id": self.run_id.get().strip(), "timestamp": datetime.now().isoformat(timespec="seconds"),
            "module_ids": ",".join(map(str, ids)), "num_modules": len(ids),
            "morphology": self.morphology.get(), "period_ms": float(self.period_ms.get()),
            "duration_s": float(self.duration_s.get()), "payload_bytes": int(self.payload_bytes.get()),
            "channel": int(self.channel.get()), "distance_m": float(self.distance_m.get()),
            "state": self.state.get(), "repetition": int(self.repetition.get()),
            "port": self.port.get(), "baud": int(self.baud.get()),
            "ping_enabled": bool(self.enable_ping.get()),
        }

    def _start(self):
        try:
            cfg = self._config()
            if not cfg["run_id"] or not self._module_ids():
                raise ValueError("Indique Run ID y al menos un módulo")
            self.measurement.start(cfg)
            self.status.config(text="Midiendo...")
            if self.enable_ping.get():
                self._send_pings()
            self.after(round(float(self.duration_s.get()) * 1000), self._auto_stop)
        except Exception as exc:
            messagebox.showerror("Inicio", str(exc))

    def _send_pings(self):
        if not self.measurement.running or not self.enable_ping.get():
            return
        for node_id in self._module_ids():
            self.ping_seq += 1
            sent_ns = time.perf_counter_ns()
            cmd = f"PING,{self.run_id.get()},{node_id},{self.ping_seq},{sent_ns}"
            try:
                self.serial.send(cmd)
                self.measurement.register_ping(self.run_id.get(), node_id, self.ping_seq, sent_ns)
            except Exception as exc:
                self.measurement.add_event("ERROR", "PING_TX", str(exc), node_id)
        self.ping_job = self.after(max(10, round(float(self.period_ms.get()))), self._send_pings)

    def _auto_stop(self):
        if self.measurement.running:
            self._stop()

    def _stop(self):
        if self.ping_job:
            self.after_cancel(self.ping_job)
            self.ping_job = None
        self.measurement.stop()
        self.status.config(text=f"Detenida: {len(self.measurement.packets)} registros")

    def _export(self):
        if not self.measurement.packets:
            messagebox.showwarning("Exportar", "No hay datos para exportar")
            return
        default = f"red_espnow_{self.run_id.get()}.xlsx"
        filename = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile=default,
                                                filetypes=[("Excel", "*.xlsx")])
        if filename:
            try:
                export_excel(Path(filename), self.measurement)
                messagebox.showinfo("Exportar", f"Archivo guardado:\n{filename}")
            except Exception as exc:
                messagebox.showerror("Exportar", str(exc))

    def _tick(self):
        for row in self.tree.get_children():
            self.tree.delete(row)
        for r in self.measurement.summaries():
            def fmt(v, d=2): return "—" if v is None else f"{v:.{d}f}"
            self.tree.insert("", "end", values=(r["node_id"], r["received"],
                             fmt(r["pdr"] * 100 if r["pdr"] is not None else None),
                             fmt(r["period_mean_ms"]), fmt(r["jitter_mean_ms"]),
                             fmt(r["rtt_mean_ms"]), fmt(r["frequency_hz"])))
        while True:
            try:
                line = self.log_q.get_nowait()
            except queue.Empty:
                break
            self.log.config(state="normal")
            self.log.insert("end", line + "\n")
            self.log.see("end")
            self.log.config(state="disabled")
        self.after(250, self._tick)

    def _close(self):
        self.serial.close()
        self.destroy()


if __name__ == "__main__":
    App().mainloop()
