# -*- coding: utf-8 -*-

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import threading
import time
import math
import random
import json
import os
import statistics
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime
from typing import Optional

import numpy as np

import serial
import serial.tools.list_ports

from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.ticker import MultipleLocator
from matplotlib.patches import Rectangle

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


# ======================================================
# Medición integrada de comunicación ESP-NOW
# ======================================================
@dataclass
class NetPacket:
    run_id: str
    node_id: int
    seq: Optional[int]
    pc_rx_ns: int
    elapsed_s: float
    interarrival_ms: Optional[float]
    tx_us: Optional[int] = None
    hub_rx_us: Optional[int] = None
    packet_type: str = "LEGACY"
    v1: Optional[float] = None
    v2: Optional[float] = None
    v3: Optional[float] = None
    payload_bytes: Optional[int] = None
    rssi_dbm: Optional[float] = None
    status: str = "RX"
    rtt_ms: Optional[float] = None
    raw: str = ""


@dataclass
class NetEvent:
    elapsed_s: float
    level: str
    node_id: Optional[int]
    event: str
    detail: str


def _opt_float(value):
    value = str(value).strip()
    return None if value == "" else float(value)


def _opt_int(value):
    value = str(value).strip()
    return None if value == "" else int(value)


def _mean(values):
    return statistics.fmean(values) if values else None


def _stdev(values):
    return statistics.stdev(values) if len(values) >= 2 else None


def _percentile(values, p):
    if not values:
        return None
    data = sorted(values)
    pos = (len(data) - 1) * p
    lo, hi = math.floor(pos), math.ceil(pos)
    if lo == hi:
        return data[lo]
    return data[lo] + (data[hi] - data[lo]) * (pos - lo)


class NetworkMeasurement:
    """Registra la comunicación durante la misma ventana del algoritmo."""
    def __init__(self):
        self.lock = threading.Lock()
        self.running = False
        self.started_ns = 0
        self.stopped_ns = 0
        self.config = {}
        self.packets = []
        self.events = []
        self.last_rx_ns = {}
        self.sequences = defaultdict(set)

    def start(self, config):
        with self.lock:
            self.running = True
            self.started_ns = time.perf_counter_ns()
            self.stopped_ns = 0
            self.config = dict(config)
            self.packets.clear()
            self.events.clear()
            self.last_rx_ns.clear()
            self.sequences.clear()

    def stop(self):
        with self.lock:
            if self.running:
                self.stopped_ns = time.perf_counter_ns()
            self.running = False

    def elapsed(self, now_ns=None):
        if not self.started_ns:
            return 0.0
        end = now_ns or self.stopped_ns or time.perf_counter_ns()
        return max(0.0, (end - self.started_ns) / 1e9)

    def add_event(self, level, event, detail, node_id=None):
        with self.lock:
            self.events.append(NetEvent(self.elapsed(), level, node_id, event, detail))

    def ingest(self, line):
        if not self.running or not line or line.startswith("#"):
            return False
        now_ns = time.perf_counter_ns()
        parts = [x.strip() for x in line.split(",")]
        try:
            if parts[0].upper() == "DATA" and len(parts) >= 13:
                run_id = parts[1]
                if not run_id or run_id == "SIN_RUN":
                    run_id = self.config.get("run_id", "")
                packet = NetPacket(
                    run_id, int(parts[2]), _opt_int(parts[3]), now_ns,
                    self.elapsed(now_ns), None, _opt_int(parts[4]),
                    _opt_int(parts[5]), parts[6], _opt_float(parts[7]),
                    _opt_float(parts[8]), _opt_float(parts[9]),
                    _opt_int(parts[10]), _opt_float(parts[11]), parts[12], None, line)
            elif parts[0].upper() == "ACK" and len(parts) >= 5:
                sent_ns = int(parts[4])
                packet = NetPacket(parts[1] or self.config.get("run_id", ""),
                                   int(parts[2]), int(parts[3]), now_ns,
                                   self.elapsed(now_ns), None, packet_type="ACK",
                                   status="ACK", rtt_ms=(now_ns - sent_ns) / 1e6,
                                   raw=line)
            else:
                return False
        except (ValueError, IndexError) as exc:
            self.add_event("WARN", "ERROR_PARSE", f"{line} | {exc}")
            return False

        with self.lock:
            previous = self.last_rx_ns.get(packet.node_id)
            if previous is not None:
                packet.interarrival_ms = (now_ns - previous) / 1e6
            self.last_rx_ns[packet.node_id] = now_ns
            if packet.seq is not None:
                if packet.seq in self.sequences[packet.node_id]:
                    self.events.append(NetEvent(packet.elapsed_s, "WARN", packet.node_id,
                                                "DUPLICADO", str(packet.seq)))
                self.sequences[packet.node_id].add(packet.seq)
            self.packets.append(packet)
        return True

    def snapshot(self):
        with self.lock:
            return list(self.packets), list(self.events), dict(self.config)

    def summaries(self):
        packets, _, cfg = self.snapshot()
        by_node = defaultdict(list)
        for packet in packets:
            by_node[packet.node_id].append(packet)
        rows = []
        for node_id, items in sorted(by_node.items()):
            items.sort(key=lambda p: p.pc_rx_ns)
            data = [p for p in items if p.packet_type != "ACK"]
            intervals = [p.interarrival_ms for p in data if p.interarrival_ms is not None]
            rtts = [p.rtt_ms for p in items if p.rtt_ms is not None]
            seqs = sorted({p.seq for p in data if p.seq is not None})
            expected = seqs[-1] - seqs[0] + 1 if seqs else None
            received = len(seqs) if seqs else len(data)
            lost = max(0, expected - received) if expected is not None else None
            duration = ((data[-1].pc_rx_ns - data[0].pc_rx_ns) / 1e9
                        if len(data) >= 2 else None)
            frequency = (len(data) - 1) / duration if duration and duration > 0 else None
            jitter = _mean([abs(intervals[i] - intervals[i - 1])
                            for i in range(1, len(intervals))])
            over_1s = sum(1 for value in intervals if value > 1000.0)
            over_1s_pct = over_1s / len(intervals) if intervals else None
            rates = cfg.get("sensor_rate_by_node", {})
            nominal_hz = float(rates.get(str(node_id),
                                         cfg.get("configured_sensor_rate_hz", 4.0)))
            nominal_ms = 1000.0 / nominal_hz
            expected_est = (round(duration * 1000.0 / nominal_ms) + 1
                            if duration and nominal_ms > 0 else None)
            lost_est = (max(0, expected_est - len(data))
                        if expected_est is not None else None)
            age_approx = (sum(value * value for value in intervals) /
                          (2.0 * sum(intervals))) if intervals and sum(intervals) > 0 else None
            bits = sum((p.payload_bytes or 0) * 8 for p in data)
            rssi = [p.rssi_dbm for p in data if p.rssi_dbm is not None]
            rows.append({
                "run_id": cfg.get("run_id"), "algorithm": cfg.get("algorithm"),
                "repetition": cfg.get("repetition"),
                "robot_state": cfg.get("robot_state"), "node_id": node_id,
                "expected": expected, "received": received, "lost": lost,
                "pdr": received / expected if expected else None,
                "expected_estimated": expected_est,
                "lost_estimated": lost_est,
                "pdr_estimated": (min(1.0, len(data) / expected_est)
                                  if expected_est else None),
                "pdr_source": ("NODE_SEQUENCE" if seqs else "NOMINAL_RATE_ESTIMATE"),
                "period_mean_ms": _mean(intervals),
                "period_median_ms": _percentile(intervals, 0.50),
                "period_sd_ms": _stdev(intervals),
                "period_p95_ms": _percentile(intervals, 0.95),
                "period_min_ms": min(intervals) if intervals else None,
                "period_max_ms": max(intervals) if intervals else None,
                "jitter_mean_ms": jitter, "frequency_hz": frequency,
                "period_cv": (_stdev(intervals) / _mean(intervals)
                              if _stdev(intervals) is not None and _mean(intervals) else None),
                "availability_pct": (1.0 - over_1s_pct
                                     if over_1s_pct is not None else None),
                "age_approx_mean_ms": age_approx,
                "intervals_gt_1s": over_1s,
                "intervals_gt_1s_pct": over_1s_pct,
                "rtt_mean_ms": _mean(rtts), "rtt_sd_ms": _stdev(rtts),
                "rtt_p95_ms": _percentile(rtts, 0.95),
                "rtt_max_ms": max(rtts) if rtts else None,
                "throughput_bps": bits / duration if duration and bits else None,
                "rssi_mean_dbm": _mean(rssi), "rssi_sd_db": _stdev(rssi),
                "rssi_min_dbm": min(rssi) if rssi else None,
                "rssi_max_dbm": max(rssi) if rssi else None
            })
        return rows


class RNANavigator:
    """Inferencia 11-32-16-8-4 exportada desde MATLAB a JSON."""
    ACTIONS = ("AV", "GI", "GD", "EV")

    def __init__(self, model_path):
        self.model_path = model_path
        self.layers = []
        self.xoffset = np.zeros(11)
        self.gain = np.ones(11)
        self.ymin = -np.ones(11)
        self.h = []
        self.previous_action = 1
        self.load()

    def load(self):
        with open(self.model_path, "r", encoding="utf-8") as fh:
            model = json.load(fh)
        self.xoffset = np.asarray(model["input_process"]["xoffset"], dtype=float)
        self.gain = np.asarray(model["input_process"]["gain"], dtype=float)
        self.ymin = np.asarray(model["input_process"]["ymin"], dtype=float)
        self.layers = [(np.asarray(x["W"], dtype=float),
                        np.asarray(x["b"], dtype=float)) for x in model["layers"]]

    def reset(self, h0=0.0):
        self.h = [float(h0)] * 3
        self.previous_action = 1

    @staticmethod
    def _tansig(z):
        return 2.0 / (1.0 + np.exp(-2.0 * np.clip(z, -40, 40))) - 1.0

    def features(self, humidity, prox, signal_c=0.0):
        self.h.append(float(humidity))
        self.h = self.h[-3:]
        h2, h1, h0 = self.h[-3], self.h[-2], self.h[-1]
        dh = h0 - h1
        mean_h = float(np.mean(self.h))
        slope = (h0 - h2) / 2.0
        pf, pl, pr = prox
        return np.asarray([h0, h1, h2, dh, mean_h, slope,
                           pf, pl, pr, signal_c,
                           float(self.previous_action)], dtype=float)

    def predict(self, x):
        a = (x - self.xoffset) * self.gain + self.ymin
        for index, (W, b) in enumerate(self.layers):
            z = W @ a + b
            a = self._tansig(z) if index < len(self.layers) - 1 else z
        exp = np.exp(a - np.max(a))
        probabilities = exp / np.sum(exp)
        action_index = int(np.argmax(probabilities))
        self.previous_action = action_index + 1
        return self.ACTIONS[action_index], probabilities.tolist(), x.tolist()


# =========================
# Parámetros base (servos)
# =========================
HOME_DX = 512
HOME_LX = 95

DX_MIN, DX_MAX = 0, 1023
LX_MIN, LX_MAX = 0, 240

LX_PER_DX = 0.22
PHI_H_DEG_DEFAULT = 90.0


# =========================
# Utilidades
# =========================
def list_serial_ports():
    return [p.device for p in serial.tools.list_ports.comports()]


def clamp_int(x: float, lo: int, hi: int) -> int:
    return max(lo, min(int(round(x)), hi))


def clamp_float(x: float, lo: float, hi: float) -> float:
    return max(lo, min(float(x), hi))


def deg2rad(deg: float) -> float:
    return deg * math.pi / 180.0


def wrap_pi(a: float) -> float:
    while a > math.pi:
        a -= 2.0 * math.pi
    while a < -math.pi:
        a += 2.0 * math.pi
    return a


def theta_from_alignment(aln: str) -> float:
    if aln == "+X":
        return 0.0
    if aln == "+Y":
        return math.pi / 2.0
    if aln == "-X":
        return math.pi
    if aln == "-Y":
        return -math.pi / 2.0
    return 0.0


def parse_id_list(s: str):
    s = s.strip()
    if not s:
        return []
    out = []
    for part in s.split(","):
        p = part.strip()
        if not p:
            continue
        out.append(int(p))
    return out


def build_manhattan_waypoints(ax, ay, bx, by, step_m, x_first=True, bounds=(0.0, 1.0, 0.0, 1.0)):
    """
    Waypoints Manhattan desde A a B, respetando límites (clamp en cada punto).
    bounds = (X_MIN, X_MAX, Y_MIN, Y_MAX)
    """
    X_MIN, X_MAX, Y_MIN, Y_MAX = bounds

    def clip(x, y):
        return (clamp_float(x, X_MIN, X_MAX), clamp_float(y, Y_MIN, Y_MAX))

    ax, ay = clip(ax, ay)
    bx, by = clip(bx, by)

    if step_m <= 0:
        step_m = 0.001

    wps = []
    x, y = ax, ay

    def walk_axis(cur, target):
        if abs(target - cur) < 1e-12:
            return []
        sgn = 1.0 if target > cur else -1.0
        pts = []
        v = cur
        while True:
            nxt = v + sgn * step_m
            if (sgn > 0 and nxt >= target) or (sgn < 0 and nxt <= target):
                pts.append(target)
                break
            pts.append(nxt)
            v = nxt
        return pts

    if x_first:
        for nx in walk_axis(x, bx):
            x, y = clip(nx, y)
            wps.append((x, y))
        for ny in walk_axis(y, by):
            x, y = clip(x, ny)
            wps.append((x, y))
    else:
        for ny in walk_axis(y, by):
            x, y = clip(x, ny)
            wps.append((x, y))
        for nx in walk_axis(x, bx):
            x, y = clip(nx, y)
            wps.append((x, y))

    if not wps or (wps[-1][0] != bx or wps[-1][1] != by):
        wps.append((bx, by))
    return wps


# =========================
# Morfologías (dibujo)
# =========================
def morphology_positions(name: str, n: int):
    n = max(1, int(n))
    pts = []

    if name == "Serpiente":
        for i in range(n):
            pts.append((i, 0))
        return pts

    if name == "L":
        k = max(2, min(n, n // 2 + 1))
        for i in range(k):
            pts.append((i, 0))
        x0, y0 = pts[-1]
        for j in range(1, n - k + 1):
            pts.append((x0, -j))
        return pts

    if name == "T":
        tv = max(2, min(n, n // 2 + 1))
        th = n - tv
        for j in range(tv):
            pts.append((0, -j))
        if th <= 0:
            return pts
        used = 0
        step = 1
        while used < th:
            pts.append((step, 0))
            used += 1
            if used >= th:
                break
            pts.append((-step, 0))
            used += 1
            step += 1
        return pts[:n]

    if name == "Cuadrúpedo":
        pts.append((0, 0))
        if n == 1:
            return pts
        legs = [(1, 0), (-1, 0), (0, 1), (0, -1)]
        i = 0
        while len(pts) < n and i < len(legs):
            pts.append(legs[i])
            i += 1
        if len(pts) >= n:
            return pts[:n]
        base = [(1, 0), (-1, 0), (0, 1), (0, -1)]
        ext_dirs = [(1, 0), (-1, 0), (0, 1), (0, -1)]
        ext_len = 2
        while len(pts) < n:
            for b, d in zip(base, ext_dirs):
                if len(pts) >= n:
                    break
                bx, by = b
                dx, dy = d
                pts.append((bx + dx * (ext_len - 1), by + dy * (ext_len - 1)))
            ext_len += 1
        return pts[:n]

    for i in range(n):
        pts.append((i, 0))
    return pts


# =========================
# Mapa de humedad (grid) + interpolación
# =========================
class HumidityMap:
    def __init__(self):
        self.nx = 60
        self.ny = 40
        self.seed = 1
        self.noise = 0.05

        self.peak_x = 1.20
        self.peak_y = 0.70
        self.peak_sigma = 0.18

        self.base_ax = 0.10
        self.base_ay = 0.05

        self.xs = []
        self.ys = []
        self.H = []
        self._generated = False

        self.bounds = (0.0, 1.90, 0.0, 1.20)

    def generate(self, nx, ny, seed, noise, peak_x, peak_y, peak_sigma, bounds):
        self.bounds = bounds
        X_MIN, X_MAX, Y_MIN, Y_MAX = bounds

        self.nx = max(10, min(int(nx), 500))
        self.ny = max(10, min(int(ny), 500))
        self.seed = int(seed)
        self.noise = max(0.0, min(float(noise), 1.0))

        self.peak_x = clamp_float(float(peak_x), X_MIN, X_MAX)
        self.peak_y = clamp_float(float(peak_y), Y_MIN, Y_MAX)
        self.peak_sigma = max(0.03, min(float(peak_sigma), 0.80))

        random.seed(self.seed)

        self.xs = [X_MIN + (X_MAX - X_MIN) * i / (self.nx - 1) for i in range(self.nx)]
        self.ys = [Y_MIN + (Y_MAX - Y_MIN) * j / (self.ny - 1) for j in range(self.ny)]

        H = []
        maxv = -1e9
        minv = 1e9

        for y in self.ys:
            row = []
            for x in self.xs:
                dx = (x - self.peak_x)
                dy = (y - self.peak_y)
                g = math.exp(-(dx * dx + dy * dy) / (2.0 * self.peak_sigma * self.peak_sigma))

                base = self.base_ax * (x - X_MIN) / (X_MAX - X_MIN) + self.base_ay * (y - Y_MIN) / (Y_MAX - Y_MIN)
                val = 0.15 + 0.80 * g + 0.20 * base
                val += self.noise * 0.08 * (2.0 * random.random() - 1.0)

                row.append(val)
                maxv = max(maxv, val)
                minv = min(minv, val)
            H.append(row)

        if maxv - minv < 1e-12:
            maxv = minv + 1e-12
        for j in range(self.ny):
            for i in range(self.nx):
                H[j][i] = (H[j][i] - minv) / (maxv - minv)

        self.H = H
        self._generated = True

    def is_ready(self):
        return self._generated and self.H and self.xs and self.ys

    def value_bilinear(self, x, y):
        if not self.is_ready():
            return 0.0

        X_MIN, X_MAX, Y_MIN, Y_MAX = self.bounds

        x = clamp_float(x, X_MIN, X_MAX)
        y = clamp_float(y, Y_MIN, Y_MAX)

        nx = self.nx
        ny = self.ny

        fx = (x - X_MIN) / (X_MAX - X_MIN) * (nx - 1)
        fy = (y - Y_MIN) / (Y_MAX - Y_MIN) * (ny - 1)
        i0 = int(math.floor(fx))
        j0 = int(math.floor(fy))
        i1 = min(i0 + 1, nx - 1)
        j1 = min(j0 + 1, ny - 1)
        i0 = max(0, min(i0, nx - 1))
        j0 = max(0, min(j0, ny - 1))

        tx = fx - i0
        ty = fy - j0

        h00 = self.H[j0][i0]
        h10 = self.H[j0][i1]
        h01 = self.H[j1][i0]
        h11 = self.H[j1][i1]

        hx0 = h00 * (1 - tx) + h10 * tx
        hx1 = h01 * (1 - tx) + h11 * tx
        hxy = hx0 * (1 - ty) + hx1 * ty
        return max(0.0, min(1.0, hxy))

    def argmax_grid(self):
        if not self.is_ready():
            return None
        best = (-1.0, None, None)
        for j, y in enumerate(self.ys):
            for i, x in enumerate(self.xs):
                v = self.H[j][i]
                if v > best[0]:
                    best = (v, x, y)
        return (best[1], best[2], best[0])


# ======================================================
# Ventana de gráficas (pestañas)
# ======================================================
class PlotWindow(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Gráficas")
        self.geometry("1180x780")

        self.nb = ttk.Notebook(self)
        self.nb.pack(fill="both", expand=True)

        self.tab_morph = ttk.Frame(self.nb)
        self.tab_map = ttk.Frame(self.nb)
        self.nb.add(self.tab_morph, text="Morfología")
        self.nb.add(self.tab_map, text="Humedad + Recorrido")

        self.fig_m = Figure(figsize=(8.5, 6.5), dpi=100)
        self.ax_m = self.fig_m.add_subplot(111)
        self.ax_m.set_aspect("equal", adjustable="box")
        self.ax_m.grid(True)
        self.canvas_m = FigureCanvasTkAgg(self.fig_m, master=self.tab_morph)
        self.canvas_m.get_tk_widget().pack(fill="both", expand=True)

        self.fig_h = Figure(figsize=(8.5, 6.5), dpi=100)
        self.ax_h = self.fig_h.add_subplot(111)
        self.ax_h.set_xlabel("X (m)")
        self.ax_h.set_ylabel("Y (m)")
        self.ax_h.grid(True)
        self.canvas_h = FigureCanvasTkAgg(self.fig_h, master=self.tab_map)
        self.canvas_h.get_tk_widget().pack(fill="both", expand=True)

        self.im = None
        self.plan_line, = self.ax_h.plot([], [], linestyle="--", alpha=0.6)
        self.traj_line, = self.ax_h.plot([], [], linewidth=2.0)
        self.robot_tip, = self.ax_h.plot([], [], linestyle="None", marker=(3, 0, 0), markersize=16)
        self.pt_goal, = self.ax_h.plot([], [], marker="x", markersize=10)
        self.pt_best, = self.ax_h.plot([], [], marker="*", markersize=14)
        self.pt_true, = self.ax_h.plot([], [], marker="o", markersize=6)
        self.txt = self.ax_h.text(0.02, 0.98, "", transform=self.ax_h.transAxes, ha="left", va="top", fontsize=10)

        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self._alive = True
        self.after(150, self._refresh)

    def _on_close(self):
        self._alive = False
        self.destroy()

    def _apply_fixed_axes(self):
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.app.get_bounds()

        try:
            tick = float(self.app.tick_var.get())
            if tick <= 0:
                tick = (X_MAX - X_MIN) / 10.0
        except Exception:
            tick = (X_MAX - X_MIN) / 10.0

        self.ax_h.set_xlim(X_MIN, X_MAX)
        self.ax_h.set_ylim(Y_MIN, Y_MAX)
        self.ax_h.xaxis.set_major_locator(MultipleLocator(tick))
        self.ax_h.yaxis.set_major_locator(MultipleLocator(tick))

    def _draw_morphology(self):
        self.ax_m.clear()
        self.ax_m.set_title("Morfología (H/V por módulo)")
        self.ax_m.grid(True)
        self.ax_m.set_aspect("equal", adjustable="box")

        ids = self.app.ids_runtime[:]
        pts = self.app.morph_positions[:]
        if not ids or not pts or len(ids) != len(pts):
            self.canvas_m.draw_idle()
            return

        xs = [p[0] for p in pts]
        ys = [p[1] for p in pts]
        xmin, xmax = min(xs) - 1, max(xs) + 1
        ymin, ymax = min(ys) - 1, max(ys) + 1
        self.ax_m.set_xlim(xmin, xmax)
        self.ax_m.set_ylim(ymin, ymax)

        for rid, (cx, cy) in zip(ids, pts):
            is_h = self.app.horiz_var_by_id.get(rid, tk.BooleanVar(value=False)).get()
            face = "#2b7bff" if is_h else "#ff9b2b"
            r = Rectangle((cx - 0.45, cy - 0.30), 0.90, 0.60, facecolor=face, edgecolor="black", linewidth=1.5)
            self.ax_m.add_patch(r)
            self.ax_m.text(cx, cy, f"ID {rid}\n{'H' if is_h else 'V'}", ha="center", va="center", fontsize=9, color="black")

        self.ax_m.text(0.02, 0.98, f"N módulos: {len(ids)}", transform=self.ax_m.transAxes, ha="left", va="top", fontsize=10)
        self.canvas_m.draw_idle()

    def _draw_humidity(self):
        hm = self.app.hmap
        if not hm.is_ready():
            return

        X_MIN, X_MAX, Y_MIN, Y_MAX = self.app.get_bounds()

        if self.im is None:
            self.im = self.ax_h.imshow(
                hm.H,
                origin="lower",
                extent=(X_MIN, X_MAX, Y_MIN, Y_MAX),
                aspect="auto",
                interpolation="nearest",
                alpha=0.85,
                cmap="jet",
                vmin=0.0, vmax=1.0
            )
            self.fig_h.colorbar(self.im, ax=self.ax_h, fraction=0.046, pad=0.04).set_label("Humedad mapa (0..1)")
        else:
            self.im.set_data(hm.H)
            self.im.set_extent((X_MIN, X_MAX, Y_MIN, Y_MAX))

        tg = hm.argmax_grid()
        if tg:
            xg, yg, _ = tg
            self.pt_true.set_data([xg], [yg])
        else:
            self.pt_true.set_data([], [])

        show_route = (self.app.show_route_var.get() == "Mostrar")
        show_goal = (self.app.show_goal_var.get() == "Mostrar")

        wps = self.app.waypoints_runtime[:]
        if show_route and wps:
            ax0, ay0 = self.app.start_A
            xs_plan = [ax0] + [p[0] for p in wps]
            ys_plan = [ay0] + [p[1] for p in wps]
            self.plan_line.set_data(xs_plan, ys_plan)
        else:
            self.plan_line.set_data([], [])

        if show_route:
            self.traj_line.set_data(self.app.traj_x[:], self.app.traj_y[:])
        else:
            self.traj_line.set_data([], [])

        x_r, y_r, th = self.app.pose_runtime
        ang_deg = math.degrees(th) - 90.0
        self.robot_tip.set_data([x_r], [y_r])
        self.robot_tip.set_marker((3, 0, ang_deg))

        gx, gy = self.app.current_goal
        if show_goal and gx is not None:
            self.pt_goal.set_data([gx], [gy])
        else:
            self.pt_goal.set_data([], [])

        bx, by, bv = self.app.best_found
        if bx is not None:
            self.pt_best.set_data([bx], [by])
        else:
            self.pt_best.set_data([], [])

        algo = self.app.algo_var.get()
        it = self.app.search_iter
        dist = self.app.distance_m
        vavg = self.app.robot_avg_speed_m_s()

        # Promedio de sensores (si existe)
        s_avg, s_n = self.app.sensors_avg_norm()
        if s_avg is None:
            s_txt = "N/A (sin sensores activos)"
        else:
            s_txt = f"{s_avg:.3f} (n={s_n})"

        if bx is not None:
            txt = (
                f"Algoritmo: {algo}\n"
                f"Iteración: {it}\n"
                f"Mejor h={bv:.3f} en ({bx:.2f},{by:.2f})\n"
                f"H_prom_sensores: {s_txt}\n"
                f"Distancia acumulada: {dist:.3f} m\n"
                f"Velocidad promedio robot: {vavg:.4f} m/s"
            )
        else:
            txt = (
                f"Algoritmo: {algo}\n"
                f"H_prom_sensores: {s_txt}\n"
                f"Distancia acumulada: {dist:.3f} m\n"
                f"Velocidad promedio robot: {vavg:.4f} m/s"
            )

        self.txt.set_text(txt)
        self._apply_fixed_axes()
        self.canvas_h.draw_idle()

    def _refresh(self):
        if not self._alive:
            return
        try:
            self._draw_morphology()
            self._draw_humidity()
        except Exception:
            pass
        finally:
            self.after(150, self._refresh)


# ======================================================
# Hilo: movimiento hacia un objetivo (waypoints)
# ======================================================
class NavigateThread(threading.Thread):
    def __init__(self, app, stop_event):
        super().__init__(daemon=True)
        self.app = app
        self.stop_event = stop_event
        self.last_t = time.perf_counter()

    def run(self):
        while not self.stop_event.is_set():
            try:
                if self.app.require_serial() and (not self.app.ser or not self.app.ser.is_open):
                    time.sleep(0.2)
                    continue

                freq = float(self.app.freq_var.get())
                rate = float(self.app.rate_var.get())
                if freq <= 0 or rate <= 0:
                    time.sleep(0.1)
                    continue

                tol = float(self.app.tol_var.get())
                dead = math.radians(float(self.app.deadband_deg_var.get()))

                now = time.perf_counter()
                dt = now - self.last_t
                if dt <= 0:
                    dt = 1.0 / max(1.0, rate)
                self.last_t = now

                kv = float(self.app.kv_var.get())
                kw = float(self.app.kw_var.get())

                if self.app.current_goal[0] is None:
                    g = self.app.algo_next_goal()
                    if g is None:
                        self.app.set_status("Búsqueda finalizada (sin nuevo objetivo).")
                        self.stop_event.set()
                        break
                    self.app.set_new_goal(gx=g[0], gy=g[1])
                    self.app.set_status("Objetivo generado por algoritmo.")
                    time.sleep(0.01)
                    continue

                wp = self.app.current_wp()
                if wp is None:
                    self.app.on_reach_goal()
                    continue

                tx_r, ty_r = wp
                x_r, y_r, th = self.app.pose_runtime

                dist = math.hypot(tx_r - x_r, ty_r - y_r)
                if dist <= tol:
                    self.app.advance_wp()
                    time.sleep(0.005)
                    continue

                desired = math.atan2(ty_r - y_r, tx_r - x_r)
                err = wrap_pi(desired - th)

                turning = abs(err) > dead
                cw = err < 0

                if turning:
                    Av = float(self.app.turn_v_ticks.get())
                    Ah = float(self.app.turn_h_ticks.get())
                    phi_v = float(self.app.turn_phase_v_deg.get())
                    phi_h = float(self.app.turn_phase_h_deg.get())
                else:
                    Av = float(self.app.move_v_ticks.get())
                    Ah = float(self.app.move_h_ticks.get())
                    phi_v = float(self.app.move_phase_v_deg.get())
                    phi_h = float(self.app.move_phase_h_deg.get())

                shift_enable = bool(self.app.turn_phase_shift_enable.get())
                shift_dir = self.app.turn_phase_shift_dir.get().strip()
                shift_deg = float(self.app.turn_phase_shift_deg.get())
                extra_phi_h = 0.0
                if turning and shift_enable:
                    if shift_dir == "CW" and cw:
                        extra_phi_h = shift_deg
                    elif shift_dir == "CCW" and (not cw):
                        extra_phi_h = shift_deg

                t_chain = time.perf_counter() - self.app.t0
                lines, cmd_positions = self.app.build_servo_lines_with_positions(
                    t=t_chain, freq=freq,
                    Av=Av, Ah=Ah,
                    phi_v_deg=phi_v,
                    phi_h_deg=phi_h + extra_phi_h,
                    turning=turning, cw=cw
                )

                self.app.update_speed_history(cmd_positions=cmd_positions, dt=dt)

                if lines and self.app.require_serial():
                    self.app.send_lines_batch(lines)

                self.app.update_pose_model(Av=Av, Ah=Ah, turning=turning, cw=cw, kv=kv, kw=kw, dt=dt)

                # Aquí se actualiza "last_humidity" con promedio de sensores si hay activos
                self.app.sample_humidity_at_pose()

                self.app.set_status("Navegación + búsqueda activa.")
                time.sleep(1.0 / rate)

            except Exception as e:
                self.app.set_status(f"Detenido por error: {e}")
                self.stop_event.set()

        self.app.network.stop()
        self.app.run_end_time = time.perf_counter()


# ======================================================
# Núcleo de algoritmos de búsqueda
# ======================================================
class SearchController:
    def __init__(self, app):
        self.app = app
        self.reset()

    def reset(self):
        self.algo = None
        self.iter = 0

        self.hc_step = 0.10

        self.ga_pop = []
        self.ga_pop_size = 18
        self.ga_elite = 4
        self.ga_mut_sigma = 0.12
        self.ga_gen = 0

        self.cg_k = 20
        self.cg_radius = 0.35

        self.bc_dir = None
        self.bc_run_len = 0
        self.bc_max_run = 10
        self.bc_step = 0.08
        self.bc_last_h = None

        self.sa_T = 1.0
        self.sa_alpha = 0.92
        self.sa_step = 0.18
        self.sa_current = None
        self.sa_current_h = None

        self.rw_step = 0.12
        self.rna = None
        self.rna_step = 0.08
        self.rna_turn_rad = math.radians(30.0)
        self.rna_last = None

    def set_algo(self, algo_name: str):
        self.reset()
        self.algo = algo_name
        self.iter = 0

        self.hc_step = float(self.app.hc_step_var.get())
        self.ga_pop_size = int(self.app.ga_pop_var.get())
        self.ga_elite = int(self.app.ga_elite_var.get())
        self.ga_mut_sigma = float(self.app.ga_mut_var.get())
        self.cg_k = int(self.app.cg_k_var.get())
        self.cg_radius = float(self.app.cg_radius_var.get())
        self.bc_step = float(self.app.bc_step_var.get())
        self.bc_max_run = int(self.app.bc_maxrun_var.get())
        self.sa_T = float(self.app.sa_T0_var.get())
        self.sa_alpha = float(self.app.sa_alpha_var.get())
        self.sa_step = float(self.app.sa_step_var.get())
        self.rw_step = float(self.app.rw_step_var.get())

        x0, y0, _ = self.app.pose_runtime
        if algo_name == "Algoritmo genético":
            self._ga_init()
        if algo_name == "Recocido simulado":
            self.sa_current = (x0, y0)
            self.sa_current_h = self.app.hmap.value_bilinear(x0, y0)
        if algo_name == "RNA de navegación":
            model_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                      "modelo_RNA_navegacion.json")
            self.rna = RNANavigator(model_path)
            self.rna.reset(self.app.last_humidity)

    def _bounds(self):
        return self.app.get_bounds()

    def _rand_point(self):
        X_MIN, X_MAX, Y_MIN, Y_MAX = self._bounds()
        return (random.uniform(X_MIN, X_MAX), random.uniform(Y_MIN, Y_MAX))

    def _rand_unit(self):
        a = random.uniform(-math.pi, math.pi)
        return (math.cos(a), math.sin(a))

    def _clip(self, x, y):
        X_MIN, X_MAX, Y_MIN, Y_MAX = self._bounds()
        return (clamp_float(x, X_MIN, X_MAX), clamp_float(y, Y_MIN, Y_MAX))

    def _hc_next(self):
        x, y, _ = self.app.pose_runtime
        step = max(0.005, self.hc_step)

        candidates = []
        for dx in (-step, 0.0, step):
            for dy in (-step, 0.0, step):
                if dx == 0.0 and dy == 0.0:
                    continue
                xx, yy = self._clip(x + dx, y + dy)
                hv = self.app.hmap.value_bilinear(xx, yy)
                candidates.append((hv, xx, yy))

        candidates.sort(reverse=True, key=lambda t: t[0])
        best = candidates[0]
        return (best[1], best[2])

    def _ga_init(self):
        self.ga_pop = []
        for _ in range(self.ga_pop_size):
            self.ga_pop.append(self._rand_point())
        self.ga_gen = 0

    def _ga_next(self):
        scored = []
        for (x, y) in self.ga_pop:
            xx, yy = self._clip(x, y)
            hv = self.app.hmap.value_bilinear(xx, yy)
            scored.append((hv, xx, yy))
        scored.sort(reverse=True, key=lambda t: t[0])

        best = scored[0]
        best_point = (best[1], best[2])

        elites = scored[:max(1, min(self.ga_elite, len(scored)))]
        new_pop = [(e[1], e[2]) for e in elites]

        def pick_parent():
            a = random.choice(scored)
            b = random.choice(scored)
            return a if a[0] > b[0] else b

        while len(new_pop) < self.ga_pop_size:
            p1 = pick_parent()
            p2 = pick_parent()
            cx = 0.5 * (p1[1] + p2[1])
            cy = 0.5 * (p1[2] + p2[2])
            cx += random.gauss(0.0, self.ga_mut_sigma)
            cy += random.gauss(0.0, self.ga_mut_sigma)
            cx, cy = self._clip(cx, cy)
            new_pop.append((cx, cy))

        self.ga_pop = new_pop
        self.ga_gen += 1
        return best_point

    def _cg_next(self):
        x, y, _ = self.app.pose_runtime
        k = max(5, self.cg_k)
        R = max(0.02, self.cg_radius)

        samples = []
        for _ in range(k):
            ang = random.uniform(-math.pi, math.pi)
            r = R * math.sqrt(random.random())
            xx = x + r * math.cos(ang)
            yy = y + r * math.sin(ang)
            xx, yy = self._clip(xx, yy)
            hv = self.app.hmap.value_bilinear(xx, yy)
            samples.append((hv, xx, yy))

        p = 2.0
        wsum = 1e-12
        xsum = 0.0
        ysum = 0.0
        for hv, xx, yy in samples:
            w = (hv ** p)
            wsum += w
            xsum += w * xx
            ysum += w * yy

        gx = xsum / wsum
        gy = ysum / wsum
        gx, gy = self._clip(gx, gy)
        return (gx, gy)

    def _bc_next(self):
        x, y, _ = self.app.pose_runtime
        step = max(0.005, self.bc_step)

        if self.bc_dir is None:
            self.bc_dir = self._rand_unit()
            self.bc_run_len = 0
            self.bc_last_h = self.app.hmap.value_bilinear(x, y)

        dx, dy = self.bc_dir
        nx, ny = self._clip(x + step * dx, y + step * dy)
        hnew = self.app.hmap.value_bilinear(nx, ny)

        if self.bc_last_h is None:
            self.bc_last_h = self.app.hmap.value_bilinear(x, y)

        if hnew >= self.bc_last_h:
            self.bc_last_h = hnew
            self.bc_run_len += 1
            if self.bc_run_len >= self.bc_max_run:
                self.bc_dir = self._rand_unit()
                self.bc_run_len = 0
            return (nx, ny)
        else:
            self.bc_dir = self._rand_unit()
            self.bc_run_len = 0
            dx, dy = self.bc_dir
            nx, ny = self._clip(x + step * dx, y + step * dy)
            self.bc_last_h = self.app.hmap.value_bilinear(nx, ny)
            return (nx, ny)

    def _sa_next(self):
        x, y, _ = self.app.pose_runtime

        if self.sa_current is None:
            self.sa_current = (x, y)
            self.sa_current_h = self.app.hmap.value_bilinear(x, y)

        T = max(1e-6, self.sa_T)
        step = max(0.005, self.sa_step)

        ang = random.uniform(-math.pi, math.pi)
        r = step * math.sqrt(random.random())
        nx, ny = self._clip(
            self.sa_current[0] + r * math.cos(ang),
            self.sa_current[1] + r * math.sin(ang)
        )
        hnew = self.app.hmap.value_bilinear(nx, ny)

        d = hnew - self.sa_current_h
        if d >= 0:
            accept = True
        else:
            p = math.exp(d / T)
            accept = (random.random() < p)

        if accept:
            self.sa_current = (nx, ny)
            self.sa_current_h = hnew

        self.sa_T = T * max(0.50, min(self.sa_alpha, 0.999))
        return self.sa_current

    def _rw_next(self):
        x, y, _ = self.app.pose_runtime
        step = max(0.005, float(self.rw_step))
        ang = random.uniform(-math.pi, math.pi)
        r = step * math.sqrt(random.random())
        nx, ny = self._clip(x + r * math.cos(ang), y + r * math.sin(ang))
        return (nx, ny)

    def _rna_next(self):
        x, y, th = self.app.pose_runtime
        humidity = self.app.get_current_humidity_value(x, y)
        prox = self.app.proximity_norm()
        action, probabilities, features = self.rna.predict(
            self.rna.features(humidity, prox, 0.0))
        if action == "GI":
            heading = th + self.rna_turn_rad
        elif action == "GD":
            heading = th - self.rna_turn_rad
        elif action == "EV":
            pf, pl, pr = prox
            heading = th + (self.rna_turn_rad if pr >= pl else -self.rna_turn_rad)
            if pf < 0.15:
                heading = th + math.pi
        else:
            heading = th
        nx, ny = self._clip(x + self.rna_step * math.cos(heading),
                            y + self.rna_step * math.sin(heading))
        self.rna_last = {"action": action, "probabilities": probabilities,
                         "features": features}
        return nx, ny

    def next_goal(self):
        self.iter += 1
        algo = self.algo
        if algo == "Ascenso a la colina":
            return self._hc_next()
        if algo == "Algoritmo genético":
            return self._ga_next()
        if algo == "Exploracion frontera":
            return self._cg_next()
        if algo == "RNA de navegación":
            return self._rna_next()
        if algo == "Recocido simulado":
            return self._sa_next()
        if algo == "Caminata aleatoria":
            return self._rw_next()
        return self._hc_next()


# ======================================================
# Aplicación principal
# ======================================================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Búsqueda en mapa de humedad + Morfologías + Sensores + Histórico (Excel)")
        self.geometry("1600x980")

        # Serial
        self.ser = None
        self.iolock = threading.Lock()

        # Thread
        self.stop_event = threading.Event()
        self.nav_thread = None

        # Reader thread (serial RX)
        self.read_stop = threading.Event()
        self.reader_thread = None

        # Plot
        self.plot_win = None

        # Humedad
        self.hmap = HumidityMap()

        # Runtime del robot
        self.ids_runtime = []
        self.morph_positions = []

        self.start_A = (0.0, 0.0)
        self.current_goal = (None, None)
        self.waypoints_runtime = []
        self.wp_index = 0

        self.pose_runtime = (0.0, 0.0, 0.0)  # (x,y,theta)
        self.traj_x = []
        self.traj_y = []
        self.t0 = time.perf_counter()

        self.best_found = (None, None, -1.0)  # (x,y,h)
        self.search_iter = 0

        self.distance_m = 0.0

        # Histórico
        self.fitness_history = []
        self.speed_acc = {}
        self.last_cmd_pos = {}
        self.speed_history = []
        self.run_config_snapshot = None
        self.run_id = None
        self.run_end_time = None
        self.network = NetworkMeasurement()

        self.humidity_series = []
        self.last_humidity = 0.0

        self.robot_series = []

        self._last_iter_dist = 0.0
        self._last_iter_time = 0.0

        self.search = SearchController(self)

        # =========================
        # SENSORES DE HUMEDAD (NUEVO)
        # =========================
        self.sensors_lock = threading.Lock()
        # sensors[sid] = {"raw": float, "norm": float, "t": float(perf_counter)}
        self.sensors = {}
        self.proximity = {}
        self.sensor_timeout_s = 5.0  # sensor "conectado" si hay lectura reciente

        # ========== Variables UI ==========
        self.port_var = tk.StringVar(value="COM7")
        self.baud_var = tk.StringVar(value="460800")
        self.eol_var = tk.StringVar(value="LF")
        self.status_var = tk.StringVar(value="Desconectado")

        self.nmods_var = tk.IntVar(value=5)
        self.ids_var = tk.StringVar(value="3,6,4,5,1")
        self.morph_var = tk.StringVar(value="Cuadrúpedo")

        # --- límites configurables ---
        self.xmin_var = tk.StringVar(value="0.0")
        self.xmax_var = tk.StringVar(value="1.90")
        self.ymin_var = tk.StringVar(value="0.0")
        self.ymax_var = tk.StringVar(value="1.20")

        self.ax_var = tk.StringVar(value="0.000")
        self.ay_var = tk.StringVar(value="1.000")
        self.align_var = tk.StringVar(value="+X")

        self.step_m_var = tk.StringVar(value="0.010")
        self.x_first_var = tk.BooleanVar(value=True)
        self.tol_var = tk.StringVar(value="0.10")
        self.deadband_deg_var = tk.StringVar(value="5")
        self.tick_var = tk.StringVar(value="0.10")

        self.freq_var = tk.StringVar(value="0.40")
        self.rate_var = tk.StringVar(value="20")

        self.move_v_ticks = tk.StringVar(value="210")
        self.move_h_ticks = tk.StringVar(value="10")
        self.move_phase_v_deg = tk.StringVar(value="0")
        self.move_phase_h_deg = tk.StringVar(value=str(PHI_H_DEG_DEFAULT))

        self.turn_v_ticks = tk.StringVar(value="180")
        self.turn_h_ticks = tk.StringVar(value="210")
        self.turn_phase_v_deg = tk.StringVar(value="0")
        self.turn_phase_h_deg = tk.StringVar(value=str(PHI_H_DEG_DEFAULT))

        self.turn_phase_shift_enable = tk.BooleanVar(value=False)
        self.turn_phase_shift_dir = tk.StringVar(value="CW")
        self.turn_phase_shift_deg = tk.StringVar(value="45")

        self.kv_var = tk.StringVar(value="0.0010")
        self.kw_var = tk.StringVar(value="0.01")

        self.send_serial_var = tk.BooleanVar(value=True)

        self.show_route_var = tk.StringVar(value="Mostrar")
        self.show_goal_var = tk.StringVar(value="Mostrar")

        # H/V por ID
        self.horiz_var_by_id = {}

        # -------- Mapa humedad UI --------
        self.hm_nx_var = tk.StringVar(value="60")
        self.hm_ny_var = tk.StringVar(value="40")
        self.hm_seed_var = tk.StringVar(value="1")
        self.hm_noise_var = tk.StringVar(value="0.05")
        self.hm_peakx_var = tk.StringVar(value="1.20")
        self.hm_peaky_var = tk.StringVar(value="0.70")
        self.hm_sigma_var = tk.StringVar(value="0.18")

        # -------- Algoritmos UI --------
        self.algo_var = tk.StringVar(value="Ascenso a la colina")
        self.repetition_var = tk.IntVar(value=1)
        self.max_iters_var = tk.StringVar(value="120")
        self.stop_h_thresh_var = tk.StringVar(value="0.995")
        self.force_max_iters_var = tk.BooleanVar(value=True)

        self.hc_step_var = tk.StringVar(value="0.20")
        self.ga_pop_var = tk.StringVar(value="18")
        self.ga_elite_var = tk.StringVar(value="4")
        self.ga_mut_var = tk.StringVar(value="0.12")
        self.cg_k_var = tk.StringVar(value="20")
        self.cg_radius_var = tk.StringVar(value="0.35")
        self.bc_step_var = tk.StringVar(value="0.08")
        self.bc_maxrun_var = tk.StringVar(value="10")
        self.sa_T0_var = tk.StringVar(value="1.0")
        self.sa_alpha_var = tk.StringVar(value="0.92")
        self.sa_step_var = tk.StringVar(value="0.18")
        self.rw_step_var = tk.StringVar(value="0.12")

        # -------- Sensores: escalamiento raw->norm (NUEVO) --------
        # norm = (raw - raw_min)/(raw_max-raw_min) y opcional invertir (1-norm)
        self.hraw_min_var = tk.StringVar(value="0")
        self.hraw_max_var = tk.StringVar(value="1023")
        self.hinv_var = tk.BooleanVar(value=False)
        self.s_timeout_var = tk.StringVar(value="5.0")
        self.s_avg_var = tk.StringVar(value="N/A")
        self.s_n_var = tk.StringVar(value="0")

        # construir UI
        self._build_ui()
        self.protocol("WM_DELETE_WINDOW", self.on_close)

        self._rebuild_hv_checks()
        self.update_morphology_runtime()

        # Inicializar con límites válidos + mapa
        self.apply_bounds_and_reinit()
        self.generate_humidity_map()

        # refresco tabla sensores
        self.after(250, self._refresh_sensors_table)

    # ================== Límites ==================
    def get_bounds(self):
        try:
            xmin = float(self.xmin_var.get())
            xmax = float(self.xmax_var.get())
            ymin = float(self.ymin_var.get())
            ymax = float(self.ymax_var.get())
        except Exception:
            xmin, xmax, ymin, ymax = 0.0, 1.90, 0.0, 1.20

        if xmax <= xmin:
            xmax = xmin + 1e-6
        if ymax <= ymin:
            ymax = ymin + 1e-6

        return (xmin, xmax, ymin, ymax)

    def clip_xy(self, x, y):
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.get_bounds()
        return (clamp_float(x, X_MIN, X_MAX), clamp_float(y, Y_MIN, Y_MAX))

    def apply_bounds_and_reinit(self):
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.get_bounds()

        try:
            ax = float(self.ax_var.get())
            ay = float(self.ay_var.get())
        except Exception:
            ax, ay = X_MIN, Y_MIN
        ax, ay = self.clip_xy(ax, ay)
        self.ax_var.set(f"{ax:.3f}")
        self.ay_var.set(f"{ay:.3f}")

        try:
            px = float(self.hm_peakx_var.get())
            py = float(self.hm_peaky_var.get())
        except Exception:
            px, py = (X_MIN + X_MAX) / 2.0, (Y_MIN + Y_MAX) / 2.0
        px, py = self.clip_xy(px, py)
        self.hm_peakx_var.set(f"{px:.3f}")
        self.hm_peaky_var.set(f"{py:.3f}")

        x, y, th = self.pose_runtime
        x, y = self.clip_xy(x, y)
        self.pose_runtime = (x, y, th)
        if not self.traj_x:
            self.traj_x = [x]
            self.traj_y = [y]
        else:
            self.traj_x[-1] = x
            self.traj_y[-1] = y

        gx, gy = self.current_goal
        if gx is not None:
            gx, gy = self.clip_xy(gx, gy)
            self.current_goal = (gx, gy)

        bx, by, bv = self.best_found
        if bx is not None:
            bx, by = self.clip_xy(bx, by)
            self.best_found = (bx, by, bv)

    # ================= UI =================
    def _build_ui(self):
        root = ttk.Frame(self, padding=10)
        root.pack(fill="both", expand=True)
        for i in range(3):
            root.columnconfigure(i, weight=1)

        col1 = ttk.Frame(root)
        col2 = ttk.Frame(root)
        col3 = ttk.Frame(root)
        col1.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        col2.grid(row=0, column=1, sticky="nsew", padx=(8, 8))
        col3.grid(row=0, column=2, sticky="nsew", padx=(8, 0))

        self._ui_serial(col1)
        self._ui_modules(col1)
        self._ui_morphology(col1)
        self._ui_humidity(col1)

        self._ui_navigation(col2)
        self._ui_motion(col2)
        self._ui_algo(col2)

        self._ui_exec(col3)
        self._ui_status(col3)
        self._ui_sensors(col3)   # NUEVO
        self._ui_history(col3)

    def _ui_serial(self, parent):
        f = ttk.LabelFrame(parent, text="Puerto Serial", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="COM").grid(row=0, column=0, sticky="w")
        self.port_cb = ttk.Combobox(f, textvariable=self.port_var, values=list_serial_ports(), width=16)
        self.port_cb.grid(row=0, column=1, padx=6, sticky="w")
        ttk.Button(f, text="Refrescar", command=self._refresh_ports).grid(row=0, column=2, padx=6)

        ttk.Label(f, text="Baudios").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.baud_var, width=10).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))

        ttk.Label(f, text="EOL").grid(row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Combobox(f, textvariable=self.eol_var, values=["LF", "CRLF"], width=6, state="readonly")\
            .grid(row=1, column=3, padx=6, sticky="w", pady=(6, 0))

        ttk.Button(f, text="Conectar", command=self.connect).grid(row=2, column=1, padx=6, pady=(8, 0), sticky="w")
        ttk.Button(f, text="Desconectar", command=self.disconnect).grid(row=2, column=2, padx=6, pady=(8, 0), sticky="w")

        ttk.Checkbutton(f, text="Enviar comandos al robot (Serial)", variable=self.send_serial_var)\
            .grid(row=3, column=0, columnspan=4, sticky="w", pady=(10, 0))

    def _ui_modules(self, parent):
        f = ttk.LabelFrame(parent, text="Módulos", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="Cantidad").grid(row=0, column=0, sticky="w")
        ttk.Combobox(f, textvariable=self.nmods_var, values=list(range(1, 25)), width=6, state="readonly")\
            .grid(row=0, column=1, padx=6, sticky="w")
        ttk.Button(f, text="Aplicar", command=self._apply_nmods).grid(row=0, column=2, padx=6, sticky="w")

        ttk.Label(f, text="IDs (coma)").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.ids_var, width=26).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0), columnspan=2)

        self.hv_frame = ttk.Frame(f)
        self.hv_frame.grid(row=2, column=0, columnspan=3, sticky="ew", pady=(10, 0))

        ttk.Label(f, text="Nota: H/V define qué onda aplica a cada ID.").grid(
            row=3, column=0, columnspan=3, sticky="w", pady=(8, 0)
        )

    def _ui_morphology(self, parent):
        f = ttk.LabelFrame(parent, text="Morfología", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="Tipo").grid(row=0, column=0, sticky="w")
        ttk.Combobox(f, textvariable=self.morph_var,
                     values=["Serpiente", "L", "T", "Cuadrúpedo"], width=12, state="readonly")\
            .grid(row=0, column=1, padx=6, sticky="w")
        ttk.Button(f, text="Actualizar dibujo", command=self.update_morphology_runtime)\
            .grid(row=0, column=2, padx=6, sticky="w")

    def _ui_humidity(self, parent):
        f = ttk.LabelFrame(parent, text="Mapa de humedad", padding=10)
        f.pack(fill="x", pady=(0, 10))

        lim = ttk.LabelFrame(f, text="Límites del espacio", padding=8)
        lim.grid(row=0, column=0, columnspan=4, sticky="ew")

        ttk.Label(lim, text="Xmin").grid(row=0, column=0, sticky="w")
        ttk.Entry(lim, textvariable=self.xmin_var, width=8).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(lim, text="Xmax").grid(row=0, column=2, sticky="w")
        ttk.Entry(lim, textvariable=self.xmax_var, width=8).grid(row=0, column=3, padx=6, sticky="w")

        ttk.Label(lim, text="Ymin").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(lim, textvariable=self.ymin_var, width=8).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Label(lim, text="Ymax").grid(row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(lim, textvariable=self.ymax_var, width=8).grid(row=1, column=3, padx=6, sticky="w", pady=(6, 0))

        ttk.Button(lim, text="Aplicar límites", command=self.on_apply_limits)\
            .grid(row=2, column=0, columnspan=4, sticky="w", pady=(8, 0))

        ttk.Label(f, text="Nx").grid(row=1, column=0, sticky="w", pady=(10, 0))
        ttk.Entry(f, textvariable=self.hm_nx_var, width=6).grid(row=1, column=1, padx=6, sticky="w", pady=(10, 0))
        ttk.Label(f, text="Ny").grid(row=1, column=2, sticky="w", pady=(10, 0))
        ttk.Entry(f, textvariable=self.hm_ny_var, width=6).grid(row=1, column=3, padx=6, sticky="w", pady=(10, 0))

        ttk.Label(f, text="Seed").grid(row=2, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.hm_seed_var, width=6).grid(row=2, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Label(f, text="Ruido (0..1)").grid(row=2, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.hm_noise_var, width=6).grid(row=2, column=3, padx=6, sticky="w", pady=(6, 0))

        ttk.Label(f, text="Pico x").grid(row=3, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.hm_peakx_var, width=8).grid(row=3, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Label(f, text="Pico y").grid(row=3, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.hm_peaky_var, width=8).grid(row=3, column=3, padx=6, sticky="w", pady=(6, 0))

        ttk.Label(f, text="Sigma").grid(row=4, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.hm_sigma_var, width=8).grid(row=4, column=1, padx=6, sticky="w", pady=(6, 0))

        ttk.Button(f, text="Generar mapa", command=self.generate_humidity_map)\
            .grid(row=5, column=0, columnspan=4, sticky="w", pady=(10, 0))

    def _ui_navigation(self, parent):
        f = ttk.LabelFrame(parent, text="Navegación (pose inicial)", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="A.x").grid(row=0, column=0, sticky="w")
        ttk.Entry(f, textvariable=self.ax_var, width=10).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(f, text="A.y").grid(row=0, column=2, sticky="w")
        ttk.Entry(f, textvariable=self.ay_var, width=10).grid(row=0, column=3, padx=6, sticky="w")

        ttk.Label(f, text="Alineación").grid(row=0, column=4, sticky="w")
        ttk.Combobox(f, textvariable=self.align_var, values=["+X", "+Y", "-X", "-Y"], width=6, state="readonly")\
            .grid(row=0, column=5, padx=6, sticky="w")

        ttk.Label(f, text="Paso (m)").grid(row=1, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(f, textvariable=self.step_m_var, width=10).grid(row=1, column=1, padx=6, sticky="w", pady=(8, 0))
        ttk.Checkbutton(f, text="Primero X luego Y", variable=self.x_first_var)\
            .grid(row=1, column=2, columnspan=2, sticky="w", pady=(8, 0))

        ttk.Label(f, text="Tol (m)").grid(row=2, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(f, textvariable=self.tol_var, width=10).grid(row=2, column=1, padx=6, sticky="w", pady=(8, 0))
        ttk.Label(f, text="Deadband (°)").grid(row=2, column=2, sticky="w", pady=(8, 0))
        ttk.Entry(f, textvariable=self.deadband_deg_var, width=10).grid(row=2, column=3, padx=6, sticky="w", pady=(8, 0))

        ttk.Label(f, text="Tick ejes (m)").grid(row=3, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(f, textvariable=self.tick_var, width=10).grid(row=3, column=1, padx=6, sticky="w", pady=(8, 0))

        ttk.Label(f, text="Ejes se ajustan a límites configurados.").grid(
            row=3, column=2, columnspan=4, sticky="w", pady=(8, 0)
        )

    def _ui_motion(self, parent):
        f = ttk.LabelFrame(parent, text="Movimiento", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="Frecuencia (Hz)").grid(row=0, column=0, sticky="w")
        ttk.Entry(f, textvariable=self.freq_var, width=10).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(f, text="Tasa envío (Hz)").grid(row=0, column=2, sticky="w")
        ttk.Entry(f, textvariable=self.rate_var, width=10).grid(row=0, column=3, padx=6, sticky="w")

        f1 = ttk.LabelFrame(f, text="Avance (recto)", padding=8)
        f1.grid(row=1, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Label(f1, text="V(A ticks)").grid(row=0, column=0, sticky="w")
        ttk.Entry(f1, textvariable=self.move_v_ticks, width=10).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(f1, text="H(A ticks)").grid(row=0, column=2, sticky="w")
        ttk.Entry(f1, textvariable=self.move_h_ticks, width=10).grid(row=0, column=3, padx=6, sticky="w")

        ttk.Label(f1, text="V(φ°)").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f1, textvariable=self.move_phase_v_deg, width=10).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Label(f1, text="H(φ°)").grid(row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(f1, textvariable=self.move_phase_h_deg, width=10).grid(row=1, column=3, padx=6, sticky="w", pady=(6, 0))

        f2 = ttk.LabelFrame(f, text="Giro", padding=8)
        f2.grid(row=2, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Label(f2, text="V(A ticks)").grid(row=0, column=0, sticky="w")
        ttk.Entry(f2, textvariable=self.turn_v_ticks, width=10).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(f2, text="H(A ticks)").grid(row=0, column=2, sticky="w")
        ttk.Entry(f2, textvariable=self.turn_h_ticks, width=10).grid(row=0, column=3, padx=6, sticky="w")

        ttk.Label(f2, text="V(φ°)").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(f2, textvariable=self.turn_phase_v_deg, width=10).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Label(f2, text="H(φ°)").grid(row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(f2, textvariable=self.turn_phase_h_deg, width=10).grid(row=1, column=3, padx=6, sticky="w", pady=(6, 0))

        f3 = ttk.LabelFrame(f, text="Giro: cambio de fase SOLO en un sentido", padding=8)
        f3.grid(row=3, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Checkbutton(f3, text="Habilitar", variable=self.turn_phase_shift_enable).grid(row=0, column=0, sticky="w")
        ttk.Label(f3, text="Sentido").grid(row=0, column=1, sticky="w")
        ttk.Combobox(f3, textvariable=self.turn_phase_shift_dir, values=["CW", "CCW"], width=6, state="readonly")\
            .grid(row=0, column=2, padx=6, sticky="w")
        ttk.Label(f3, text="Δφ (°) aplicado a H").grid(row=0, column=3, sticky="w")
        ttk.Entry(f3, textvariable=self.turn_phase_shift_deg, width=10).grid(row=0, column=4, padx=6, sticky="w")

        f4 = ttk.LabelFrame(f, text="Modelo (pose) para la gráfica", padding=8)
        f4.grid(row=4, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Label(f4, text="k_v").grid(row=0, column=0, sticky="w")
        ttk.Entry(f4, textvariable=self.kv_var, width=10).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(f4, text="k_w").grid(row=0, column=2, sticky="w")
        ttk.Entry(f4, textvariable=self.kw_var, width=10).grid(row=0, column=3, padx=6, sticky="w")

    def _ui_algo(self, parent):
        f = ttk.LabelFrame(parent, text="Búsqueda (selección de algoritmo)", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Label(f, text="Algoritmo").grid(row=0, column=0, sticky="w")
        ttk.Combobox(
            f,
            textvariable=self.algo_var,
            values=[
                "Ascenso a la colina",
                "Algoritmo genético",
                "Exploracion frontera",
                "RNA de navegación",
                "Recocido simulado",
                "Caminata aleatoria"
            ],
            width=24,
            state="readonly"
        ).grid(row=0, column=1, padx=6, sticky="w")

        ttk.Label(f, text="Máx iteraciones").grid(row=0, column=2, sticky="w")
        ttk.Entry(f, textvariable=self.max_iters_var, width=8).grid(row=0, column=3, padx=6, sticky="w")

        ttk.Label(f, text="Parar si h ≥").grid(row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Entry(f, textvariable=self.stop_h_thresh_var, width=8).grid(row=1, column=3, padx=6, sticky="w", pady=(6, 0))
        ttk.Checkbutton(
            f,
            text="Forzar máximo de iteraciones (ignorar convergencia temprana)",
            variable=self.force_max_iters_var
        ).grid(row=2, column=0, columnspan=4, padx=6, pady=(8, 0), sticky="w")

        box = ttk.Frame(f)
        box.grid(row=2, column=0, columnspan=4, sticky="ew", pady=(10, 0))

        hf = ttk.LabelFrame(box, text="Hill", padding=8)
        hf.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        ttk.Label(hf, text="step").grid(row=0, column=0, sticky="w")
        ttk.Entry(hf, textvariable=self.hc_step_var, width=8).grid(row=0, column=1, padx=6, sticky="w")

        gf = ttk.LabelFrame(box, text="GA", padding=8)
        gf.grid(row=0, column=1, sticky="ew", padx=(0, 8))
        ttk.Label(gf, text="pop").grid(row=0, column=0, sticky="w")
        ttk.Entry(gf, textvariable=self.ga_pop_var, width=6).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(gf, text="elite").grid(row=0, column=2, sticky="w")
        ttk.Entry(gf, textvariable=self.ga_elite_var, width=6).grid(row=0, column=3, padx=6, sticky="w")
        ttk.Label(gf, text="mutσ").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(gf, textvariable=self.ga_mut_var, width=6).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))

        cf = ttk.LabelFrame(box, text="Centroide", padding=8)
        cf.grid(row=0, column=2, sticky="ew", padx=(0, 8))
        ttk.Label(cf, text="k").grid(row=0, column=0, sticky="w")
        ttk.Entry(cf, textvariable=self.cg_k_var, width=6).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(cf, text="R").grid(row=0, column=2, sticky="w")
        ttk.Entry(cf, textvariable=self.cg_radius_var, width=6).grid(row=0, column=3, padx=6, sticky="w")

        bf = ttk.LabelFrame(box, text="Bacteria", padding=8)
        bf.grid(row=1, column=0, sticky="ew", padx=(0, 8), pady=(8, 0))
        ttk.Label(bf, text="step").grid(row=0, column=0, sticky="w")
        ttk.Entry(bf, textvariable=self.bc_step_var, width=8).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(bf, text="maxrun").grid(row=0, column=2, sticky="w")
        ttk.Entry(bf, textvariable=self.bc_maxrun_var, width=6).grid(row=0, column=3, padx=6, sticky="w")

        sf = ttk.LabelFrame(box, text="SA", padding=8)
        sf.grid(row=1, column=1, sticky="ew", padx=(0, 8), pady=(8, 0))
        ttk.Label(sf, text="T0").grid(row=0, column=0, sticky="w")
        ttk.Entry(sf, textvariable=self.sa_T0_var, width=6).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(sf, text="α").grid(row=0, column=2, sticky="w")
        ttk.Entry(sf, textvariable=self.sa_alpha_var, width=6).grid(row=0, column=3, padx=6, sticky="w")
        ttk.Label(sf, text="step").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(sf, textvariable=self.sa_step_var, width=6).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))

        rf = ttk.LabelFrame(box, text="Random walk", padding=8)
        rf.grid(row=1, column=2, sticky="ew", padx=(0, 8), pady=(8, 0))
        ttk.Label(rf, text="step").grid(row=0, column=0, sticky="w")
        ttk.Entry(rf, textvariable=self.rw_step_var, width=8).grid(row=0, column=1, padx=6, sticky="w")

    def _ui_exec(self, parent):
        f = ttk.LabelFrame(parent, text="Ejecución", padding=10)
        f.pack(fill="x", pady=(0, 10))

        ttk.Button(f, text="Abrir gráficas", command=self.open_plot_window).pack(side="left", padx=6)
        ttk.Button(f, text="Actualizar Morfología", command=self.update_morphology_runtime).pack(side="left", padx=6)
        ttk.Button(f, text="HOME", command=self.home_all).pack(side="left", padx=6)

        ttk.Button(f, text="Iniciar búsqueda", command=self.start_search).pack(side="left", padx=6)
        ttk.Button(f, text="Detener", command=self.stop).pack(side="left", padx=6)
        ttk.Label(f, text="Repetición:").pack(side="left", padx=(16, 4))
        ttk.Spinbox(f, from_=1, to=100, textvariable=self.repetition_var, width=5).pack(side="left")

        f2 = ttk.Frame(parent)
        f2.pack(fill="x", pady=(0, 10))
        box = ttk.LabelFrame(f2, text="Vista (gráfica)", padding=10)
        box.pack(fill="x")

        ttk.Label(box, text="Ruta").grid(row=0, column=0, sticky="w")
        ttk.Combobox(box, textvariable=self.show_route_var, values=["Mostrar", "Ocultar"], width=10, state="readonly")\
            .grid(row=0, column=1, padx=6, sticky="w")

        ttk.Label(box, text="Punto final").grid(row=0, column=2, sticky="w")
        ttk.Combobox(box, textvariable=self.show_goal_var, values=["Mostrar", "Ocultar"], width=10, state="readonly")\
            .grid(row=0, column=3, padx=6, sticky="w")

    def _ui_status(self, parent):
        f = ttk.LabelFrame(parent, text="Estado", padding=10)
        f.pack(fill="x")
        ttk.Label(f, textvariable=self.status_var, wraplength=520, justify="left").pack(anchor="w")

    # ================== SENSORES UI (NUEVO) ==================
    def _ui_sensors(self, parent):
        f = ttk.LabelFrame(parent, text="Sensores de humedad (Serial)", padding=10)
        f.pack(fill="both", pady=(10, 10), expand=False)

        top = ttk.Frame(f)
        top.pack(fill="x")

        ttk.Label(top, text="RAW min").grid(row=0, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.hraw_min_var, width=8).grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(top, text="RAW max").grid(row=0, column=2, sticky="w")
        ttk.Entry(top, textvariable=self.hraw_max_var, width=8).grid(row=0, column=3, padx=6, sticky="w")
        ttk.Checkbutton(top, text="Invertir", variable=self.hinv_var).grid(row=0, column=4, padx=6, sticky="w")

        ttk.Label(top, text="timeout (s)").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(top, textvariable=self.s_timeout_var, width=8).grid(row=1, column=1, padx=6, sticky="w", pady=(6, 0))
        ttk.Button(top, text="Aplicar", command=self._apply_sensor_params).grid(row=1, column=2, padx=6, sticky="w", pady=(6, 0))

        mid = ttk.Frame(f)
        mid.pack(fill="x", pady=(8, 6))

        ttk.Label(mid, text="Promedio (0..1):").pack(side="left")
        ttk.Label(mid, textvariable=self.s_avg_var, width=10).pack(side="left", padx=(6, 12))
        ttk.Label(mid, text="Sensores activos:").pack(side="left")
        ttk.Label(mid, textvariable=self.s_n_var, width=6).pack(side="left", padx=(6, 0))

        # Tabla
        cols = ("sid", "raw", "norm", "edad_s")
        self.sensor_tv = ttk.Treeview(f, columns=cols, show="headings", height=8)
        self.sensor_tv.heading("sid", text="Sensor ID")
        self.sensor_tv.heading("raw", text="H raw")
        self.sensor_tv.heading("norm", text="H (0..1)")
        self.sensor_tv.heading("edad_s", text="Edad (s)")

        self.sensor_tv.column("sid", width=80, anchor="center")
        self.sensor_tv.column("raw", width=110, anchor="center")
        self.sensor_tv.column("norm", width=110, anchor="center")
        self.sensor_tv.column("edad_s", width=100, anchor="center")

        self.sensor_tv.pack(fill="both", expand=True)

        ttk.Button(f, text="Limpiar lecturas", command=self.clear_sensors).pack(anchor="w", pady=(6, 0))

    def _apply_sensor_params(self):
        try:
            self.sensor_timeout_s = max(0.2, float(self.s_timeout_var.get()))
        except Exception:
            self.sensor_timeout_s = 5.0
            self.s_timeout_var.set("5.0")

    def _ui_history(self, parent):
        f = ttk.LabelFrame(parent, text="Exportación", padding=10)
        f.pack(fill="x", pady=(0, 0))

        ttk.Button(f, text="Exportar a Excel (XLSX + JSON)", command=self.export_history).pack(side="left", padx=6)
        ttk.Button(f, text="Limpiar histórico", command=self.clear_history).pack(side="left", padx=6)

    # ================== Eventos límites ==================
    def on_apply_limits(self):
        self.apply_bounds_and_reinit()
        self.generate_humidity_map()
        self.set_status("Límites aplicados. Mapa regenerado y coordenadas validadas.")

    # ================== H/V dinámico ==================
    def _apply_nmods(self):
        self._rebuild_hv_checks()
        self.update_morphology_runtime()

    def _rebuild_hv_checks(self):
        for w in self.hv_frame.winfo_children():
            w.destroy()

        n = int(self.nmods_var.get())
        try:
            ids = parse_id_list(self.ids_var.get())
        except Exception:
            ids = []

        if len(ids) != n:
            if len(ids) > n:
                ids = ids[:n]
            else:
                used = set(ids)
                nxt = 1
                while len(ids) < n:
                    if nxt not in used:
                        ids.append(nxt)
                        used.add(nxt)
                    nxt += 1
            self.ids_var.set(",".join(str(i) for i in ids))

        self.horiz_var_by_id = {}
        for rid in ids:
            self.horiz_var_by_id[rid] = tk.BooleanVar(value=False)

        ttk.Label(self.hv_frame, text="Marcar IDs horizontales (H):").grid(row=0, column=0, columnspan=6, sticky="w")
        for i, rid in enumerate(ids):
            r = 1 + (i // 6)
            c = i % 6
            ttk.Checkbutton(self.hv_frame, text=f"ID {rid}", variable=self.horiz_var_by_id[rid])\
                .grid(row=r, column=c, padx=(0, 10), pady=(6, 0), sticky="w")

    def horizontal_ids_set(self):
        return {rid for rid, var in self.horiz_var_by_id.items() if var.get()}

    # ================== Morfología runtime ==================
    def update_morphology_runtime(self):
        try:
            ids = parse_id_list(self.ids_var.get())
        except Exception:
            ids = []
        if not ids:
            self.set_status("IDs vacíos.")
            return
        self.ids_runtime = ids[:]
        self.morph_positions = morphology_positions(self.morph_var.get().strip(), len(ids))
        if self.plot_win and self.plot_win.winfo_exists():
            self.plot_win.lift()

    # ================= Humedad (mapa) =================
    def generate_humidity_map(self):
        try:
            nx = int(self.hm_nx_var.get())
            ny = int(self.hm_ny_var.get())
            seed = int(self.hm_seed_var.get())
            noise = float(self.hm_noise_var.get())
            px = float(self.hm_peakx_var.get())
            py = float(self.hm_peaky_var.get())
            sg = float(self.hm_sigma_var.get())
        except Exception as e:
            messagebox.showerror("Humedad", f"Parámetros inválidos: {e}")
            return

        bounds = self.get_bounds()
        px, py = self.clip_xy(px, py)
        self.hm_peakx_var.set(f"{px:.3f}")
        self.hm_peaky_var.set(f"{py:.3f}")

        self.hmap.generate(nx=nx, ny=ny, seed=seed, noise=noise, peak_x=px, peak_y=py, peak_sigma=sg, bounds=bounds)
        mx = self.hmap.argmax_grid()
        if mx:
            xg, yg, vg = mx
            self.set_status(f"Mapa generado. Máximo grid h={vg:.3f} en ({xg:.2f},{yg:.2f}).")
        else:
            self.set_status("Mapa generado.")
        if self.plot_win and self.plot_win.winfo_exists():
            self.plot_win.lift()

    # ================= Plot =================
    def open_plot_window(self):
        if self.plot_win and self.plot_win.winfo_exists():
            self.plot_win.lift()
            return
        self.plot_win = PlotWindow(self)

    # ================= Serial =================
    def _refresh_ports(self):
        self.port_cb["values"] = list_serial_ports()

    def _eol_bytes(self) -> bytes:
        return b"\r\n" if self.eol_var.get() == "CRLF" else b"\n"

    def send_lines_batch(self, lines):
        if not self.ser or not self.ser.is_open:
            raise RuntimeError("Maestro no conectado")
        eol = self._eol_bytes()
        payload = b"".join([ln.strip().encode("utf-8") + eol for ln in lines])
        with self.iolock:
            self.ser.write(payload)

    def connect(self):
        port = self.port_var.get().strip()
        if not port:
            messagebox.showwarning("Serial", "Seleccione el COM.")
            return
        try:
            baud = int(self.baud_var.get().strip())
        except ValueError:
            messagebox.showwarning("Serial", "Baudios inválidos.")
            return

        try:
            self.ser = serial.Serial(port, baud, timeout=0.2)
            time.sleep(1.2)
        except Exception as e:
            messagebox.showerror("Serial", f"No fue posible abrir el puerto: {e}")
            self.ser = None
            return

        self.status_var.set(f"Conectado a {port} @ {baud} | EOL={self.eol_var.get()}")

        # Arrancar lector serial (NUEVO)
        self.start_serial_reader()

    def disconnect(self):
        self.stop()

        # Detener lector serial (NUEVO)
        self.stop_serial_reader()

        try:
            if self.ser and self.ser.is_open:
                with self.iolock:
                    self.ser.close()
        except Exception:
            pass
        self.ser = None
        self.status_var.set("Desconectado")

    def set_status(self, s: str):
        self.status_var.set(s)

    def require_serial(self):
        return bool(self.send_serial_var.get())

    # ================= Lector serial (NUEVO) =================
    def start_serial_reader(self):
        self.stop_serial_reader()
        if not self.ser or not self.ser.is_open:
            return
        self.read_stop.clear()
        self.reader_thread = threading.Thread(target=self._serial_read_loop, daemon=True)
        self.reader_thread.start()

    def stop_serial_reader(self):
        self.read_stop.set()
        th = self.reader_thread
        self.reader_thread = None
        if th and th.is_alive():
            try:
                th.join(timeout=0.6)
            except Exception:
                pass

    def _serial_read_loop(self):
        """
        Lee líneas de la UART y extrae humedad por sensor.

        Formatos aceptados:
          - "sid,valor"            ej: "7,627"
          - "H,sid,valor"          ej: "H,7,627"
          - "HUM,sid,valor"        ej: "HUM,7,627"
        """
        while (not self.read_stop.is_set()) and self.ser and self.ser.is_open:
            try:
                with self.iolock:
                    raw = self.ser.readline()
                if not raw:
                    continue
                try:
                    line = raw.decode("utf-8", errors="ignore").strip()
                except Exception:
                    continue
                if not line:
                    continue
                self._handle_rx_line(line)
            except Exception:
                time.sleep(0.05)

    def _handle_rx_line(self, line: str):
        upper = line.upper()
        if upper.startswith("DATA,") or upper.startswith("ACK,"):
            self.network.ingest(line)
        prox = self.parse_proximity_line(line)
        if prox is not None:
            sid, front, left, right = prox
            with self.sensors_lock:
                self.proximity[sid] = {"f": front, "l": left, "r": right,
                                       "t": time.perf_counter()}
        parsed = self.parse_humidity_line(line)
        if parsed is None:
            return
        sid, hraw = parsed
        hnorm = self.humidity_raw_to_norm(hraw)
        now = time.perf_counter()
        with self.sensors_lock:
            self.sensors[int(sid)] = {"raw": float(hraw), "norm": float(hnorm), "t": now}

    def parse_proximity_line(self, line: str):
        parts = [p.strip() for p in line.replace(";", ",").split(",")]
        if len(parts) != 4 or not all(parts):
            return None
        try:
            return int(float(parts[0])), float(parts[1]), float(parts[2]), float(parts[3])
        except ValueError:
            return None

    def proximity_norm(self):
        """Devuelve cercanía F,L,R en [0,1]; 1 significa obstáculo próximo."""
        now = time.perf_counter()
        active = []
        with self.sensors_lock:
            for value in self.proximity.values():
                if now - value["t"] <= self.sensor_timeout_s:
                    active.append(value)
        if not active:
            return 0.0, 0.0, 0.0
        # VCNL4010: umbral 140, saturación 220, según la metodología.
        normalize = lambda z: max(0.0, min(1.0, (z - 140.0) / 80.0))
        return tuple(max(normalize(v[k]) for v in active) for k in ("f", "l", "r"))

    def parse_humidity_line(self, line: str):
        s = line.strip()
        if not s:
            return None

        # normalización básica
        s = s.replace(";", ",")
        parts = [p.strip() for p in s.split(",") if p.strip() != ""]
        if len(parts) < 2:
            return None

        # Prefijos opcionales
        p0 = parts[0].upper()
        if p0 in ("H", "HUM", "HUMID", "HUMEDAD", "HU"):
            if len(parts) < 3:
                return None
            try:
                sid = int(float(parts[1]))
                val = float(parts[2])
                return (sid, val)
            except Exception:
                return None

        # Sin prefijo: "sid,valor"
        if len(parts) == 2:
            try:
                sid = int(float(parts[0]))
                val = float(parts[1])
                return (sid, val)
            except Exception:
                return None

        # Cuatro campos numéricos son proximidad, no humedad.
        return None

    def humidity_raw_to_norm(self, raw: float) -> float:
        try:
            rmin = float(self.hraw_min_var.get())
            rmax = float(self.hraw_max_var.get())
        except Exception:
            rmin, rmax = 0.0, 1023.0
            self.hraw_min_var.set("0")
            self.hraw_max_var.set("1023")

        if rmax <= rmin:
            rmax = rmin + 1e-9

        n = (float(raw) - rmin) / (rmax - rmin)
        n = max(0.0, min(1.0, n))
        if bool(self.hinv_var.get()):
            n = 1.0 - n
        return max(0.0, min(1.0, n))

    def sensors_active_list(self):
        try:
            tout = max(0.2, float(self.s_timeout_var.get()))
        except Exception:
            tout = self.sensor_timeout_s
        self.sensor_timeout_s = tout

        now = time.perf_counter()
        out = []
        with self.sensors_lock:
            for sid, d in self.sensors.items():
                age = now - float(d.get("t", 0.0))
                if age <= tout:
                    out.append((sid, d.get("raw", 0.0), d.get("norm", 0.0), age))
        out.sort(key=lambda x: x[0])
        return out

    def sensors_avg_norm(self):
        act = self.sensors_active_list()
        if not act:
            return (None, 0)
        vals = [float(a[2]) for a in act]
        if not vals:
            return (None, 0)
        return (sum(vals) / len(vals), len(vals))

    def clear_sensors(self):
        with self.sensors_lock:
            self.sensors = {}
        self.s_avg_var.set("N/A")
        self.s_n_var.set("0")

    def _refresh_sensors_table(self):
        """
        Refresca tabla y promedio. Se ejecuta periódicamente con after().
        """
        try:
            act = self.sensors_active_list()
            avg, n = self.sensors_avg_norm()

            # promedio UI
            if avg is None:
                self.s_avg_var.set("N/A")
                self.s_n_var.set("0")
            else:
                self.s_avg_var.set(f"{avg:.3f}")
                self.s_n_var.set(str(n))

            # actualizar treeview
            # limpiar y repoblar (simple y robusto)
            for item in self.sensor_tv.get_children():
                self.sensor_tv.delete(item)

            for sid, raw, norm, age in act:
                self.sensor_tv.insert(
                    "", "end",
                    values=(int(sid), f"{float(raw):.3f}", f"{float(norm):.3f}", f"{float(age):.2f}")
                )
        except Exception:
            pass
        finally:
            self.after(250, self._refresh_sensors_table)

    # ================= HOME =================
    def home_all(self):
        if self.require_serial() and (not self.ser or not self.ser.is_open):
            messagebox.showwarning("HOME", "Conecte el Maestro o desactive 'Enviar comandos al robot'.")
            return

        self.stop_event.set()

        ids = parse_id_list(self.ids_var.get())
        if not ids:
            messagebox.showwarning("HOME", "IDs vacíos.")
            return

        if self.require_serial():
            lines = []
            for rid in ids:
                if rid <= 3:
                    lines.append(f"N {rid} POS {HOME_DX}")
                else:
                    lines.append(f"N {rid} LX {HOME_LX}")
            try:
                self.send_lines_batch(lines)
            except Exception as e:
                self.set_status(f"Error HOME: {e}")
                return

        self.apply_bounds_and_reinit()

        ax = float(self.ax_var.get())
        ay = float(self.ay_var.get())
        th0 = theta_from_alignment(self.align_var.get().strip())
        self.pose_runtime = (ax, ay, th0)

        self.traj_x = [ax]
        self.traj_y = [ay]
        self.t0 = time.perf_counter()

        self.distance_m = 0.0
        self._last_iter_dist = 0.0
        self._last_iter_time = 0.0

        # humedad inicial: sensores si hay, si no mapa
        hv0 = self.get_current_humidity_value(ax, ay)
        self.best_found = (ax, ay, hv0)
        self.search_iter = 0

        self.current_goal = (None, None)
        self.waypoints_runtime = []
        self.wp_index = 0

        self.last_cmd_pos = {}
        self.speed_acc = {}
        self.speed_history = []
        self.humidity_series = []
        self.robot_series = []
        self.last_humidity = hv0

        self.set_status("HOME enviado (si aplica) y pose/búsqueda reiniciadas (con límites).")

    # ================= Movimiento: líneas + posiciones =================
    def build_servo_lines_with_positions(self, t, freq, Av, Ah, phi_v_deg, phi_h_deg, turning, cw):
        sign_v = -1.0
        if turning:
            sign_h = sign_v if cw else -sign_v
        else:
            sign_h = sign_v

        Av_lx = Av * LX_PER_DX
        Ah_lx = Ah * LX_PER_DX

        ids = self.ids_runtime
        n = len(ids)
        if n <= 0:
            return [], {}

        horizontals = self.horizontal_ids_set()
        phase_origin_deg = 0.0

        out = []
        posmap = {}

        for i, rid in enumerate(ids):
            phi_idx_deg = phase_origin_deg + (360.0 * i / n)

            if rid in horizontals:
                phase = 2 * math.pi * freq * t + deg2rad(phi_h_deg + sign_h * phi_idx_deg)
                ss = math.sin(phase)
                if rid <= 3:
                    pos = clamp_int(HOME_DX + Ah * ss, DX_MIN, DX_MAX)
                    out.append(f"N {rid} POS {pos}")
                    posmap[rid] = float(pos)
                else:
                    ang = clamp_int(HOME_LX + Ah_lx * ss, LX_MIN, LX_MAX)
                    out.append(f"N {rid} LX {ang}")
                    posmap[rid] = float(ang)
            else:
                phase = 2 * math.pi * freq * t + deg2rad(phi_v_deg + sign_v * phi_idx_deg)
                ss = math.sin(phase)
                if rid <= 3:
                    pos = clamp_int(HOME_DX + Av * ss, DX_MIN, DX_MAX)
                    out.append(f"N {rid} POS {pos}")
                    posmap[rid] = float(pos)
                else:
                    ang = clamp_int(HOME_LX + Av_lx * ss, LX_MIN, LX_MAX)
                    out.append(f"N {rid} LX {ang}")
                    posmap[rid] = float(ang)

        return out, posmap

    # ================= Histórico: velocidades por módulo =================
    def init_speed_acc_if_needed(self):
        for rid in self.ids_runtime:
            if rid not in self.speed_acc:
                self.speed_acc[rid] = {"sum_abs": 0.0, "sum_dt": 0.0, "avg": 0.0}

    def update_speed_history(self, cmd_positions: dict, dt: float):
        if dt <= 0:
            return
        self.init_speed_acc_if_needed()

        for rid, pos in cmd_positions.items():
            if rid in self.last_cmd_pos:
                d = abs(pos - self.last_cmd_pos[rid])
                acc = self.speed_acc.setdefault(rid, {"sum_abs": 0.0, "sum_dt": 0.0, "avg": 0.0})
                acc["sum_abs"] += d
                acc["sum_dt"] += dt
                if acc["sum_dt"] > 1e-12:
                    acc["avg"] = acc["sum_abs"] / acc["sum_dt"]
            self.last_cmd_pos[rid] = pos

        now = time.perf_counter()
        if not hasattr(self, "_last_speed_snap_t"):
            self._last_speed_snap_t = now
        if (now - self._last_speed_snap_t) >= 0.5:
            self._last_speed_snap_t = now
            avg_by_id = {rid: round(self.speed_acc[rid]["avg"], 6) for rid in sorted(self.speed_acc.keys())}
            self.speed_history.append({
                "t": round(now - self.t0, 6),
                "avg_by_id": avg_by_id,
                "distance_m": round(self.distance_m, 6)
            })

    # ================= Pose + distancia + serie robot =================
    def update_pose_model(self, Av, Ah, turning, cw, kv, kw, dt):
        v = kv * max(0.0, Av - Ah)
        w = 0.0
        if turning:
            w = kw * max(0.0, Ah - Av) * (-1.0 if cw else +1.0)

        x, y, th = self.pose_runtime
        x_prev, y_prev = x, y

        th = th + w * dt
        x = x + v * math.cos(th) * dt
        y = y + v * math.sin(th) * dt

        x, y = self.clip_xy(x, y)

        dd = math.hypot(x - x_prev, y - y_prev)
        self.distance_m += dd

        self.pose_runtime = (x, y, th)
        self.traj_x.append(x)
        self.traj_y.append(y)

        now = time.perf_counter()
        if not hasattr(self, "_last_robot_snap_t"):
            self._last_robot_snap_t = now
        if (now - self._last_robot_snap_t) >= 0.5:
            self._last_robot_snap_t = now
            t = now - self.t0
            v_inst = dd / max(1e-9, dt)
            self.robot_series.append({"t": round(t, 6), "distance_m": round(self.distance_m, 6), "v_m_s": round(v_inst, 6)})

    # ================= Humedad efectiva (NUEVO) =================
    def get_current_humidity_value(self, x, y):
        """
        Retorna humedad efectiva (0..1):
        - Si hay sensores activos: promedio de sensores
        - Si no: mapa sintético (bilinear)
        """
        s_avg, s_n = self.sensors_avg_norm()
        if s_avg is not None and s_n > 0:
            return float(max(0.0, min(1.0, s_avg)))
        if self.hmap.is_ready():
            return float(self.hmap.value_bilinear(x, y))
        return 0.0

    def sample_humidity_at_pose(self):
        x, y, _ = self.pose_runtime
        hv = self.get_current_humidity_value(x, y)

        self.last_humidity = hv

        bx, by, bv = self.best_found
        if hv > bv:
            self.best_found = (x, y, hv)

        now = time.perf_counter()
        if not hasattr(self, "_last_hum_snap_t"):
            self._last_hum_snap_t = now
        if (now - self._last_hum_snap_t) >= 0.5:
            self._last_hum_snap_t = now
            self.humidity_series.append({"t": round(now - self.t0, 6), "x": round(x, 6), "y": round(y, 6),
                                         "h": round(hv, 6), "distance_m": round(self.distance_m, 6)})

    # ================= Ruta hacia objetivo =================
    def set_new_goal(self, gx, gy):
        gx, gy = self.clip_xy(gx, gy)
        self.current_goal = (gx, gy)

        x, y, _ = self.pose_runtime
        step_m = float(self.step_m_var.get())
        x_first = bool(self.x_first_var.get())
        self.start_A = (x, y)

        bounds = self.get_bounds()
        self.waypoints_runtime = build_manhattan_waypoints(x, y, gx, gy, step_m, x_first=x_first, bounds=bounds)
        self.wp_index = 0

    def current_wp(self):
        if 0 <= self.wp_index < len(self.waypoints_runtime):
            return self.waypoints_runtime[self.wp_index]
        return None

    def advance_wp(self):
        if self.wp_index + 1 < len(self.waypoints_runtime):
            self.wp_index += 1
            return True
        self.wp_index = len(self.waypoints_runtime)
        return False

    # ================= Fitness por iteración (histórico) =================
    def record_fitness_iteration(self):
        bx, by, bv = self.best_found
        if bx is None:
            return

        elapsed_s = float(time.perf_counter() - self.t0)
        dist_total = float(self.distance_m)

        dt = max(1e-9, elapsed_s - float(getattr(self, "_last_iter_time", 0.0)))
        dd = dist_total - float(getattr(self, "_last_iter_dist", 0.0))
        v_robot = dd / dt

        self._last_iter_time = elapsed_s
        self._last_iter_dist = dist_total

        s_avg, s_n = self.sensors_avg_norm()
        rna = self.search.rna_last or {}
        probs = rna.get("probabilities", [None, None, None, None])
        prox_f, prox_l, prox_r = self.proximity_norm()

        self.fitness_history.append({
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "iter": int(self.search_iter),
            "algorithm": self.algo_var.get(),
            "best_humidity": float(bv), "best_x": float(bx), "best_y": float(by),
            "humidity_now": float(self.last_humidity),
            "sensors_avg": None if s_avg is None else float(s_avg),
            "sensors_n": int(s_n),
            "distance_m": float(dist_total), "elapsed_s": float(elapsed_s),
            "delta_distance_m": float(dd), "delta_time_s": float(dt), "robot_speed_m_s": float(v_robot),
            "rna_action": rna.get("action"),
            "rna_p_AV": probs[0], "rna_p_GI": probs[1],
            "rna_p_GD": probs[2], "rna_p_EV": probs[3],
            "prox_F": prox_f, "prox_Izq": prox_l, "prox_Der": prox_r
        })

    def on_reach_goal(self):
        gx, gy = self.current_goal
        if gx is None:
            return

        self.sample_humidity_at_pose()
        self.record_fitness_iteration()

        self.search_iter += 1
        try:
            max_iters = int(self.max_iters_var.get())
        except Exception:
            max_iters = 120

        try:
            h_stop = float(self.stop_h_thresh_var.get())
        except Exception:
            h_stop = 0.995

        bx, by, bv = self.best_found
        force_max_iters = bool(self.force_max_iters_var.get())
        reached_limit = self.search_iter >= max_iters
        reached_humidity = bv >= h_stop

        if reached_limit or (reached_humidity and not force_max_iters):
            self.set_status(f"Búsqueda finalizada. Mejor h={bv:.3f} en ({bx:.2f},{by:.2f}).")
            self.current_goal = (None, None)
            self.waypoints_runtime = []
            self.wp_index = 0
            self.stop_event.set()
            return

        self.current_goal = (None, None)
        self.waypoints_runtime = []
        self.wp_index = 0

    # ================= Control del algoritmo =================
    def algo_next_goal(self):
        if not self.hmap.is_ready():
            self.set_status("Mapa no está listo.")
            return None

        self.apply_bounds_and_reinit()
        self.update_morphology_runtime()

        if self.search.algo != self.algo_var.get():
            self.search.set_algo(self.algo_var.get())

        g = self.search.next_goal()
        if g is None:
            return None

        gx, gy = self.clip_xy(g[0], g[1])
        return (gx, gy)

    # ================= Ejecución =================
    def start_search(self):
        if self.require_serial() and (not self.ser or not self.ser.is_open):
            messagebox.showwarning("Inicio", "Conecte el Maestro o desactive 'Enviar comandos al robot'.")
            return
        if not self.hmap.is_ready():
            messagebox.showwarning("Inicio", "Genere primero el mapa de humedad.")
            return

        self.update_morphology_runtime()
        self.apply_bounds_and_reinit()

        self.run_id = datetime.now().strftime("R%Y%m%d_%H%M%S")
        self.run_end_time = None
        self.run_config_snapshot = self.snapshot_config()
        self.run_config_snapshot["run_id"] = self.run_id
        self.run_config_snapshot["repetition"] = int(self.repetition_var.get())

        self.network.start({
            "run_id": self.run_id,
            "timestamp": self.run_config_snapshot["timestamp"],
            "algorithm": self.algo_var.get(),
            "repetition": int(self.repetition_var.get()),
            "robot_state": "Movimiento" if self.require_serial() else "Simulacion",
            "module_ids": ",".join(str(x) for x in self.ids_runtime),
            "num_modules": len(self.ids_runtime),
            "morphology": self.morph_var.get(),
            "port": self.port_var.get(), "baud": self.baud_var.get(),
            "channel": 1, "configured_motion_rate_hz": float(self.rate_var.get())
            , "configured_sensor_rate_hz": 4.0,
            "sensor_rate_by_node": {"1": 4.0, "3": 4.0, "4": 4.0,
                                    "5": 4.0, "6": 4.0, "7": 0.8,
                                    "8": 0.8, "9": 0.8}
        })
        if self.ser and self.ser.is_open:
            try:
                self.send_lines_batch([
                    f"RUN,{self.run_id},{self.algo_var.get().replace(',', '_')}"
                ])
            except Exception as exc:
                self.network.add_event("WARN", "RUN_CONTEXT", str(exc))

        self.fitness_history = []
        self.speed_acc = {}
        self.last_cmd_pos = {}
        self.speed_history = []
        self.humidity_series = []
        self.robot_series = []
        self.distance_m = 0.0
        self._last_iter_dist = 0.0
        self._last_iter_time = 0.0

        x0 = float(self.ax_var.get())
        y0 = float(self.ay_var.get())
        th0 = theta_from_alignment(self.align_var.get().strip())
        self.pose_runtime = (x0, y0, th0)
        self.traj_x = [x0]
        self.traj_y = [y0]
        self.t0 = time.perf_counter()

        hv0 = self.get_current_humidity_value(x0, y0)
        self.last_humidity = hv0
        self.best_found = (x0, y0, hv0)
        self.search_iter = 0

        self.current_goal = (None, None)
        self.waypoints_runtime = []
        self.wp_index = 0

        self.search.set_algo(self.algo_var.get())

        if self.nav_thread and self.nav_thread.is_alive():
            return

        self.stop_event.clear()
        self.nav_thread = NavigateThread(self, self.stop_event)
        self.nav_thread.start()
        self.set_status("Búsqueda iniciada.")

        if not (self.plot_win and self.plot_win.winfo_exists()):
            self.open_plot_window()
        else:
            self.plot_win.lift()

    def stop(self):
        self.stop_event.set()
        self.network.stop()
        if self.run_end_time is None:
            self.run_end_time = time.perf_counter()
        self.set_status("Detenido")

    def robot_avg_speed_m_s(self):
        end = self.run_end_time if self.run_end_time is not None else time.perf_counter()
        dur = end - self.t0
        if dur <= 1e-9:
            return 0.0
        return self.distance_m / dur

    # ================= Histórico: configuración =================
    def snapshot_config(self):
        ids = parse_id_list(self.ids_var.get())
        horiz = sorted(list(self.horizontal_ids_set()))
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.get_bounds()

        cfg = {
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "bounds": {"xmin": X_MIN, "xmax": X_MAX, "ymin": Y_MIN, "ymax": Y_MAX},
            "serial": {
                "port": self.port_var.get(),
                "baud": self.baud_var.get(),
                "eol": self.eol_var.get(),
                "send_enabled": bool(self.send_serial_var.get())
            },
            "modules": {
                "n": int(self.nmods_var.get()),
                "ids": ids,
                "horizontal_ids": horiz,
                "morphology": self.morph_var.get()
            },
            "navigation": {
                "A": {"x": float(self.ax_var.get()), "y": float(self.ay_var.get())},
                "alignment": self.align_var.get(),
                "step_m": float(self.step_m_var.get()),
                "x_first": bool(self.x_first_var.get()),
                "tol_m": float(self.tol_var.get()),
                "deadband_deg": float(self.deadband_deg_var.get()),
                "tick_axes_m": float(self.tick_var.get())
            },
            "motion": {
                "freq_hz": float(self.freq_var.get()),
                "rate_hz": float(self.rate_var.get()),
                "move": {"Av": float(self.move_v_ticks.get()), "Ah": float(self.move_h_ticks.get()),
                         "phiV_deg": float(self.move_phase_v_deg.get()), "phiH_deg": float(self.move_phase_h_deg.get())},
                "turn": {"Av": float(self.turn_v_ticks.get()), "Ah": float(self.turn_h_ticks.get()),
                         "phiV_deg": float(self.turn_phase_v_deg.get()), "phiH_deg": float(self.turn_phase_h_deg.get())},
                "phase_shift": {"enable": bool(self.turn_phase_shift_enable.get()),
                                "dir": self.turn_phase_shift_dir.get(),
                                "deg": float(self.turn_phase_shift_deg.get())},
                "pose_model": {"kv": float(self.kv_var.get()), "kw": float(self.kw_var.get())}
            },
            "humidity_map": {
                "nx": int(self.hm_nx_var.get()), "ny": int(self.hm_ny_var.get()),
                "seed": int(self.hm_seed_var.get()), "noise": float(self.hm_noise_var.get()),
                "peak_x": float(self.hm_peakx_var.get()), "peak_y": float(self.hm_peaky_var.get()),
                "sigma": float(self.hm_sigma_var.get())
            },
            "sensors": {
                "raw_min": float(self.hraw_min_var.get() or 0),
                "raw_max": float(self.hraw_max_var.get() or 1023),
                "invert": bool(self.hinv_var.get()),
                "timeout_s": float(self.s_timeout_var.get() or 5.0)
            },
            "search": {
                "algorithm": self.algo_var.get(),
                "max_iters": int(self.max_iters_var.get()),
                "stop_threshold": float(self.stop_h_thresh_var.get()),
                "force_max_iterations": bool(self.force_max_iters_var.get()),
                "params": {
                    "hc_step": float(self.hc_step_var.get()),
                    "ga_pop": int(self.ga_pop_var.get()),
                    "ga_elite": int(self.ga_elite_var.get()),
                    "ga_mut_sigma": float(self.ga_mut_var.get()),
                    "cg_k": int(self.cg_k_var.get()),
                    "cg_radius": float(self.cg_radius_var.get()),
                    "bc_step": float(self.bc_step_var.get()),
                    "bc_maxrun": int(self.bc_maxrun_var.get()),
                    "sa_T0": float(self.sa_T0_var.get()),
                    "sa_alpha": float(self.sa_alpha_var.get()),
                    "sa_step": float(self.sa_step_var.get()),
                    "rw_step": float(self.rw_step_var.get())
                }
            }
        }
        return cfg

    # ================= Exportar / limpiar histórico =================
    def clear_history(self):
        self.fitness_history = []
        self.speed_history = []
        self.speed_acc = {}
        self.last_cmd_pos = {}
        self.humidity_series = []
        self.robot_series = []
        self.run_config_snapshot = None
        self.run_id = None
        self.run_end_time = None
        self.network = NetworkMeasurement()
        self.set_status("Histórico limpiado.")

    def _autosize_ws(self, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    v = "" if cell.value is None else str(cell.value)
                    max_len = max(max_len, len(v))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(60, max(10, max_len + 2))

    def export_history(self):
        if self.run_config_snapshot is None and not self.fitness_history and not self.speed_history:
            messagebox.showinfo("Exportar", "No hay histórico para exportar.")
            return

        folder = filedialog.askdirectory(title="Seleccione carpeta para exportar")
        if not folder:
            return

        run_id = self.run_id or datetime.now().strftime("R%Y%m%d_%H%M%S")
        method_slug = "".join(c if c.isalnum() else "_" for c in self.algo_var.get()).strip("_")
        base = os.path.join(folder, f"resultados_algoritmo_{method_slug}_{run_id}")

        # snapshot de sensores (último estado)
        with self.sensors_lock:
            sensors_snapshot = {str(k): {"raw": v.get("raw"), "norm": v.get("norm"), "t": v.get("t")} for k, v in self.sensors.items()}

        data = {
            "config": self.run_config_snapshot,
            "fitness_history": self.fitness_history,
            "speed_history": self.speed_history,
            "speed_avg_final_by_id": {str(rid): self.speed_acc[rid]["avg"] for rid in self.speed_acc} if self.speed_acc else {},
            "distance_m_final": self.distance_m,
            "robot_avg_speed_m_s": self.robot_avg_speed_m_s(),
            "best_found": {"x": self.best_found[0], "y": self.best_found[1], "h": self.best_found[2]},
            "humidity_series": self.humidity_series,
            "robot_series": self.robot_series,
            "sensors_last": sensors_snapshot
        }
        json_path = base + ".json"
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

        xlsx_path = base + ".xlsx"
        wb = Workbook()

        ws = wb.active
        ws.title = "Config"
        ws["A1"] = "config_json"
        ws["A2"] = json.dumps(self.run_config_snapshot, ensure_ascii=False, indent=2) if self.run_config_snapshot else ""
        self._autosize_ws(ws)

        ws_it = wb.create_sheet("Iteraciones")
        headers_it = [
            "timestamp", "iter", "algorithm",
            "best_humidity", "best_x", "best_y",
            "humidity_now",
            "sensors_avg", "sensors_n",
            "distance_m", "elapsed_s",
            "delta_distance_m", "delta_time_s", "robot_speed_m_s",
            "rna_action", "rna_p_AV", "rna_p_GI", "rna_p_GD", "rna_p_EV",
            "prox_F", "prox_Izq", "prox_Der"
        ]
        ws_it.append(headers_it)
        for r in self.fitness_history:
            ws_it.append([r.get(h, None) for h in headers_it])
        self._autosize_ws(ws_it)

        ws_sp = wb.create_sheet("SpeedSeries")
        ids = sorted(self.speed_acc.keys()) if self.speed_acc else []
        headers_sp = ["t", "distance_m"] + [f"id_{rid}_avg_ticks_per_s" for rid in ids]
        ws_sp.append(headers_sp)
        for s in self.speed_history:
            row = [s.get("t", None), s.get("distance_m", None)]
            avg_by_id = s.get("avg_by_id", {})
            for rid in ids:
                row.append(avg_by_id.get(rid, None))
            ws_sp.append(row)
        self._autosize_ws(ws_sp)

        ws_sf = wb.create_sheet("SpeedAvgFinal")
        ws_sf.append(["id", "avg_ticks_per_s"])
        for rid in ids:
            ws_sf.append([rid, float(self.speed_acc[rid]["avg"])])
        self._autosize_ws(ws_sf)

        ws_rs = wb.create_sheet("RobotSeries")
        ws_rs.append(["t", "distance_m", "v_m_s"])
        for s in self.robot_series:
            ws_rs.append([s.get("t"), s.get("distance_m"), s.get("v_m_s")])
        self._autosize_ws(ws_rs)

        ws_hs = wb.create_sheet("HumiditySeries")
        ws_hs.append(["t", "x", "y", "h", "distance_m"])
        for s in self.humidity_series:
            ws_hs.append([s.get("t"), s.get("x"), s.get("y"), s.get("h"), s.get("distance_m")])
        self._autosize_ws(ws_hs)

        # Sensores (último estado)
        ws_s = wb.create_sheet("SensoresLast")
        ws_s.append(["sensor_id", "raw", "norm", "age_s"])
        now = time.perf_counter()
        with self.sensors_lock:
            for sid in sorted(self.sensors.keys()):
                d = self.sensors[sid]
                age = now - float(d.get("t", 0.0))
                ws_s.append([sid, d.get("raw"), d.get("norm"), age])
        self._autosize_ws(ws_s)

        ws_sum = wb.create_sheet("Resumen")
        bx, by, bh = self.best_found
        end = self.run_end_time if self.run_end_time is not None else time.perf_counter()
        duration = end - self.t0
        robot_avg_speed = (self.distance_m / duration) if duration > 1e-9 else 0.0
        X_MIN, X_MAX, Y_MIN, Y_MAX = self.get_bounds()
        s_avg, s_n = self.sensors_avg_norm()

        ws_sum.append(["metric", "value"])
        ws_sum.append(["bounds_xmin", X_MIN])
        ws_sum.append(["bounds_xmax", X_MAX])
        ws_sum.append(["bounds_ymin", Y_MIN])
        ws_sum.append(["bounds_ymax", Y_MAX])
        ws_sum.append(["algorithm", self.algo_var.get()])
        ws_sum.append(["iterations", self.search_iter])
        ws_sum.append(["best_humidity", bh])
        ws_sum.append(["best_x", bx])
        ws_sum.append(["best_y", by])
        final_humidity = self.humidity_series[-1].get("h") if self.humidity_series else self.last_humidity
        ws_sum.append(["final_humidity", final_humidity])
        ws_sum.append(["distance_m_total", self.distance_m])
        ws_sum.append(["duration_s", duration])
        ws_sum.append(["robot_avg_speed_m_s", robot_avg_speed])
        ws_sum.append(["num_modules", len(self.ids_runtime)])
        ws_sum.append(["sensors_avg_now", None if s_avg is None else float(s_avg)])
        ws_sum.append(["sensors_n_now", int(s_n)])
        self._autosize_ws(ws_sum)

        wb.save(xlsx_path)

        comm_path = os.path.join(folder, f"resultados_comunicacion_{method_slug}_{run_id}.xlsx")
        try:
            self._export_network_excel(comm_path)
        except Exception as exc:
            self.set_status(f"Error exportando comunicación: {exc}")
            messagebox.showerror(
                "Exportación de comunicación",
                f"El archivo principal se guardó, pero falló el Excel de comunicación:\n{exc}"
            )
            return
        self.set_status(
            f"Exportados: {os.path.basename(xlsx_path)} y {os.path.basename(comm_path)} (+ JSON)."
        )

    def _export_network_excel(self, path):
        packets, events, config = self.network.snapshot()
        summary = self.network.summaries()
        wb = Workbook()

        ws_cfg = wb.active
        ws_cfg.title = "Configuracion"
        ws_cfg.append(["Parametro", "Valor"])
        for key, value in config.items():
            # Excel no admite dict, list, tuple o set directamente en una celda.
            if isinstance(value, (dict, list, tuple, set)):
                value = json.dumps(value, ensure_ascii=False, sort_keys=True)
            ws_cfg.append([key, value])
        ws_cfg.append(["duration_effective_s", self.network.elapsed()])
        self._autosize_ws(ws_cfg)

        ws_p = wb.create_sheet("Paquetes")
        packet_headers = list(asdict(NetPacket("", 0, None, 0, 0.0, None)).keys())
        ws_p.append(packet_headers)
        for packet in packets:
            row = asdict(packet)
            ws_p.append([row.get(h) for h in packet_headers])
        self._autosize_ws(ws_p)

        ws_n = wb.create_sheet("Resumen_nodo")
        summary_headers = list(summary[0].keys()) if summary else [
            "run_id", "algorithm", "repetition", "robot_state", "node_id"
        ]
        ws_n.append(summary_headers)
        for item in summary:
            ws_n.append([item.get(h) for h in summary_headers])
        self._autosize_ws(ws_n)

        ws_e = wb.create_sheet("Eventos")
        event_headers = list(asdict(NetEvent(0.0, "", None, "", "")).keys())
        ws_e.append(event_headers)
        for event in events:
            row = asdict(event)
            ws_e.append([row.get(h) for h in event_headers])
        self._autosize_ws(ws_e)

        ws_s = wb.create_sheet("Series_nodo")
        ws_s.append(["run_id", "algorithm", "repetition", "node_id", "sample",
                     "elapsed_s", "interarrival_ms", "rtt_ms", "rssi_dbm",
                     "packet_type"])
        sample_by_node = defaultdict(int)
        for packet in packets:
            sample_by_node[packet.node_id] += 1
            ws_s.append([packet.run_id, config.get("algorithm"), config.get("repetition"),
                         packet.node_id, sample_by_node[packet.node_id], packet.elapsed_s,
                         packet.interarrival_ms, packet.rtt_ms, packet.rssi_dbm,
                         packet.packet_type])
        self._autosize_ws(ws_s)
        wb.save(path)

    # ================= Cierre =================
    def on_close(self):
        try:
            self.disconnect()
        except Exception:
            pass
        self.destroy()


if __name__ == "__main__":
    App().mainloop()
